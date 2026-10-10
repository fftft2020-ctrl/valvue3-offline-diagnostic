from __future__ import annotations

import json
import hashlib
import shutil
from pathlib import Path

import fitz
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from offline_engine import pipeline
from offline_engine.format_revision import (
    compare_format_only_content,
    docx_content_identity,
)


def _write_project(project: Path) -> None:
    (project / "raw").mkdir(parents=True)
    (project / "project.toml").write_text(
        """
[project]
customer = "Demo"
report_title = "Report"
report_date = "2026-07-25"
[report_periods]
active = "completed"
completed_label = "2026年7月18日至7月21日"
[paths]
raw = "raw"
work = "work/shared_engine"
output = "output"
output_docx = "legacy.docx"
output_pdf = "legacy.pdf"
""".strip(),
        encoding="utf-8",
    )


def _write_docx(path: Path, *, text: str = "已确认的诊断内容") -> None:
    document = Document()
    document.add_heading("FV-101 | 设备概览", level=1)
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "评级结论"
    table.cell(0, 1).text = "3级，有缺陷"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def _write_wps_pdf(path: Path) -> None:
    document = fitz.open()
    page = document.new_page(width=595.28, height=841.89)
    page.insert_text((72, 100), "ValVue3 Offline Diagnostic Report")
    document.set_metadata({"creator": "WPS 文字"})
    document.save(path)
    document.close()


def test_docx_content_identity_ignores_formatting_but_rejects_text_changes(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.docx"
    formatted = tmp_path / "formatted.docx"
    changed = tmp_path / "changed.docx"
    _write_docx(source)
    shutil.copy2(source, formatted)
    document = Document(formatted)
    document.paragraphs[1].paragraph_format.space_after = Pt(3)
    document.tables[0].rows[0].cells[1].paragraphs[0].runs[0].font.size = Pt(8)
    document.save(formatted)
    _write_docx(changed, text="被改动的诊断内容")

    assert docx_content_identity(source) == docx_content_identity(formatted)
    comparison = compare_format_only_content(source, formatted)
    assert comparison["ok"] is True
    assert compare_format_only_content(source, changed)["ok"] is False


def test_prepare_format_only_revision_copies_parent_without_regeneration(
    tmp_path: Path,
    monkeypatch,
) -> None:
    project = tmp_path / "DemoProject"
    _write_project(project)
    source = (
        project
        / "output/review"
        / "DemoProject-2026年7月18日至7月21日-人工校验版-v3.docx"
    )
    _write_docx(source)

    def forbidden_run_project(*args, **kwargs):
        raise AssertionError("format-only flow must not run report generation")

    monkeypatch.setattr(pipeline, "run_project", forbidden_run_project)

    result = pipeline.prepare_format_only_revision(
        project,
        source_docx=source,
        report_period="completed",
    )

    target = result.config.output_docx
    assert result.status == "format_revision_prepared"
    assert target.name.endswith("待人工校验版-v4.docx")
    assert target.read_bytes() == source.read_bytes()
    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert manifest["status"] == "format_revision_prepared"
    assert manifest["parent_docx"] == str(source.resolve())
    assert manifest["candidate_docx"] == str(target.resolve())
    assert manifest["content_identity"]["ok"] is True


def test_prepare_format_only_revision_accepts_configured_root_output(
    tmp_path: Path,
) -> None:
    project = tmp_path / "DemoProject"
    _write_project(project)
    source = project / "output/legacy.docx"
    _write_docx(source)

    result = pipeline.prepare_format_only_revision(
        project,
        source_docx=source,
        report_period="completed",
    )

    assert result.config.output_docx.parent == project / "output/business"
    assert result.config.output_docx.name.endswith("最终商务版-v1.docx")
    assert result.config.output_docx.read_bytes() == source.read_bytes()


def test_wps_delivery_rejects_content_change_in_format_only_revision(
    tmp_path: Path,
    monkeypatch,
) -> None:
    project = tmp_path / "DemoProject"
    _write_project(project)
    source = project / "output/review/DemoProject-人工校验版-v1.docx"
    candidate = project / "output/review/DemoProject-人工校验版-v2.docx"
    _write_docx(source)
    _write_docx(candidate, text="诊断内容已被修改")
    pdf = candidate.with_suffix(".pdf")
    _write_wps_pdf(pdf)
    monkeypatch.setattr(
        pipeline,
        "validate_report",
        lambda *args, **kwargs: {"ok": True, "failures": []},
    )

    with pytest.raises(RuntimeError, match="format-only content mismatch"):
        pipeline.validate_wps_delivery(
            project,
            docx=candidate,
            wps_pdf=pdf,
            format_only_source=source,
        )


def test_prepared_format_revision_auto_binds_parent_during_wps_validation(
    tmp_path: Path,
    monkeypatch,
) -> None:
    project = tmp_path / "DemoProject"
    _write_project(project)
    source = (
        project
        / "output/review"
        / "DemoProject-2026年7月18日至7月21日-人工校验版-v1.docx"
    )
    _write_docx(source)
    prepared = pipeline.prepare_format_only_revision(
        project,
        source_docx=source,
        report_period="completed",
    )
    candidate = prepared.config.output_docx
    document = Document(candidate)
    document.paragraphs[1].paragraph_format.space_after = Pt(2)
    document.save(candidate)
    pdf = candidate.with_suffix(".pdf")
    _write_wps_pdf(pdf)
    monkeypatch.setattr(
        pipeline,
        "validate_report",
        lambda *args, **kwargs: {"ok": True, "failures": []},
    )

    result = pipeline.validate_wps_delivery(
        project,
        docx=candidate,
        wps_pdf=pdf,
        report_period="completed",
    )

    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert result.status == "wps_layout_review_required"
    assert manifest["format_only_validation"]["ok"] is True
    assert manifest["format_only_validation"]["source"] == str(source.resolve())


def test_report_cli_exposes_format_only_revision_options() -> None:
    script = Path(__file__).resolve().parents[1] / "run_offline_report.py"
    result = __import__("subprocess").run(
        [__import__("sys").executable, str(script), "--help"],
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0
    assert "--prepare-format-only" in result.stdout
    assert "--format-only-source" in result.stdout


@pytest.mark.parametrize("change", ["anchor", "bookmark_end", "bookmark_range"])
def test_format_only_protects_internal_link_and_bookmark_scope(tmp_path: Path, change: str) -> None:
    source, candidate = tmp_path / "source.docx", tmp_path / "candidate.docx"
    document = Document()
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "valve_one")
    document.add_paragraph("跳转至阀门")._p.append(link)
    cell = document.add_table(rows=1, cols=2).cell(0, 0)
    cell.text = "原始章节"
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:name"), "valve_one")
    start.set(qn("w:id"), "7")
    start.set(qn("w:colFirst"), "0")
    start.set(qn("w:colLast"), "0")
    end.set(qn("w:id"), "7")
    cell.paragraphs[0]._p.insert(0, start)
    cell.paragraphs[0]._p.append(end)
    document.save(source)
    # Word processors may renumber paired IDs without changing the bookmark.
    start.set(qn("w:id"), "12")
    end.set(qn("w:id"), "12")
    document.save(candidate)
    assert compare_format_only_content(source, candidate)["ok"] is True

    if change == "anchor":
        link.set(qn("w:anchor"), "valve_two")
    elif change == "bookmark_end":
        end.getparent().remove(end)
    else:
        start.set(qn("w:colLast"), "1")
    document.save(candidate)
    assert compare_format_only_content(source, candidate)["ok"] is False


def test_format_only_rejects_parent_changed_to_match_child(tmp_path: Path) -> None:
    project = tmp_path / "DemoProject"
    _write_project(project)
    source = project / "output/review/confirmed-v1.docx"
    _write_docx(source)
    prepared = pipeline.prepare_format_only_revision(project, source_docx=source)
    candidate = prepared.config.output_docx
    _write_docx(source, text="母版与候选被同时改写")
    shutil.copy2(source, candidate)
    pdf = candidate.with_suffix(".pdf")
    _write_wps_pdf(pdf)

    with pytest.raises(RuntimeError, match="parent changed"):
        pipeline.validate_wps_delivery(project, docx=candidate, wps_pdf=pdf)
    manifest = project / "work/wps-validation" / candidate.stem / "delivery_manifest.json"
    assert json.loads(manifest.read_text())["status"] == "wps_delivery_validation_failed"


@pytest.mark.parametrize("failure", [
    "content_exception", "content_rejection", "layout_exception", "missing_pdf", "input_changed",
])
def test_failed_wps_revalidation_revokes_success_and_preserves_history(
    tmp_path: Path, monkeypatch, failure: str,
) -> None:
    project = tmp_path / "DemoProject"
    _write_project(project)
    docx = project / "output/review/report.docx"
    _write_docx(docx)
    pdf = docx.with_suffix(".pdf")
    _write_wps_pdf(pdf)
    monkeypatch.setattr(pipeline, "validate_report", lambda *args, **kwargs: {"ok": True})
    prior = pipeline.validate_wps_delivery(
        project, docx=docx, wps_pdf=pdf, approve_visual=True,
        visual_review_note="匿名夹具的审核记录",
    )
    previous = prior.manifest_path.read_bytes()
    old_sheet = Path(json.loads(previous)["automatic_layout"]["contact_sheet"])
    old_sheet_bytes = old_sheet.read_bytes()
    _write_docx(docx, text="本次待校验内容")

    def fail(*args, **kwargs):
        raise RuntimeError("injected validation failure")

    if failure == "content_exception":
        monkeypatch.setattr(pipeline, "validate_report", fail)
    elif failure == "content_rejection":
        monkeypatch.setattr(pipeline, "validate_report", lambda *args, **kwargs: {
            "ok": False, "summary_issues": ["关键诊断缺失"],
        })
    elif failure == "layout_exception":
        monkeypatch.setattr(pipeline, "inspect_wps_pdf_layout", fail)
    elif failure == "missing_pdf":
        pdf.unlink()
    else:
        def change_during_validation(*args, **kwargs):
            _write_docx(docx, text="校验过程中修改了文件")
            return {"ok": True}
        monkeypatch.setattr(pipeline, "validate_report", change_during_validation)

    with pytest.raises(RuntimeError):
        pipeline.validate_wps_delivery(project, docx=docx, wps_pdf=pdf)
    manifest = json.loads(prior.manifest_path.read_text())
    assert manifest["status"] == "wps_delivery_validation_failed"
    assert manifest["delivery"] == {}
    assert manifest["visual_review"]["approved"] is False
    assert manifest["inputs"]["docx"] == str(docx)
    history = prior.manifest_path.parent / "history" / f"{hashlib.sha256(previous).hexdigest()}.json"
    assert history.read_bytes() == previous
    assert old_sheet.read_bytes() == old_sheet_bytes
