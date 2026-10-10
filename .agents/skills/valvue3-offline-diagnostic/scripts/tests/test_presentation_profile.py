from dataclasses import replace
import json
from pathlib import Path

import pytest
from docx import Document

from offline_engine import business_report, generator, pipeline
from offline_engine.config import ConfigError, load_project


def _config(tmp_path: Path):
    (tmp_path / "project.toml").write_text(
        '[project]\ncustomer="Demo Customer"\nreport_title="Demo Report"\n'
        'report_date="2026-07-22"\nunit_name="Demo Unit"\n'
        '[paths]\noutput_docx="report.docx"\n',
        encoding="utf-8",
    )
    return load_project(tmp_path)


@pytest.mark.parametrize("profile", [None, "business_dashboard"])
def test_default_and_explicit_profile_use_current_renderer(monkeypatch, tmp_path: Path, profile) -> None:
    captured = {}

    def fake_renderer(valves, **kwargs):
        captured["valves"] = valves
        captured.update(kwargs)

    monkeypatch.setattr(business_report, "generate_business_dashboard_docx", fake_renderer)
    config = _config(tmp_path)
    config = replace(config, presentation={**config.presentation, "profile": profile})
    valves = [
        generator.ValveReport(
            folder="FV-101",
            rating=1,
            final_rating=1,
            rating_conclusion_override="冻结评级结论。",
            overall_assessment="冻结综合诊断结论。",
            actions=["冻结建议。"],
        )
    ]

    pipeline._render_prepared_report(config, valves)

    assert captured["valves"] == valves
    assert captured["output_path"] == config.output_docx
    assert captured["project_name"] == "Demo Customer"
    assert captured["report_name"] == "Demo Report"
    assert captured["report_date"] == "2026-07-22"
    assert captured["unit_name"] == "Demo Unit"
    assert captured["rating_conclusions"] == {
        "FV-101": "冻结评级结论。"
    }
    assert "本批共诊断1台阀门：1级1台" in captured["overall_conclusion"]
    assert "冻结综合诊断结论。" not in captured["overall_conclusion"]
    assert "对应单阀章节" in captured["recommendations"][0]


def test_batch_summary_does_not_concatenate_valve_prose(monkeypatch, tmp_path):
    captured = {}
    monkeypatch.setattr(
        business_report, "generate_business_dashboard_docx",
        lambda valves, **kwargs: captured.update(kwargs),
    )
    config = _config(tmp_path)
    valves = [generator.ValveReport(
        folder=f"FV-{index}", rating=grade, final_rating=grade,
        overall_assessment="单阀长结论" * 200,
        actions=["单阀专用建议" * 100],
    ) for index, grade in enumerate([2] * 5 + [3] * 56 + [4] * 16)]
    pipeline._render_prepared_report(config, valves)
    assert "2级5台、3级56台、4级16台" in captured["overall_conclusion"]
    assert len(captured["overall_conclusion"]) < 200
    assert "单阀专用建议" not in "".join(captured["recommendations"])
    config = replace(config, presentation={
        **config.presentation, "overall_conclusion_override": "已核对的整批摘要。",
        "recommendation_overrides": ["已核对的整批建议。"],
    })
    pipeline._render_prepared_report(config, valves)
    assert captured["overall_conclusion"] == "已核对的整批摘要。"
    assert captured["recommendations"] == ["已核对的整批建议。"]


def test_unknown_presentation_profile_is_rejected(tmp_path: Path) -> None:
    config = replace(_config(tmp_path), presentation={"profile": "unknown"})
    with pytest.raises(ValueError, match="unknown presentation profile"):
        pipeline._render_prepared_report(config, [])


def test_retired_standard_renderer_fails_without_creating_output(tmp_path: Path) -> None:
    config = replace(_config(tmp_path), presentation={"profile": "standard"})
    with pytest.raises(ValueError, match="retired.*business_dashboard"):
        pipeline._render_prepared_report(config, [])
    assert not config.output_docx.exists()


def test_two_project_configs_render_isolated_docx_in_one_process(
    tmp_path: Path, monkeypatch
) -> None:
    projects = []
    for name, tag, data_only in (
        ("Alpha", "FV-A", True),
        ("Beta", "FV-B", True),
    ):
        project = tmp_path / name
        (project / "raw").mkdir(parents=True)
        (project / "project.toml").write_text(
            f'''[project]
customer = "{name} Customer"
report_title = "{name} Report"
report_date = "2026-09-05"
[paths]
raw = "raw"
work = "work"
output_docx = "{name}.docx"
[report_layout]
only_valves = ["{tag}"]
[presentation]
data_only = {str(data_only).lower()}
blank_ratings = {str(data_only).lower()}
blank_conclusions = {str(data_only).lower()}
blank_actions = {str(data_only).lower()}
''',
            encoding="utf-8",
        )
        projects.append(project)

    valves = [
        generator.ValveReport(
            folder="FV-A", rating=3, final_rating=3,
            rating_conclusion_override="Alpha diagnosis",
            main_issue="Alpha issue", action="Alpha action",
        ),
        generator.ValveReport(
            folder="FV-B", rating=1, final_rating=1,
            rating_conclusion_override="Beta diagnosis",
            main_issue="Beta issue", action="Beta action",
        ),
    ]
    monkeypatch.setattr(generator, "prepare_valve_data", lambda **_kwargs: valves)
    monkeypatch.setattr(pipeline, "load_specification_pages", lambda *args: {})
    monkeypatch.setattr(pipeline, "validate_report", lambda *args, **kwargs: {"ok": True})

    results = [pipeline.run_project(project, skip_pdf=True) for project in projects]

    documents = [Document(result.config.output_docx) for result in results]
    visible = [
        "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        for document in documents
    ]
    assert "Alpha Report" in visible[0] and "FV-A" in visible[0]
    assert "FV-B" not in visible[0] and "Beta Customer" not in visible[0]
    assert "Beta Report" in visible[1] and "FV-B" in visible[1]
    assert "FV-A" not in visible[1] and "Alpha Customer" not in visible[1]
    assert "诊断结论" not in visible[0]
    assert "诊断结论" not in visible[1]
    for result, expected in zip(results, ("FV-A", "FV-B")):
        summary = json.loads(
            (result.config.work_dir / "analysis_summary.json").read_text(encoding="utf-8")
        )
        assert [item["folder"] for item in summary if item["included_in_report"]] == [expected]
        from formal_diagnosis_tools.single_valve_workflow import select_project_targets

        assert [item["folder"] for item in select_project_targets(summary)] == [expected]


@pytest.mark.parametrize("retired_config", [
    '[presentation]\nprofile="standard"',
    '[report_layout]\ntoc_append_label="old batch"',
    '[report_periods]\nactive="batch"\nbatch_label="batch"\n'
    '[report_period_overrides.batch]\ntoc_append_label="old batch"',
])
def test_preflight_rejects_retired_report_configuration(tmp_path: Path, retired_config: str) -> None:
    (tmp_path / "project.toml").write_text(
        '[project]\ncustomer="Demo"\nreport_title="Report"\nreport_date="2026-09-05"\n'
        '[paths]\noutput_docx="report.docx"\n' + retired_config,
        encoding="utf-8",
    )
    with pytest.raises(ConfigError, match="retired"):
        load_project(tmp_path)
