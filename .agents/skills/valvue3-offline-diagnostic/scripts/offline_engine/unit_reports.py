"""Assemble device-unit volumes from archived, unchanged DOCX chapters."""

from __future__ import annotations

import hashlib
import json
import re
import tomllib
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docxcompose.composer import Composer
import fitz

from .business_report import RATING_COLORS
from .report_text import missing_pdf_body_text, normalize_report_text


def _hash(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def _header(element, document, tags: set[str]) -> tuple[str, int] | None:
    if element.tag != qn("w:tbl"):
        return None
    table = Table(element, document)
    if len(table.rows) != 1 or len(table.columns) < 2:
        return None
    tag = (table.cell(0, 0).text.strip().splitlines() or [""])[0].strip()
    if tag not in tags:
        return None
    rating = re.search(r"(?<!\d)([1-5])级", table.cell(0, 1).text)
    if not rating:
        return None
    return tag, int(rating.group(1))


def _chapters(document, tags: set[str]) -> dict[str, tuple[int, int, int]]:
    body = list(document.element.body)
    starts = []
    for index, element in enumerate(body):
        found = _header(element, document, tags)
        if found:
            tag, rating = found
            start = index
            if index and body[index - 1].tag == qn("w:p"):
                previous = body[index - 1]
                visible = "".join(previous.xpath(".//w:t/text()"))
                if not visible.strip("\u200b \t\n") and (
                    previous.xpath('.//w:br[@w:type="page"]')
                    or previous.xpath(".//w:pageBreakBefore")
                ):
                    start -= 1
            starts.append((start, index, tag, rating))
    result = {}
    for position, (start, header_index, tag, rating) in enumerate(starts):
        if tag in result:
            raise ValueError(f"duplicate valve chapter: {tag}")
        stop = starts[position + 1][0] if position + 1 < len(starts) else len(body) - 1
        result[tag] = (start, stop, rating)
    if set(result) != tags:
        raise ValueError(f"chapter coverage mismatch: missing={sorted(tags - set(result))}")
    return result


def _chapter_content(document, start: int, stop: int) -> tuple[str, Counter[str]]:
    """Keep the human wording and each embedded image unchanged across composition."""
    nodes = list(document.element.body)[start:stop]
    wording = "\n".join("".join(node.xpath(".//w:t/text()")) for node in nodes)
    images: Counter[str] = Counter()
    for node in nodes:
        for blip in node.xpath(".//a:blip"):
            relation = blip.get(qn("r:embed"))
            if relation:
                blob = document.part.related_parts[relation].blob
                images[hashlib.sha256(blob).hexdigest()] += 1
    return wording, images


def _shade(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _subset(document, chapters, tags: set[str]) -> Document:
    body = document.element.body
    nodes = list(body)
    keep = {id(nodes[index]) for tag in tags for index in range(chapters[tag][0], chapters[tag][1])}
    for node in nodes:
        if node.tag != qn("w:sectPr") and id(node) not in keep:
            body.remove(node)
    for node in body:
        found = _header(node, document, tags)
        if found:
            table = Table(node, document)
            _shade(table.cell(0, 1), RATING_COLORS[found[1]])
    relation_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    used = {value for node in document.element.iter() for name, value in node.attrib.items()
            if name.startswith(relation_ns)}
    for rid, relation in list(document.part.rels.items()):
        if relation.reltype.endswith("/image") and rid not in used:
            del document.part.rels[rid]
    return document


def _add_front(document, title: str, unit: str, ordered: list[str], ratings: dict[str, int]) -> None:
    body = document.element.body
    front = []

    def paragraph(text: str, style: str | None = None):
        item = document.add_paragraph(text, style=style)
        front.append(item._p)
        return item

    paragraph(title, "Title")
    paragraph(f"{unit}｜{len(ordered)}台阀门")
    paragraph("报告日期：" + datetime.now(ZoneInfo("Asia/Shanghai")).date().isoformat())
    paragraph("诊断总览", "Heading 1")
    counts = Counter(ratings[tag] for tag in ordered)
    paragraph("本装置共诊断{}台；{}。".format(
        len(ordered), "、".join(f"{level}级{counts[level]}台" for level in range(1, 6) if counts[level])
    ))
    paragraph("各阀诊断依据、问题和处置建议见对应章节。")
    paragraph("阀门目录", "Heading 1")
    listing = document.add_table(rows=1, cols=3)
    front.append(listing._tbl)
    listing.style = "Table Grid"
    for cell, value in zip(listing.rows[0].cells, ("序号", "位号", "等级")):
        cell.text = value
    listing.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for number, tag in enumerate(ordered, 1):
        cells = listing.add_row().cells
        cells[0].text = str(number)
        cells[1].text = tag
        cells[2].text = f"{ratings[tag]}级"
        _shade(cells[2], RATING_COLORS[ratings[tag]])
    for item in enumerate(front):
        index, node = item
        body.remove(node)
        body.insert(index, node)


def assemble_unit_reports(project_dir: Path, selection: Path) -> Path:
    """Create one DOCX per source unit, retaining all selected chapters."""
    project_dir = project_dir.resolve()
    config = tomllib.loads((project_dir / "project.toml").read_text(encoding="utf-8"))
    selection = selection.resolve()
    if not selection.is_relative_to(project_dir / "work"):
        raise ValueError("unit source selection must be under project work")
    request = json.loads(selection.read_text(encoding="utf-8"))
    title = str(request.get("title") or "").strip()
    if not title:
        raise ValueError("unit report title is required")
    context_path = (project_dir / request["unit_context"]).resolve()
    if not context_path.is_relative_to(project_dir / "experience"):
        raise ValueError("unit map must be project experience")
    context = json.loads(context_path.read_text(encoding="utf-8"))
    source_records = []
    all_tags: set[str] = set()
    ratings: dict[str, int] = {}
    original_content = {}
    chapters_by_source = []
    clean_footer = None
    footer_reference_count = float("inf")
    for item in request["sources"]:
        period = item["period"]
        tags = set(config["report_period_overrides"][period]["only_valves"])
        if all_tags & tags:
            raise ValueError("source periods overlap")
        path = (project_dir / item["docx"]).resolve()
        if not any(path.is_relative_to(project_dir / "experience" / folder)
                   for folder in ("feedback_inbox", "agent_synced")):
            raise ValueError("source must be archived in feedback_inbox or agent_synced")
        provenance = item.get("provenance", {})
        if not isinstance(provenance, dict):
            raise ValueError("source provenance must be an object")
        if path.is_relative_to(project_dir / "experience/agent_synced") and not provenance:
            raise ValueError("agent-synced source requires provenance")
        document = Document(path)
        section = document.sections[0]
        references = section._sectPr.xpath("./w:footerReference")
        if references and len(references) < footer_reference_count:
            clean_footer = deepcopy(section.footer._element)
            footer_reference_count = len(references)
        chapters = _chapters(document, tags)
        for tag in tags:
            unit = str((context.get(tag) or {}).get("unit") or "").strip()
            if not unit:
                raise ValueError(f"missing unit for {tag}")
            ratings[tag] = chapters[tag][2]
            original_content[tag] = _chapter_content(document, *chapters[tag][:2])
        chapters_by_source.append(chapters)
        source_records.append({"period": period, "docx": str(path.relative_to(project_dir)), "sha256": _hash(path), "count": len(tags), "provenance": provenance})
        all_tags.update(tags)
    if not source_records or not all_tags.issubset(context):
        raise ValueError("unit map and selected valve scope differ")
    grouped = defaultdict(list)
    for tag in all_tags:
        grouped[str(context[tag]["unit"]).strip()].append(tag)

    output_root = project_dir / "output/business"
    staging_root = project_dir / "work/unit-report-assembly"
    used_versions = []
    for root, pattern in ((output_root, r"按装置分册-v(\d+)"),
                          (staging_root, r"pending-v(\d+)")):
        for child in root.iterdir() if root.exists() else ():
            if found := re.fullmatch(pattern, child.name):
                used_versions.append(int(found.group(1)))
    version = max(used_versions, default=0) + 1
    output_dir = output_root / f"按装置分册-v{version}"
    pending_dir = staging_root / f"pending-v{version}"
    pending_dir.mkdir(parents=True)
    completed = []
    names = set()
    for unit, tags in sorted(grouped.items()):
        ordered = [tag for source in source_records for tag in config["report_period_overrides"][source["period"]]["only_valves"] if tag in tags]
        volumes = []
        for source, chapters in zip(source_records, chapters_by_source):
            selected = set(tags) & set(chapters)
            if selected:
                # Reload the immutable archived source for every volume.
                fresh = Document(project_dir / source["docx"])
                volumes.append(_subset(fresh, chapters, selected))
        master = volumes[0]
        _add_front(master, title, unit, ordered, ratings)
        composer = Composer(master)
        for volume in volumes[1:]:
            composer.append(volume)
        safe_unit = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', "_", unit).strip(" .")
        name = f"{safe_unit}-阀门离线诊断报告.docx"
        if name in names:
            raise ValueError(f"unit report filename collision: {name}")
        names.add(name)
        staged = pending_dir / name
        composer.save(staged)
        rendered = Document(staged)
        footer_references = rendered.sections[0]._sectPr.xpath("./w:footerReference")
        if len(footer_references) > 1 and clean_footer is not None:
            section = rendered.sections[0]
            for reference in footer_references:
                if reference.get(qn("w:type")) != "default":
                    section._sectPr.remove(reference)
            footer = section.footer._element
            for child in list(footer):
                footer.remove(child)
            for child in clean_footer:
                footer.append(deepcopy(child))
            rendered.save(staged)
            rendered = Document(staged)
        verified = _chapters(rendered, set(tags))
        if {tag: value[2] for tag, value in verified.items()} != {tag: ratings[tag] for tag in tags}:
            raise RuntimeError(f"rating coverage changed in {staged}")
        for tag, (start, stop, _) in verified.items():
            if _chapter_content(rendered, start, stop) != original_content[tag]:
                raise RuntimeError(f"source chapter text or images changed: {tag} in {staged}")
        completed.append({"unit": unit, "docx": str((output_dir / name).relative_to(project_dir)), "sha256": _hash(staged),
                          "valves": ordered, "ratings": dict(Counter(ratings[tag] for tag in tags))})
    manifest = {"sources": source_records, "unit_context": str(context_path.relative_to(project_dir)),
                "unit_context_sha256": _hash(context_path), "valve_count": len(all_tags),
                "unit_count": len(completed), "reports": completed}
    (pending_dir / "分册清单.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    pending_dir.rename(output_dir)
    return output_dir / "分册清单.json"


def validate_unit_report_pdf(project_dir: Path, docx: Path, pdf: Path) -> dict | None:
    """Check a WPS unit-volume PDF against its bound DOCX and assembly manifest."""
    manifest_path = docx.parent / "分册清单.json"
    if not manifest_path.is_file():
        return None
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    record = next(
        (item for item in manifest.get("reports", [])
         if (project_dir / item["docx"]).resolve() == docx.resolve()),
        None,
    )
    if record is None:
        return None

    failures = []
    if _hash(docx) != record["sha256"]:
        failures.append("unit DOCX differs from its assembly manifest")
    tags = set(record["valves"])
    document = Document(docx)
    try:
        chapters = _chapters(document, tags)
    except ValueError as exc:
        failures.append(str(exc))
        chapters = {}
    if Counter(chapter[2] for chapter in chapters.values()) != Counter(
        {int(level): count for level, count in record["ratings"].items()}
    ):
        failures.append("unit DOCX rating counts differ from its assembly manifest")

    with fitz.open(pdf) as rendered:
        pdf_text = "\n".join(page.get_text() for page in rendered)
    compact_pdf = normalize_report_text(pdf_text)
    for tag, (_, _, rating) in chapters.items():
        if normalize_report_text(f"{tag}{rating}级") not in compact_pdf:
            failures.append(f"PDF missing valve/rating: {tag}")
    checked_paragraphs, missing = missing_pdf_body_text(document, pdf_text)
    if missing:
        failures.append(f"PDF missing {len(missing)} DOCX body paragraphs")
    return {
        "ok": not failures,
        "unit": record["unit"],
        "valve_count": len(tags),
        "checked_paragraphs": checked_paragraphs,
        "failures": failures,
    }
