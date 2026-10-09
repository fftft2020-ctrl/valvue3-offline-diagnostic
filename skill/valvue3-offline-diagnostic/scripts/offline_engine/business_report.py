"""The shared DOCX renderer for review, data-only, and business reports."""

from __future__ import annotations

import math
import re
from collections import Counter
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Mapping
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

from .curve_identity import curve_review_marker, sort_test_images, spec_insertion_marker
from .presentation import (
    CUSTOMER_METRIC_LABELS,
    TEST_EXECUTION_STATUS_ZH,
    TEST_LABELS,
    basic_information_rows,
    clean_fault,
    curve_execution_time_label,
    display_actuator_style,
    display_air_action,
    display_manufacturer,
    display_model_spec,
    display_positioner_model,
    display_specification_value,
    format_friction_metric,
    format_number,
    normalize_customer_terms,
    replace_customer_terms,
    sanitize_customer_text,
    test_execution_issue_text,
)


FONT_LATIN = "Arial"
FONT_CJK = "Microsoft YaHei"
PAGE_TITLE_SIZE = 18.0
SECTION_TITLE_SIZE = 11.0
BODY_SIZE = 8.8
COMPACT_BODY_SIZE = 8.5
DASHBOARD_BODY_SIZE = 8.0
DASHBOARD_COMPACT_SIZE = 7.8
DASHBOARD_SECTION_SIZE = 10.3
DASHBOARD_TITLE_SIZE = 17.5
WHITE = "FFFFFF"
BRAND_PRIMARY = "004B46"
BRAND_SECONDARY = "006F66"
TEXT = "213B37"
MUTED = "496A64"
MINT = "EAF4F1"
PALE = "F6F9F8"
LINE = "B9D4CE"
RATING_COLORS = {
    1: "2E8B57",
    2: "C79200",
    3: "D96C0B",
    4: "C9342C",
    5: "6B7280",
}
RATING_SHORT = {
    1: "无影响，维持运行并纳入周期点检",
    2: "轻微缺陷，建议安排计划复核",
    3: "有缺陷，建议计划复核或维护",
    4: "严重缺陷，建议优先复核和检修",
    5: "无数据 / 未知影响，需补充测试",
}

OWNER_PENDING = "待业主补充"

REPORT_FIELD_MANIFEST_PATHS = (
    "positioner.option_config", "positioner.device_mode", "positioner.valve_position",
    "positioner.manual_setpoint", "positioner.input_signal", "positioner.board_temperature",
    "positioner.min_temperature", "positioner.max_temperature", "runtime.strokes",
    "runtime.cycles", "runtime.time_open", "runtime.time_closed", "runtime.time_near_closed",
    "file.report_start_time", "file.report_end_time", "file.report_time",
    "status.active_faults", "status.historical_faults", "status.task_execution",
    "metrics.run_ramp.hysteresis.max", "metrics.run_ramp.hysteresis.average",
    "metrics.run_ramp.hysteresis.min", "metrics.run_ramp.dead_band.max",
    "metrics.run_ramp.dead_band.average", "metrics.run_ramp.dead_band.min",
    "metrics.run_ramp.linearity.max", "metrics.run_ramp.linearity.average",
    "metrics.run_ramp.linearity.min",
)


def _mapping_value(mapping: Any, *path: str) -> Any:
    """Read a normalized evidence value without exposing its wrapper."""
    current = mapping
    for key in path:
        if not isinstance(current, dict) or key not in current:
            return None
        current = current[key]
    if isinstance(current, dict):
        for key in ("normalized_value", "value", "raw_value"):
            if current.get(key) not in (None, ""):
                return current[key]
    return current


def _first_value(*values: Any, default: str = "未提供") -> Any:
    return next((value for value in values if value not in (None, "", [], {})), default)


def _owner_context(value: Any) -> str:
    text = str(value or "").strip()
    return text if text and text not in {"未提供", "数据不足"} else OWNER_PENDING


def _with_unit(value: Any, unit: str) -> str:
    if value in (None, "", [], {}):
        return "未提供"
    text = str(value).strip()
    if not text:
        return "未提供"
    if text.startswith("-0.0"):
        text = text[1:]
    return text if re.search(r"[A-Za-z%]$", text) else f"{text} {unit}"


def _normalize_actuation_text(value: Any) -> str:
    text = str(value or "").strip()
    replacements = (
        (r"\bDouble(?:\s+Acting)?\b", "双作用"),
        (r"\bSingle(?:\s+Acting)?\b", "单作用"),
        (r"\bATO\b", "气开"),
        (r"\bATC\b", "气关"),
    )
    for pattern, replacement in replacements:
        text = re.sub(pattern, replacement, text, flags=re.I)
    return text or "未提供"


def _optimized_png(data: bytes) -> bytes:
    with Image.open(BytesIO(data)) as image:
        if image.mode == "RGB":
            method = Image.Quantize.MEDIANCUT
        elif image.mode == "RGBA":
            method = Image.Quantize.FASTOCTREE
        else:
            return data
        optimized = image.quantize(
            colors=256,
            method=method,
            dither=Image.Dither.NONE,
        )
        buffer = BytesIO()
        optimized.save(buffer, format="PNG", optimize=True, compress_level=9)
    candidate = buffer.getvalue()
    return candidate if len(candidate) < len(data) else data


def add_customer_picture(
    container: Any,
    image_path: str | Path,
    *,
    width: Any,
    picture_name: str,
):
    """Insert a report image without retaining its source filename."""
    safe_name = re.sub(r"[^A-Za-z0-9_.-]+", "-", picture_name).strip("-.")
    if not safe_name:
        safe_name = "report-image.png"
    inline_shape = container.add_picture(str(image_path), width=width)
    properties = inline_shape._inline.xpath(".//pic:cNvPr")
    if len(properties) != 1:
        raise ValueError("inserted picture is missing non-visual properties")
    properties[0].set("name", safe_name)
    properties[0].attrib.pop("descr", None)
    properties[0].attrib.pop("title", None)
    return inline_shape


def add_bookmark(paragraph: Any, name: str) -> None:
    part = paragraph.part
    bookmark_id = getattr(part, "_offline_bookmark_counter", 0) + 1
    part._offline_bookmark_counter = bookmark_id
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_hyperlink(paragraph: Any, text: str, anchor: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Microsoft YaHei")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r_pr.extend((r_fonts, color, underline))
    text_el = OxmlElement("w:t")
    text_el.text = sanitize_customer_text(text)
    run.extend((r_pr, text_el))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _compact_test_execution_issue_text(item: dict[str, Any]) -> str:
    text = test_execution_issue_text(item)
    text = re.sub(r"（原始报告状态 ([^)]+)）", r" [\1]", text)
    return text.replace("（第 ", "（第").replace(" 次）", "次）")


def _optimize_docx_png_media(path: Path) -> None:
    temporary = path.with_name(path.name + ".media-optimized.tmp")
    try:
        with ZipFile(path, "r") as source, ZipFile(temporary, "w") as destination:
            for item in source.infolist():
                data = source.read(item.filename)
                if (
                    item.filename.startswith("word/media/")
                    and item.filename.lower().endswith(".png")
                ):
                    data = _optimized_png(data)
                destination.writestr(item, data)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def _font(run, size: float = 8, *, bold: bool = False, color: str = TEXT) -> None:
    normalized_text = normalize_customer_terms(run.text)
    if normalized_text != run.text:
        run.text = normalized_text
    run.font.name = FONT_LATIN
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), FONT_CJK)


def _set_column_widths(table, widths: list[float]) -> None:
    table.autofit = False
    for index, width in enumerate(widths):
        value = Inches(width)
        table.columns[index].width = value
        for cell in table.columns[index].cells:
            cell.width = value


def _keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _set_row_min_height(row, height: float) -> None:
    row.height = Inches(height)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def _shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def _borders(cell, color: str = LINE, size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:color"), color)


def _no_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "nil")


def _margins(cell, top: int = 50, start: int = 60, bottom: int = 50, end: int = 60) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _cell_text(
    cell,
    text: Any,
    size: float = 7.2,
    *,
    bold: bool = False,
    color: str = TEXT,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    _font(paragraph.add_run(sanitize_customer_text(str(text))), size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_internal_link(
    paragraph,
    text: str,
    anchor: str,
    *,
    size: float = 7.7,
    bold: bool = True,
) -> None:
    add_internal_hyperlink(paragraph, text, anchor)
    hyperlink = paragraph._p[-1]
    run = hyperlink.find(qn("w:r"))
    r_pr = run.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    for attr in ("ascii", "hAnsi", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), FONT_CJK)
    for name in ("sz", "szCs"):
        node = r_pr.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            r_pr.append(node)
        node.set(qn("w:val"), str(round(size * 2)))
    if bold and r_pr.find(qn("w:b")) is None:
        r_pr.append(OxmlElement("w:b"))


def _cell_link(
    cell,
    text: str,
    anchor: str,
    *,
    size: float = 7.7,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.CENTER,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    _add_internal_link(paragraph, text, anchor, size=size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _page_break_before_next(doc: Document) -> None:
    """Start the next block on a new page without creating a stranded break page."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    _font(paragraph.add_run("\u200b"), 1.0, color=WHITE)


def _section_label(doc: Document, text: str, *, size: float = 10.5) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    _font(paragraph.add_run(text), size, bold=True, color=BRAND_PRIMARY)
    p_pr = paragraph._p.get_or_add_pPr()
    p_borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), BRAND_SECONDARY)
    p_borders.append(bottom)
    p_pr.append(p_borders)


def _cell_bullets(
    cell,
    items: list[str],
    *,
    size: float,
    limit: int | None = 3,
) -> None:
    """Render customer-facing bullets without changing the source wording."""
    cell.text = ""
    visible_items = items if limit is None else items[:limit]
    for index, item in enumerate(visible_items):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        _font(paragraph.add_run(f"• {sanitize_customer_text(item)}"), size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _apply_term_replacements(doc: Document, replacements: dict[str, str] | None) -> None:
    """Apply project-configured terminology at the final customer-render boundary."""
    for text_node in doc.element.body.iter(qn("w:t")):
        text_node.text = replace_customer_terms(text_node.text or "", replacements)


def _without_unavailable_retest_recommendations(text: str) -> str:
    """Remove retest-only exits while preserving an actionable leading clause."""
    banned = re.compile(r"复测|补采|重新测试")
    rewritten: list[str] = []
    for sentence in re.split(r"(?<=[。！？])", text):
        match = banned.search(sentence)
        if match is None:
            rewritten.append(sentence)
            continue
        prefix = sentence[:match.start()]
        boundary = max(prefix.rfind("，"), prefix.rfind("；"))
        if boundary >= 0:
            kept = prefix[:boundary].rstrip("，； ")
            if kept.strip("•  "):
                rewritten.append(kept + "。")
    return "".join(rewritten)


def _remove_unavailable_retest_recommendations(doc: Document) -> None:
    for text_node in doc.element.body.iter(qn("w:t")):
        text_node.text = _without_unavailable_retest_recommendations(
            text_node.text or ""
        )


def _remove_field_context_overviews(doc: Document) -> None:
    """Remove only the project-wide field-context section, not per-valve context."""
    body = doc.element.body
    children = list(body)
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        text = "".join(node.text or "" for node in child.iter(qn("w:t")))
        if text.strip() != "现场信息清单":
            continue
        body.remove(child)
        for following in children[index + 1:]:
            if following.getparent() is not body:
                continue
            if following.tag == qn("w:tbl"):
                body.remove(following)
            break


def _remove_heading_and_following_tables(
    doc: Document,
    heading: str,
    *,
    table_count: int,
) -> None:
    body = doc.element.body
    children = list(body)
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        text = "".join(node.text or "" for node in child.iter(qn("w:t"))).strip()
        if text != heading:
            continue
        body.remove(child)
        removed = 0
        for following in children[index + 1:]:
            if following.getparent() is not body:
                continue
            if following.tag == qn("w:tbl"):
                body.remove(following)
                removed += 1
                if removed == table_count:
                    break


def _remove_overview_priority(doc: Document) -> None:
    body = doc.element.body
    for table in list(body.findall(qn("w:tbl"))):
        text = "".join(node.text or "" for node in table.iter(qn("w:t")))
        if "首要关注" in text:
            body.remove(table)


def _remove_recommendation_section(doc: Document, heading: str) -> None:
    body = doc.element.body
    children = list(body)
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        text = "".join(node.text or "" for node in child.iter(qn("w:t"))).strip()
        if text != heading:
            continue
        body.remove(child)
        for following in children[index + 1:]:
            if following.getparent() is not body or following.tag != qn("w:p"):
                break
            following_text = "".join(
                node.text or "" for node in following.iter(qn("w:t"))
            ).strip()
            if not following_text.startswith("•"):
                break
            body.remove(following)


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def _configure_word_compatibility(doc: Document) -> None:
    settings = doc.settings._element
    compatibility = settings.find(qn("w:compat"))
    if compatibility is None:
        compatibility = OxmlElement("w:compat")
        settings.append(compatibility)
    compatibility_mode = next(
        (
            node
            for node in compatibility.findall(qn("w:compatSetting"))
            if node.get(qn("w:name")) == "compatibilityMode"
        ),
        None,
    )
    if compatibility_mode is None:
        compatibility_mode = OxmlElement("w:compatSetting")
        compatibility_mode.set(qn("w:name"), "compatibilityMode")
        compatibility_mode.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
        compatibility.append(compatibility_mode)
    compatibility_mode.set(qn("w:val"), "15")

    theme_language = settings.find(qn("w:themeFontLang"))
    if theme_language is None:
        theme_language = OxmlElement("w:themeFontLang")
        settings.append(theme_language)
    theme_language.set(qn("w:val"), "en-US")
    theme_language.set(qn("w:eastAsia"), "zh-CN")

    styles = doc.styles.element
    default_run_properties = styles.xpath(
        "./w:docDefaults/w:rPrDefault/w:rPr"
    )[0]
    default_fonts = default_run_properties.find(qn("w:rFonts"))
    if default_fonts is None:
        default_fonts = OxmlElement("w:rFonts")
        default_run_properties.insert(0, default_fonts)
    for attribute in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
        default_fonts.attrib.pop(qn(f"w:{attribute}"), None)
    for attribute in ("ascii", "hAnsi", "cs"):
        default_fonts.set(qn(f"w:{attribute}"), FONT_LATIN)
    default_fonts.set(qn("w:eastAsia"), FONT_CJK)

    default_language = default_run_properties.find(qn("w:lang"))
    if default_language is None:
        default_language = OxmlElement("w:lang")
        default_run_properties.append(default_language)
    default_language.set(qn("w:val"), "en-US")
    default_language.set(qn("w:eastAsia"), "zh-CN")


def _setup_document(
    doc: Document,
    *,
    project_name: str,
    data_only: bool = False,
) -> None:
    _configure_word_compatibility(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.50)
    section.right_margin = Inches(0.50)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.23)
    section.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(8.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    normal_fonts = normal._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "cs"):
        normal_fonts.set(qn(f"w:{attribute}"), FONT_LATIN)
    normal_fonts.set(qn("w:eastAsia"), FONT_CJK)

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(7.25))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(4.9)
    table.columns[1].width = Inches(2.35)
    _cell_text(table.cell(0, 0), "BAKER HUGHES  |  MASONEILAN · VALVUE3", 8.1, bold=True, color=BRAND_PRIMARY)
    _cell_text(table.cell(0, 1), project_name, 8.1, color=MUTED, align=WD_ALIGN_PARAGRAPH.RIGHT)
    for cell in table.rows[0].cells:
        _borders(cell, BRAND_PRIMARY, "0")
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = tc_pr.first_child_found_in("w:tcBorders")
        bottom = borders.find(qn("w:bottom"))
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "10")
        bottom.set(qn("w:color"), BRAND_PRIMARY)
        _margins(cell, 0, 20, 45, 20)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_label = "阀门离线检测" if data_only else "阀门离线诊断"
    _font(
        footer.add_run(f"Baker Hughes Confidential   |   {footer_label}   |   "),
        7.2,
        color=MUTED,
    )
    _add_page_number(footer)


def _add_cover(
    doc: Document,
    *,
    project_name: str,
    report_name: str,
    report_date: str,
    report_period_label: str = "",
    review_label: str = "",
    data_only: bool = False,
) -> None:
    spacer = doc.add_paragraph()
    spacer.add_run("\n\n\n")
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(7.2)
    cell = table.cell(0, 0)
    _shade(cell, BRAND_PRIMARY)
    _borders(cell, BRAND_PRIMARY, "0")
    _margins(cell, 300, 340, 300, 340)
    paragraph = cell.paragraphs[0]
    _font(paragraph.add_run("BAKER HUGHES | MASONEILAN"), 10.5, bold=True, color="A9D8CE")
    paragraph.add_run("\n")
    _font(paragraph.add_run(report_name), 23, bold=True, color=WHITE)
    paragraph.add_run("\n")
    subtitle = "阀门离线检测数据报告" if data_only else "阀门离线诊断报告"
    _font(paragraph.add_run(subtitle), 13.5, color="D2ECE6")
    if review_label.strip():
        paragraph.add_run("\n")
        _font(paragraph.add_run(review_label.strip()), 12.0, bold=True, color="A9D8CE")

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(24)
    _font(paragraph.add_run(project_name), 16, bold=True, color=BRAND_PRIMARY)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(24)
    _font(
        paragraph.add_run(
            ("检测周期" if data_only else "诊断周期")
            if report_period_label.strip()
            else "报告日期"
        ),
        8,
        bold=True,
        color=MUTED,
    )
    paragraph.add_run("\n")
    _font(paragraph.add_run(report_period_label.strip() or report_date), 11, color=TEXT)


def _rating_cards(
    doc: Document,
    valves: list[Any],
    *,
    blank_ratings: bool = False,
) -> None:
    counts = Counter(int(valve.rating) for valve in valves)
    cards = doc.add_table(rows=1, cols=6)
    cards.alignment = WD_TABLE_ALIGNMENT.CENTER
    cards.autofit = False
    data = [(str(len(valves)), "诊断阀门", BRAND_PRIMARY)] + [
        (
            "" if blank_ratings else str(counts.get(level, 0)),
            f"{level}级",
            RATING_COLORS[level],
        )
        for level in range(1, 6)
    ]
    for index, (number, label, color) in enumerate(data):
        cell = cards.cell(0, index)
        _shade(cell, color)
        _borders(cell, WHITE, "0")
        _margins(cell, 150, 50, 145, 50)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(paragraph.add_run(number), 17.5, bold=True, color=WHITE)
        paragraph.add_run("\n")
        _font(paragraph.add_run(label), 8.2, color=WHITE)


def _add_overview(
    doc: Document,
    valves: list[Any],
    *,
    unit_name: str = "",
    blank_ratings: bool = False,
    blank_actions: bool = False,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    _font(paragraph.add_run("诊断总览"), PAGE_TITLE_SIZE, bold=True, color=BRAND_PRIMARY)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    _font(paragraph.add_run(f"全量评级概览 | {len(valves)} 台阀门"), 9.2, color=MUTED)
    _rating_cards(doc, valves, blank_ratings=blank_ratings)

    _section_label(doc, f"{len(valves)} 台阀门评级清单", size=11.3)
    group_size = max(1, math.ceil(len(valves) / 3))
    row_count = min(group_size, len(valves))
    matrix = doc.add_table(rows=row_count + 1, cols=9)
    matrix.alignment = WD_TABLE_ALIGNMENT.CENTER
    matrix.autofit = False
    for index, width in enumerate([0.35, 1.45, 0.55] * 3):
        matrix.columns[index].width = Inches(width)
    for index, text in enumerate(["序号", "位号", "评级"] * 3):
        cell = matrix.cell(0, index)
        _shade(cell, BRAND_PRIMARY)
        _borders(cell)
        _margins(cell, 100, 40, 100, 40)
        _cell_text(cell, text, COMPACT_BODY_SIZE, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    _repeat_table_header(matrix.rows[0])
    _set_row_min_height(matrix.rows[0], 0.36)
    for row_index in range(row_count):
        _set_row_min_height(matrix.rows[row_index + 1], 0.45)
        for group in range(3):
            item_index = group * group_size + row_index
            base = group * 3
            if item_index >= len(valves):
                for offset in range(3):
                    cell = matrix.cell(row_index + 1, base + offset)
                    _shade(cell, WHITE)
                    _borders(cell)
                continue
            valve = valves[item_index]
            values = [
                str(item_index + 1),
                valve.folder,
                "" if blank_ratings else f"{valve.rating}级",
            ]
            for offset, value in enumerate(values):
                cell = matrix.cell(row_index + 1, base + offset)
                _borders(cell)
                _margins(cell, 130, 40, 130, 40)
                if offset == 2:
                    _shade(
                        cell,
                        PALE if blank_ratings else RATING_COLORS[int(valve.rating)],
                    )
                    _cell_text(
                        cell,
                        value,
                        BODY_SIZE,
                        bold=True,
                        color=TEXT if blank_ratings else WHITE,
                        align=WD_ALIGN_PARAGRAPH.CENTER,
                    )
                else:
                    _shade(cell, PALE if row_index % 2 == 0 else WHITE)
                    if offset == 1:
                        _cell_link(cell, value, f"valve_{valve.folder}", size=BODY_SIZE)
                    else:
                        _cell_text(
                            cell,
                            value,
                            BODY_SIZE,
                            align=WD_ALIGN_PARAGRAPH.CENTER,
                        )

    _section_label(doc, "现场信息清单", size=11.3)
    field_context = doc.add_table(rows=len(valves) + 1, cols=6)
    field_context.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(field_context, [1.15, 1.7, 0.9, 1.3, 0.7, 1.45])
    headers = ["位号", "描述", "装置", "失效后果", "有无副线", "现场备注"]
    for index, text in enumerate(headers):
        cell = field_context.cell(0, index)
        _shade(cell, BRAND_PRIMARY)
        _borders(cell)
        _margins(cell, 90, 40, 90, 40)
        _cell_text(
            cell,
            text,
            COMPACT_BODY_SIZE,
            bold=True,
            color=WHITE,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    _repeat_table_header(field_context.rows[0])
    _set_row_min_height(field_context.rows[0], 0.34)
    for row_index, valve in enumerate(valves, start=1):
        values = [
            valve.folder,
            _owner_context(getattr(valve, "description", "")),
            _owner_context(getattr(valve, "unit", "")),
            _owner_context(getattr(valve, "failure_consequence", "")),
            _owner_context(getattr(valve, "bypass", "")),
            _owner_context(getattr(valve, "site_note", "")),
        ]
        for column_index, value in enumerate(values):
            cell = field_context.cell(row_index, column_index)
            _shade(cell, PALE if row_index % 2 else WHITE)
            _borders(cell)
            _margins(cell, 80, 35, 80, 35)
            if column_index == 0:
                _cell_link(
                    cell,
                    value,
                    f"valve_{valve.folder}",
                    size=COMPACT_BODY_SIZE,
                )
            else:
                _cell_text(cell, value, COMPACT_BODY_SIZE)
        _keep_row_together(field_context.rows[row_index])
        _set_row_min_height(field_context.rows[row_index], 0.32)

    _section_label(doc, "整体判断", size=11.3)
    levels = [level for level in range(1, 5) if any(int(valve.rating) == level for valve in valves)]
    summary = doc.add_table(rows=1, cols=max(1, len(levels)))
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, level in enumerate(levels or [5]):
        cell = summary.cell(0, index)
        _shade(cell, MINT)
        _borders(cell)
        _margins(cell, 120, 110, 120, 110)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        if not blank_ratings:
            count = sum(1 for valve in valves if int(valve.rating) == level)
            _font(paragraph.add_run(f"{level}级 | {count}台"), 10.0, bold=True, color=RATING_COLORS[level])
            paragraph.add_run("\n")
            _font(paragraph.add_run(RATING_SHORT[level]), BODY_SIZE, color=TEXT)
    _set_row_min_height(summary.rows[0], 0.58)

    priority = None if blank_ratings else max(
        valves,
        key=lambda valve: (int(valve.rating), bool(getattr(valve, "basic_action_failed", False))),
    )
    alert = doc.add_table(rows=1, cols=2)
    alert.alignment = WD_TABLE_ALIGNMENT.CENTER
    alert.autofit = False
    alert.columns[0].width = Inches(1.12)
    alert.columns[1].width = Inches(5.96)
    _shade(alert.cell(0, 0), BRAND_SECONDARY)
    _shade(alert.cell(0, 1), PALE)
    _cell_text(alert.cell(0, 0), "首要关注", BODY_SIZE, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    alert_cell = alert.cell(0, 1)
    alert_cell.text = ""
    alert_paragraph = alert_cell.paragraphs[0]
    alert_paragraph.paragraph_format.space_before = Pt(0)
    alert_paragraph.paragraph_format.space_after = Pt(0)
    alert_paragraph.paragraph_format.line_spacing = 1.0
    if priority is not None:
        _add_internal_link(alert_paragraph, priority.folder, f"valve_{priority.folder}", size=BODY_SIZE)
    if priority is not None and not blank_actions:
        _font(alert_paragraph.add_run(f"：{priority.action}"), BODY_SIZE)
    alert_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_row_min_height(alert.rows[0], 0.44)
    for cell in alert.rows[0].cells:
        _borders(cell)
        _margins(cell, 90, 100, 90, 100)


def _add_data_overview(
    doc: Document,
    valves: list[Any],
    *,
    unit_name: str = "",
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    _font(
        paragraph.add_run("设备与检测信息总览"),
        PAGE_TITLE_SIZE,
        bold=True,
        color=BRAND_PRIMARY,
    )
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    _font(
        paragraph.add_run(f"全量设备与检测信息 | {len(valves)} 台阀门"),
        9.2,
        color=MUTED,
    )

    count = doc.add_table(rows=1, cols=1)
    count.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = count.cell(0, 0)
    _shade(cell, BRAND_PRIMARY)
    _borders(cell, WHITE, "0")
    _margins(cell, 150, 50, 145, 50)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(paragraph.add_run(str(len(valves))), 17.5, bold=True, color=WHITE)
    paragraph.add_run("\n")
    _font(paragraph.add_run("检测阀门"), 8.2, color=WHITE)

    _section_label(doc, f"{len(valves)} 台阀门清单", size=11.3)
    group_size = max(1, math.ceil(len(valves) / 3))
    row_count = min(group_size, len(valves))
    matrix = doc.add_table(rows=row_count + 1, cols=6)
    matrix.alignment = WD_TABLE_ALIGNMENT.CENTER
    matrix.autofit = False
    for index, width in enumerate([0.45, 1.95] * 3):
        matrix.columns[index].width = Inches(width)
    for index, text in enumerate(["序号", "位号"] * 3):
        cell = matrix.cell(0, index)
        _shade(cell, BRAND_PRIMARY)
        _borders(cell)
        _margins(cell, 100, 40, 100, 40)
        _cell_text(
            cell,
            text,
            COMPACT_BODY_SIZE,
            bold=True,
            color=WHITE,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    _repeat_table_header(matrix.rows[0])
    _set_row_min_height(matrix.rows[0], 0.36)
    for row_index in range(row_count):
        _set_row_min_height(matrix.rows[row_index + 1], 0.45)
        for group in range(3):
            item_index = group * group_size + row_index
            base = group * 2
            for offset in range(2):
                cell = matrix.cell(row_index + 1, base + offset)
                _borders(cell)
                _margins(cell, 130, 40, 130, 40)
                _shade(cell, PALE if row_index % 2 == 0 else WHITE)
            if item_index >= len(valves):
                continue
            valve = valves[item_index]
            _cell_text(
                matrix.cell(row_index + 1, base),
                str(item_index + 1),
                BODY_SIZE,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
            _cell_link(
                matrix.cell(row_index + 1, base + 1),
                valve.folder,
                f"valve_{valve.folder}",
                size=BODY_SIZE,
            )

    _section_label(doc, "现场信息清单", size=11.3)
    field_context = doc.add_table(rows=len(valves) + 1, cols=6)
    field_context.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(field_context, [1.15, 1.7, 0.9, 1.3, 0.7, 1.45])
    headers = ["位号", "描述", "装置", "失效后果", "有无副线", "现场备注"]
    for index, text in enumerate(headers):
        cell = field_context.cell(0, index)
        _shade(cell, BRAND_PRIMARY)
        _borders(cell)
        _margins(cell, 90, 40, 90, 40)
        _cell_text(
            cell,
            text,
            COMPACT_BODY_SIZE,
            bold=True,
            color=WHITE,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    _repeat_table_header(field_context.rows[0])
    _set_row_min_height(field_context.rows[0], 0.34)
    for row_index, valve in enumerate(valves, start=1):
        values = [
            valve.folder,
            _owner_context(getattr(valve, "description", "")),
            _owner_context(getattr(valve, "unit", "")),
            _owner_context(getattr(valve, "failure_consequence", "")),
            _owner_context(getattr(valve, "bypass", "")),
            _owner_context(getattr(valve, "site_note", "")),
        ]
        for column_index, value in enumerate(values):
            cell = field_context.cell(row_index, column_index)
            _shade(cell, PALE if row_index % 2 else WHITE)
            _borders(cell)
            _margins(cell, 80, 35, 80, 35)
            if column_index == 0:
                _cell_link(
                    cell,
                    value,
                    f"valve_{valve.folder}",
                    size=COMPACT_BODY_SIZE,
                )
            else:
                _cell_text(cell, value, COMPACT_BODY_SIZE)
        _keep_row_together(field_context.rows[row_index])
        _set_row_min_height(field_context.rows[row_index], 0.32)


def _add_overall_conclusion(
    doc: Document,
    valves: list[Any],
    *,
    blank_ratings: bool = False,
    blank_conclusions: bool = False,
    blank_actions: bool = False,
    show_blank_recommendations: bool = False,
    overall_conclusion: str,
    recommendations: list[str],
    recommendations_heading: str = "后续复测、维护和检修建议",
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    _font(paragraph.add_run("总体结论与检修建议"), PAGE_TITLE_SIZE, bold=True, color=BRAND_PRIMARY)
    conclusion_text = "" if blank_conclusions else overall_conclusion
    conclusion = doc.add_table(rows=1, cols=1)
    conclusion.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = conclusion.cell(0, 0)
    _shade(cell, MINT)
    _borders(cell, "9FC8BE")
    _margins(cell, 160, 160, 160, 160)
    _cell_text(cell, conclusion_text, 10.2, bold=True, color=TEXT)
    _set_row_min_height(conclusion.rows[0], 0.62)

    _section_label(doc, "建议优先处理清单", size=11.3)
    priority = [] if blank_ratings else [
        valve
        for valve in valves
        if int(valve.rating) >= 3
        or bool(getattr(valve, "human_verification_status", ""))
    ]
    table = doc.add_table(rows=1 + max(1, len(priority)), cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, width in enumerate([0.42, 1.0, 0.72, 2.35, 2.68]):
        table.columns[index].width = Inches(width)
    for index, text in enumerate(["评级", "位号", "校验来源", "主要问题", "建议行动"]):
        cell = table.cell(0, index)
        _shade(cell, BRAND_PRIMARY)
        _borders(cell)
        _margins(cell, 70, 60, 70, 60)
        _cell_text(cell, text, BODY_SIZE, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_row_min_height(table.rows[0], 0.36)
    _repeat_table_header(table.rows[0])
    if priority:
        from .presentation import simplify_action_items

        for row_index, valve in enumerate(priority, 1):
            _set_row_min_height(table.rows[row_index], 0.43)
            issue_summary = list(getattr(valve, "summary_issues", []))
            action_summary = list(getattr(valve, "summary_actions", []))
            issue_value: str | list[str] = issue_summary or valve.main_issue
            if blank_actions:
                action_value: str | list[str] = ""
            else:
                action_source = action_summary or valve.action
                normalized_actions = simplify_action_items(action_source)
                # Keep an existing concise scalar action backward-compatible;
                # long or compound actions become explicit bullets.
                if (
                    not action_summary
                    and isinstance(valve.action, str)
                    and len(normalized_actions) == 1
                    and len(normalized_actions[0]) <= 80
                ):
                    action_value = valve.action
                else:
                    action_value = normalized_actions
            source_text = (
                "人工校验"
                if getattr(valve, "human_verification_status", "")
                else "诊断评估"
            )
            for column, value in enumerate((
                f"{valve.rating}级",
                valve.folder,
                source_text,
                issue_value,
                action_value,
            )):
                cell = table.cell(row_index, column)
                _borders(cell)
                _margins(cell, 100, 65, 100, 65)
                if column == 0:
                    _shade(cell, RATING_COLORS[int(valve.rating)])
                    _cell_text(cell, value, COMPACT_BODY_SIZE, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
                else:
                    _shade(cell, PALE if row_index % 2 else WHITE)
                    if column == 1:
                        _cell_link(
                            cell,
                            value,
                            f"valve_{valve.folder}",
                            size=COMPACT_BODY_SIZE,
                            align=WD_ALIGN_PARAGRAPH.LEFT,
                        )
                    elif column in (3, 4) and isinstance(value, list):
                        # The overall table is the customer action summary;
                        # show every normalized bullet so a priority, field
                        # verification, retest, or maintenance-window action
                        # cannot disappear through a display limit.
                        _cell_bullets(
                            cell,
                            value,
                            size=COMPACT_BODY_SIZE,
                            limit=None if column == 4 else 3,
                        )
                    else:
                        _cell_text(cell, value, COMPACT_BODY_SIZE)
    elif blank_ratings:
        for cell in table.rows[1].cells:
            _shade(cell, WHITE)
            _borders(cell)
            _cell_text(cell, "", COMPACT_BODY_SIZE)
    else:
        merged = table.cell(1, 0).merge(table.cell(1, 4))
        _cell_text(merged, "未识别3级或以上阀门。", 8)
    for row in table.rows:
        _keep_row_together(row)

    if not blank_actions or show_blank_recommendations:
        _section_label(doc, recommendations_heading, size=11.3)
    if not blank_actions:
        for index, text in enumerate(recommendations):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.12)
            paragraph.paragraph_format.first_line_indent = Inches(-0.1)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.keep_with_next = (
                index < len(recommendations) - 1
            )
            _font(
                paragraph.add_run("•  "),
                8.9,
                bold=True,
                color=BRAND_SECONDARY,
            )
            _font(paragraph.add_run(text), 8.8)


def _positioner_info_value(info: dict[str, Any], key: str, *sections: str) -> Any:
    """Read structured Positioner Info values with flat-shape compatibility."""
    for section in sections:
        values = info.get(section)
        if isinstance(values, dict) and key in values:
            value = values[key]
            if isinstance(value, dict):
                value = next(
                    (
                        value[name]
                        for name in ("normalized_value", "value", "raw_value")
                        if value.get(name) not in (None, "")
                    ),
                    None,
                )
            if value not in (None, ""):
                return value
    value = info.get(key)
    if isinstance(value, dict):
        return next(
            (
                value[name]
                for name in ("normalized_value", "value", "raw_value")
                if value.get(name) not in (None, "")
            ),
            None,
        )
    return value


def _positioner_fault_values(info: dict[str, Any], key: str) -> list[str]:
    values = _positioner_info_value(info, key, "status") or []
    if isinstance(values, str):
        values = [values]
    return [
        cleaned
        for value in values
        if (cleaned := clean_fault(str(value)))
        and cleaned.lower() not in {"reset", "no error", "none", "未见活动故障"}
    ]


def historical_runtime_note(info: dict[str, Any]) -> str:
    history = info.get("historical_runtime")
    if not isinstance(history, dict):
        return ""

    def counters(values: Any) -> str:
        if not isinstance(values, dict):
            return "未提供"
        return (
            f"行程 {values.get('strokes', '未提供')}；循环 {values.get('cycles', '未提供')}；"
            f"全开/全关/近关 {values.get('time_open', '未提供')} / "
            f"{values.get('time_closed', '未提供')} / {values.get('time_near_closed', '未提供')} h"
        )

    parts = [str(history.get("note") or "").strip()]
    if history.get("historical"):
        parts.append(f"{history.get('historical_period') or '历史'}：{counters(history['historical'])}")
    if history.get("total"):
        parts.append(f"累计：{counters(history['total'])}")
    if history.get("historical_positioner_serial"):
        parts.append(f"历史定位器序列号 {history['historical_positioner_serial']}（当前身份待确认）")
    return "；".join(part for part in parts if part)


def positioner_report_metric_rows(valve: Any) -> list[list[str]]:
    info = valve.positioner_report_info
    if not info:
        return []

    def faults_text(key: str, empty_text: str) -> str:
        values = _positioner_fault_values(info, key)
        if not values and key == "active_faults":
            values = _positioner_fault_values({"status": {key: valve.active_faults}}, key)
        return "；".join(values) or empty_text

    def value(key: str, *sections: str) -> Any:
        return _positioner_info_value(info, key, *sections)

    device_mode = value("device_mode", "positioner")
    valve_position = value("valve_position", "positioner")
    manual_setpoint = value("manual_setpoint", "positioner")
    input_signal = value("input_signal", "positioner")
    supply_pressure = value("supply_pressure", "positioner")
    min_temperature = value("min_temperature", "positioner")
    max_temperature = value("max_temperature", "positioner")
    strokes = value("strokes", "runtime")
    cycles = value("cycles", "runtime")
    time_open = value("time_open", "runtime")
    time_closed = value("time_closed", "runtime")
    time_near_closed = value("time_near_closed", "runtime")

    snapshot = (
        f"模式 {device_mode if device_mode not in (None, '') else '未提供'}；阀位 {valve_position if valve_position not in (None, '') else '未提供'}；"
        f"手动设定 {manual_setpoint if manual_setpoint not in (None, '') else '未提供'}；输入 {input_signal if input_signal not in (None, '') else '未提供'}；"
        f"供气 {supply_pressure if supply_pressure not in (None, '') else '未提供'}；"
        f"温度 {min_temperature if min_temperature not in (None, '') else '未提供'} 至 {max_temperature if max_temperature not in (None, '') else '未提供'}"
    )
    status = info.get("status") if isinstance(info.get("status"), dict) else {}
    fault_data_present = (
        info.get("fault_data_status") == "present"
        or any(key in info for key in ("active_faults", "historical_faults"))
        or any(key in status for key in ("active_faults", "historical_faults"))
        or bool(valve.active_faults)
    )
    history_note = historical_runtime_note(info)
    return [
        ["累计行程 / 循环次数", f"{strokes if strokes not in (None, '') else '未提供'} / {cycles if cycles not in (None, '') else '未提供'}"],
        [
            "全开 / 全关 / 近关时间",
            f"{time_open if time_open not in (None, '') else '未提供'} h / {time_closed if time_closed not in (None, '') else '未提供'} h / {time_near_closed if time_near_closed not in (None, '') else '未提供'} h",
        ],
        ["运行快照", snapshot],
        [
            "当前故障",
            faults_text("active_faults", "未见当前故障")
            if fault_data_present
            else "源文件未提供当前故障字段",
        ],
        [
            "历史故障",
            faults_text("historical_faults", "未见历史故障")
            if fault_data_present
            else "源文件未提供历史故障字段",
        ],
        *([["历史数据备注", history_note]] if history_note else []),
    ]

def valve_model_spec(valve: Any) -> str:
    return display_model_spec(valve.valve_model, valve.spec_fields, valve.valve_spec)


def positioner_information_row(
    positioner_model: str, positioner_type: str
) -> tuple[str, str]:
    if str(positioner_model or "").strip() not in {"", "未提供"}:
        return "定位器型号", display_positioner_model(positioner_model)
    if str(positioner_type or "").strip():
        return "定位器类型", str(positioner_type).strip()
    return "定位器型号", "未提供"


def report_air_action(value: str) -> str:
    return {"ATC": "气关", "ATO": "气开"}.get(
        str(value or "").strip().upper(), display_air_action(value)
    )


def report_test_execution_issues(valve: Any) -> list[dict[str, Any]]:
    issues = list(getattr(valve, "test_execution_issues", ()) or ())
    if issues:
        return issues
    info = getattr(valve, "positioner_report_info", {}) or {}
    tasks = _positioner_info_value(info, "task_execution", "status") or []
    if not isinstance(tasks, list):
        return []
    failed = {"failed", "cancelled", "aborted", "incomplete", "not completed"}
    return [
        {
            "test_key": str(task.get("test_key") or ""),
            "test_name": str(task.get("task_name") or ""),
            "occurrence": int(task.get("occurrence") or 1),
            "status": str(task.get("result") or ""),
        }
        for task in tasks
        if str(task.get("result") or "").lower() in failed
    ]


def report_test_execution_evidence_present(valve: Any) -> bool:
    info = getattr(valve, "positioner_report_info", {}) or {}
    status = info.get("status")
    return "task_execution" in info or (
        isinstance(status, dict) and "task_execution" in status
    )


def positioner_report_basic_information_rows(valve: Any) -> list[tuple[str, str]]:
    info = valve.positioner_report_info
    if not info:
        return []
    firmware = _positioner_info_value(info, "firmware", "positioner", "file")
    software = _positioner_info_value(
        info, "diagnostic_software_version", "positioner", "file"
    )
    level = _positioner_info_value(info, "diagnostics_level", "positioner")
    option = _positioner_info_value(info, "option_config", "positioner")
    style = _positioner_info_value(info, "actuator_style", "positioner")
    if style in (None, ""):
        style = _positioner_info_value(info, "actuator_type", "positioner")
    action = _positioner_info_value(info, "air_action", "positioner")
    return [
        ("固件版本 / 诊断组件版本", f"{firmware or '未提供'} / {software or '未提供'}"),
        (
            "诊断等级 / 配置",
            f"{level or valve.spec_fields.get('positioner_diagnostics_level') or valve.spec_fields.get('positioner_diagnostics_level_note') or '原始报告未提供'} / {option or '未提供'}",
        ),
        ("执行机构形式 / 气路动作", f"{display_actuator_style(style)} / {report_air_action(action or '')}"),
    ]


def specification_supplement_rows(valve: Any) -> list[tuple[str, str]]:
    fields = valve.spec_fields
    leakage_class = display_specification_value(fields.get("leakage_class"))
    if re.fullmatch(r"0\.001%\s*/\s*Max\.?\s*Diff\.?\s*Pr\.?", leakage_class, re.I):
        leakage_class = "V"
    actuator_model = display_specification_value(fields.get("actuator_model"))
    candidates = [
        ("端连接 / 额定 Cv", fields.get("end_connection"), fields.get("rated_cv")),
        ("流向 / 泄漏等级", fields.get("flow_action"), leakage_class),
        ("执行机构型号 / 尺寸", actuator_model, fields.get("actuator_size")),
        ("执行机构供气 / 行程", fields.get("supply_pressure"), fields.get("actuator_travel")),
    ]
    return [
        (label, f"{display_specification_value(left)} / {display_specification_value(right)}")
        for label, left, right in candidates
    ]


def dtm_basic_information_rows(valve: Any) -> list[tuple[str, str]]:
    info = valve.dtm_info
    if info.get("dtm_report_type") != "svi_ii_ap":
        return []
    return [
        ("DTM报告时间", info.get("date_time") or "未提供"),
        ("固件版本 / DTM版本", f"{info.get('firmware') or '未提供'} / {info.get('dtm_revision') or '未提供'}"),
        ("执行机构形式 / 气路动作", f"{display_actuator_style(info.get('actuator_style'))} / {display_air_action(info.get('air_action'))}"),
    ]


def dtm_metric_rows(valve: Any) -> list[list[str]]:
    info = valve.dtm_info
    if info.get("dtm_report_type") != "svi_ii_ap":
        return []
    return [
        ["累计行程 / 循环次数", f"{info.get('strokes') or '未提供'} / {info.get('cycles') or '未提供'}"],
        ["全开 / 全关 / 近关时间", f"{info.get('time_open') or '未提供'} h / {info.get('time_closed') or '未提供'} h / {info.get('time_near_closed') or '未提供'} h"],
        ["DTM运行快照", f"模式 {info.get('device_mode') or '未提供'}；阀位 {info.get('valve_position') or '未提供'}；输入 {info.get('input_signal') or '未提供'}；供气 {info.get('supply_pressure') or '未提供'}；温度 {info.get('min_temperature') or '未提供'} 至 {info.get('max_temperature') or '未提供'}"],
        ["DTM活动故障", info.get("faults") or "未见活动故障"],
    ]


def _device_rows(valve: Any, unit_name: str) -> list[tuple[str, str]]:

    info = getattr(valve, "positioner_report_info", {}) or {}
    start_time = _first_value(
        _mapping_value(info, "file", "report_start_time"),
        info.get("report_start_time"),
        info.get("start_time"),
        default="",
    )
    end_time = _first_value(
        _mapping_value(info, "file", "report_end_time"),
        info.get("report_end_time"),
        info.get("end_time"),
        default="",
    )
    report_time = _first_value(
        _mapping_value(info, "file", "report_time"),
        info.get("report_time"),
        f"{start_time} – {end_time}" if start_time and end_time else "",
        getattr(valve, "detection_date", ""),
    )
    rows = [
        ("位号", valve.folder),
        ("描述", _owner_context(getattr(valve, "description", ""))),
        ("装置", _owner_context(getattr(valve, "unit", ""))),
        ("失效后果", _owner_context(getattr(valve, "failure_consequence", ""))),
        ("有无副线", _owner_context(getattr(valve, "bypass", ""))),
        ("现场备注", _owner_context(getattr(valve, "site_note", ""))),
        ("检测时间", report_time),
        ("厂家", display_manufacturer(valve.manufacturer)),
        ("设备 ID / 序列号", f"{valve.device_id or '未提供'} / {valve.serial_number or '未提供'}"),
        positioner_information_row(valve.positioner_model, valve.positioner_type),
    ]
    basic = basic_information_rows(
        manufacturer=valve.manufacturer,
        positioner_model=valve.positioner_model,
        valve_model=valve.valve_model,
        valve_type=valve.valve_type,
        spec_fields=valve.spec_fields,
        valve_spec=valve.valve_spec,
        spring_range=valve.spring_range,
    )
    rows.extend(row for row in basic if row[0] not in {"厂家", "定位器型号"})
    rows.extend(specification_supplement_rows(valve))
    rows.extend(
        positioner_report_basic_information_rows(valve)
        if valve.positioner_report_info else dtm_basic_information_rows(valve)
    )
    return [
        (label, _normalize_actuation_text(value) if label == "执行机构形式 / 气路动作" else value)
        for label, value in rows
    ]


def _metric_rows(valve: Any) -> tuple[list[tuple[str, str]], dict[str, str]]:
    metrics = valve.metrics
    supply = metrics.get("Supply Pressure") or {}
    key_rows: list[tuple[str, str]] = [
        (CUSTOMER_METRIC_LABELS["Average Error"], format_number(metrics.get("Average Error"), "%")),
        (CUSTOMER_METRIC_LABELS["T86"], format_number(metrics.get("T86"), "s")),
        (CUSTOMER_METRIC_LABELS["T63"], format_number(metrics.get("T63"), "s")),
        (CUSTOMER_METRIC_LABELS["Td"], format_number(metrics.get("Td"), "s")),
        (CUSTOMER_METRIC_LABELS["Overshoot"], format_number(metrics.get("Overshoot"), "%")),
        (
            CUSTOMER_METRIC_LABELS["Supply Pressure"],
            f"{format_number(supply.get('min'), ' psi')} - {format_number(supply.get('max'), ' psi')}；"
            f"平均 {format_number(supply.get('avg'), ' psi')}",
        ),
    ]
    ramp = (
        metrics.get("run_ramp")
        or metrics.get("Run Ramp")
        or _mapping_value(
            getattr(valve, "positioner_report_info", {}) or {},
            "metrics",
            "run_ramp",
        )
        or {}
    )
    ramp_labels = (
        ("hysteresis", "回差"),
        ("dead_band", "死区"),
        ("linearity", "线性度"),
    )
    statistic_labels = (("max", "最大值"), ("average", "平均值"), ("min", "最小值"))
    for metric_key, metric_label in ramp_labels:
        for statistic_key, statistic_label in statistic_labels:
            value = _first_value(
                _mapping_value(ramp, metric_key, statistic_key),
                metrics.get(f"run_ramp.{metric_key}.{statistic_key}"),
                metrics.get(f"Run Ramp {metric_key} {statistic_key}"),
                default=None,
            )
            key_rows.append(
                (f"斜坡测试{metric_label}{statistic_label}", format_number(value, "%"))
            )
    remaining = {
        CUSTOMER_METRIC_LABELS["Max Friction"]: format_friction_metric(metrics, "max"),
        CUSTOMER_METRIC_LABELS["Min Friction"]: format_friction_metric(metrics, "min"),
        CUSTOMER_METRIC_LABELS["Mean Friction"]: format_friction_metric(metrics, "mean"),
    }
    report_rows = _positioner_runtime_rows(valve) if valve.positioner_report_info else (
        [["定位器使用时间 / 运行时间", "未提供"]]
        + dtm_metric_rows(valve)
    )
    remaining.update({key: value for key, value in report_rows})
    return key_rows, remaining


def _performance_test_status(
    valve: Any,
    friction_values: tuple[str, str, str],
) -> str:
    if any(value not in {"", "-", "未提供"} for value in friction_values):
        return "有数据"
    issues = report_test_execution_issues(valve)
    for test_key in ("extended_signature", "standard_signature"):
        issue = next(
            (
                item
                for item in issues
                if str(item.get("test_key") or "") == test_key
            ),
            None,
        )
        if issue is None:
            continue
        status = str(issue.get("status") or issue.get("result") or "").strip()
        status_zh = TEST_EXECUTION_STATUS_ZH.get(
            status.lower(), status or "未完成"
        )
        return f"{status_zh} / 无有效曲线"
    return "无数据"


def _metric_groups(
    valve: Any,
) -> tuple[dict[str, Any], dict[str, str]]:
    key_rows, remaining = _metric_rows(valve)
    values = dict(key_rows)
    step_labels = (
        CUSTOMER_METRIC_LABELS["Average Error"],
        CUSTOMER_METRIC_LABELS["Overshoot"],
        CUSTOMER_METRIC_LABELS["T63"],
        CUSTOMER_METRIC_LABELS["T86"],
        CUSTOMER_METRIC_LABELS["Td"],
    )
    step_rows = [(label, values.get(label, "未提供")) for label in step_labels]
    ramp_rows = [
        (
            metric_label,
            values.get(f"斜坡测试{metric_label}最大值", "未提供"),
            values.get(f"斜坡测试{metric_label}平均值", "未提供"),
            values.get(f"斜坡测试{metric_label}最小值", "未提供"),
        )
        for metric_label in ("回差", "死区", "线性度")
    ]
    friction_labels = (
        CUSTOMER_METRIC_LABELS["Max Friction"],
        CUSTOMER_METRIC_LABELS["Mean Friction"],
        CUSTOMER_METRIC_LABELS["Min Friction"],
    )
    friction_values = tuple(
        remaining.pop(label, "未提供") for label in friction_labels
    )
    performance_rows = [
        ("性能测试", _performance_test_status(valve, friction_values)),
        (
            CUSTOMER_METRIC_LABELS["Supply Pressure"],
            values.get(
                CUSTOMER_METRIC_LABELS["Supply Pressure"],
                "未提供",
            ),
        ),
        *list(zip(friction_labels, friction_values)),
    ]
    return (
        {
            "step": step_rows,
            "ramp_headers": ("斜坡指标", "最大", "平均", "最小"),
            "ramp": ramp_rows,
            "performance": performance_rows,
        },
        remaining,
    )


def _positioner_runtime_rows(valve: Any) -> list[list[str]]:
    info = valve.positioner_report_info or {}

    def value(section: str, key: str, *, unit: str = "") -> str:
        raw = _first_value(
            _mapping_value(info, section, key),
            info.get(key),
            default=None,
        )
        if raw in (None, ""):
            return "未提供"
        return _with_unit(raw, unit) if unit else str(raw)

    def faults(key: str, empty: str) -> str:
        raw = _first_value(_mapping_value(info, "status", key), info.get(key), default=[])
        if isinstance(raw, str):
            raw = [raw]
        values = [
            clean_fault(str(item))
            for item in raw
            if item
            and str(item).strip().lower() not in {"reset", "no error", "none", "未见活动故障"}
        ]
        if key == "active_faults" and not values:
            values = [
                clean_fault(str(item))
                for item in getattr(valve, "active_faults", ()) or ()
                if item
                and str(item).strip().lower() not in {"reset", "no error", "none", "未见活动故障"}
            ]
        return "；".join(value for value in values if value) or empty

    snapshot = (
        f"配置 {value('positioner', 'option_config')}；"
        f"模式 {value('positioner', 'device_mode')}；"
        f"阀位 {value('positioner', 'valve_position', unit='%')}；"
        f"手动给定 {value('positioner', 'manual_setpoint', unit='%')}；"
        f"输入信号 {value('positioner', 'input_signal', unit='mA')}；"
        f"板卡温度 {value('positioner', 'board_temperature', unit='degC')}；"
        f"最低温度 {value('positioner', 'min_temperature', unit='degC')}；"
        f"最高温度 {value('positioner', 'max_temperature', unit='degC')}"
    )
    return [
        ["累计行程次数", value("runtime", "strokes", unit="counts")],
        ["累计循环次数", value("runtime", "cycles", unit="counts")],
        ["累计全开时间", value("runtime", "time_open", unit="h")],
        ["累计全关时间", value("runtime", "time_closed", unit="h")],
        ["累计近关时间", value("runtime", "time_near_closed", unit="h")],
        ["定位器运行快照", snapshot],
        ["活动故障", faults("active_faults", "未见活动故障")],
        ["历史故障", faults("historical_faults", "未见历史故障")],
    ]


def _manifest_contract_evidence(structured: dict[str, Any]) -> dict[str, Any]:
    source_record = structured.get("source_record")
    if not isinstance(source_record, dict):
        return structured
    report_info = source_record.get("positioner_report_info")
    if not isinstance(report_info, dict):
        report_info = {}
    return {
        "positioner": report_info.get("positioner", {}),
        "runtime": report_info.get("runtime", {}),
        "file": report_info.get("file", {}),
        "status": report_info.get("status", {}),
        "metrics": source_record.get("metrics", {}),
    }


def _manifest_scalar(value: Any) -> Any:
    if isinstance(value, dict):
        for key in ("normalized_value", "value", "selected"):
            if key in value and value[key] not in (None, ""):
                return _manifest_scalar(value[key])
        return {str(key): _manifest_scalar(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_manifest_scalar(item) for item in value]
    return value


_REPORT_EVIDENCE_ALIASES: dict[str, tuple[str, ...]] = {
    "positioner.option_config": (
        "positioner_report_info.positioner.option_config",
        "positioner_report_info.option_config",
        "positioner.option_config",
        "option_config",
    ),
    "positioner.device_mode": (
        "positioner_report_info.positioner.device_mode",
        "positioner_report_info.device_mode",
        "positioner.device_mode",
        "device_mode",
    ),
    "positioner.valve_position": (
        "positioner_report_info.positioner.valve_position",
        "positioner_report_info.valve_position",
        "positioner.valve_position",
        "valve_position",
    ),
    "positioner.manual_setpoint": (
        "positioner_report_info.positioner.manual_setpoint",
        "positioner_report_info.manual_setpoint",
        "positioner.manual_setpoint",
        "manual_setpoint",
    ),
    "positioner.input_signal": (
        "positioner_report_info.positioner.input_signal",
        "positioner_report_info.input_signal",
        "positioner.input_signal",
        "input_signal",
    ),
    "positioner.board_temperature": (
        "positioner_report_info.positioner.board_temperature",
        "positioner_report_info.board_temperature",
        "positioner.board_temperature",
        "board_temperature",
    ),
    "positioner.min_temperature": (
        "positioner_report_info.positioner.min_temperature",
        "positioner_report_info.min_temperature",
        "positioner.min_temperature",
        "min_temperature",
    ),
    "positioner.max_temperature": (
        "positioner_report_info.positioner.max_temperature",
        "positioner_report_info.max_temperature",
        "positioner.max_temperature",
        "max_temperature",
    ),
    "runtime.strokes": (
        "positioner_report_info.runtime.strokes",
        "positioner_report_info.strokes",
        "runtime.strokes",
        "strokes",
    ),
    "runtime.cycles": (
        "positioner_report_info.runtime.cycles",
        "positioner_report_info.cycles",
        "runtime.cycles",
        "cycles",
    ),
    "runtime.time_open": (
        "positioner_report_info.runtime.time_open",
        "positioner_report_info.time_open",
        "runtime.time_open",
        "time_open",
    ),
    "runtime.time_closed": (
        "positioner_report_info.runtime.time_closed",
        "positioner_report_info.time_closed",
        "runtime.time_closed",
        "time_closed",
    ),
    "runtime.time_near_closed": (
        "positioner_report_info.runtime.time_near_closed",
        "positioner_report_info.time_near_closed",
        "runtime.time_near_closed",
        "time_near_closed",
    ),
    "file.report_start_time": (
        "positioner_report_info.file.report_start_time",
        "positioner_report_info.report_start_time",
        "file.report_start_time",
        "report_start_time",
    ),
    "file.report_end_time": (
        "positioner_report_info.file.report_end_time",
        "positioner_report_info.report_end_time",
        "positioner_report_info.report_time",
        "file.report_end_time",
        "report_end_time",
        "report_time",
        "detection_time",
    ),
    "file.report_time": (
        "positioner_report_info.file.report_time",
        "positioner_report_info.report_time",
        "file.report_time",
        "report_time",
    ),
    "status.active_faults": (
        "positioner_report_info.status.active_faults",
        "positioner_report_info.active_faults",
        "status.active_faults",
        "active_faults",
    ),
    "status.historical_faults": (
        "positioner_report_info.status.historical_faults",
        "positioner_report_info.historical_faults",
        "status.historical_faults",
        "historical_faults",
    ),
    "status.task_execution": (
        "positioner_report_info.status.task_execution",
        "positioner_report_info.status.test_execution",
        "positioner_report_info.task_execution",
        "status.task_execution",
        "status.test_execution",
        "task_execution",
        "test_execution",
        "test_execution_issues",
    ),
}


def _manifest_nested_value(mapping: Any, path: str) -> tuple[bool, Any]:
    current = mapping
    for part in path.split("."):
        if not isinstance(current, dict) or part not in current:
            return False, None
        current = current[part]
    if current is None or current == "":
        return False, None
    return True, _manifest_scalar(current)


def resolve_report_evidence_value(
    structured: dict[str, Any] | None, path: str
) -> tuple[bool, Any]:
    """Resolve one report-contract value from the frozen structured evidence.

    ``material_extraction.fields`` is authoritative because it records the
    selected normalized value and its source.  The source-record views are
    compatibility fallbacks for older bundles and flat synthetic fixtures.
    Presence is tested explicitly so ``0`` and ``[]`` remain valid evidence.
    """
    if not isinstance(structured, dict) or not path:
        return False, None

    material_extraction = structured.get("material_extraction")
    fields = (
        material_extraction.get("fields", {})
        if isinstance(material_extraction, dict)
        else {}
    )
    if isinstance(fields, dict):
        field = fields.get(path)
        if isinstance(field, dict):
            selected = field.get("selected")
            if isinstance(selected, dict) and "normalized_value" in selected:
                value = selected.get("normalized_value")
                if value is not None and value != "":
                    return True, _manifest_scalar(value)
            if "normalized_value" in field:
                value = field.get("normalized_value")
                if value is not None and value != "":
                    return True, _manifest_scalar(value)

    candidates = (path,) + _REPORT_EVIDENCE_ALIASES.get(path, ())
    source_record = structured.get("source_record")
    views: list[dict[str, Any]] = [structured]
    if isinstance(source_record, dict):
        views.extend((source_record, _manifest_contract_evidence(structured)))
    for view in views:
        for candidate in candidates:
            present, value = _manifest_nested_value(view, candidate)
            if present:
                return True, value
    return False, None


def report_field_manifest(valve: Any, structured: dict[str, Any]) -> list[dict[str, Any]]:
    """Bind required customer-visible values to the renderer's formatted text."""
    present_values: dict[str, Any] = {}
    for path in REPORT_FIELD_MANIFEST_PATHS:
        _, present_values[path] = resolve_report_evidence_value(structured, path)

    device_rows = dict(_device_rows(valve, ""))
    key_rows, remaining = _metric_rows(valve)
    key_values = dict(key_rows)
    runtime_rows = dict(_positioner_runtime_rows(valve)) if valve.positioner_report_info else {}
    report_time = device_rows.get("检测时间", "未提供")
    snapshot = runtime_rows.get("定位器运行快照", "未提供")
    task_issues = report_test_execution_issues(valve)
    task_display = (
        "；  ".join(
            f"•  {_compact_test_execution_issue_text(item)}"
            for item in task_issues
        )
        or "未发现非成功任务"
    )
    display_values: dict[str, str] = {
        "positioner.option_config": snapshot,
        "positioner.device_mode": snapshot,
        "positioner.valve_position": snapshot,
        "positioner.manual_setpoint": snapshot,
        "positioner.input_signal": snapshot,
        "positioner.board_temperature": snapshot,
        "positioner.min_temperature": snapshot,
        "positioner.max_temperature": snapshot,
        "runtime.strokes": runtime_rows.get("累计行程次数", remaining.get("累计行程次数", "未提供")),
        "runtime.cycles": runtime_rows.get("累计循环次数", remaining.get("累计循环次数", "未提供")),
        "runtime.time_open": runtime_rows.get("累计全开时间", remaining.get("累计全开时间", "未提供")),
        "runtime.time_closed": runtime_rows.get("累计全关时间", remaining.get("累计全关时间", "未提供")),
        "runtime.time_near_closed": runtime_rows.get("累计近关时间", remaining.get("累计近关时间", "未提供")),
        "file.report_start_time": report_time,
        "file.report_end_time": report_time,
        "file.report_time": report_time,
        "status.active_faults": runtime_rows.get("活动故障", "未见活动故障"),
        "status.historical_faults": re.sub(
            r"\bAuto\s+tune\b",
            "Autotune",
            runtime_rows.get("历史故障", "未见历史故障"),
        ),
        "status.task_execution": task_display,
    }
    ramp_labels = {
        ("hysteresis", "max"): "斜坡测试回差最大值",
        ("hysteresis", "average"): "斜坡测试回差平均值",
        ("hysteresis", "min"): "斜坡测试回差最小值",
        ("dead_band", "max"): "斜坡测试死区最大值",
        ("dead_band", "average"): "斜坡测试死区平均值",
        ("dead_band", "min"): "斜坡测试死区最小值",
        ("linearity", "max"): "斜坡测试线性度最大值",
        ("linearity", "average"): "斜坡测试线性度平均值",
        ("linearity", "min"): "斜坡测试线性度最小值",
    }
    for (_metric, _statistic), label in ramp_labels.items():
        display_values[next(path for path in REPORT_FIELD_MANIFEST_PATHS if path.endswith("." + _metric + "." + _statistic))] = key_values.get(label, "未提供")

    return [
        {
            "path": path,
            "evidence_value": present_values[path],
            "display_text": str(display_values.get(path, "未提供")),
        }
        for path in REPORT_FIELD_MANIFEST_PATHS
    ]


def _add_device_table(doc: Document, valve: Any, unit_name: str, *, compact: bool = False) -> None:
    rows = _device_rows(valve, unit_name)
    table = doc.add_table(rows=math.ceil(len(rows) / 2), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, width in enumerate([1.32, 2.02, 1.45, 2.41]):
        table.columns[index].width = Inches(width)
    for index, (label, value) in enumerate(rows):
        row = index // 2
        column = (index % 2) * 2
        for cell in (table.cell(row, column), table.cell(row, column + 1)):
            _borders(cell)
            vertical_margin = 8 if compact else 35
            _margins(cell, vertical_margin, 55, vertical_margin, 55)
        _shade(table.cell(row, column), MINT)
        _shade(table.cell(row, column + 1), WHITE)
        device_size = (
            DASHBOARD_COMPACT_SIZE if compact else COMPACT_BODY_SIZE
        )
        _cell_text(
            table.cell(row, column),
            label,
            device_size,
            bold=True,
            color=BRAND_PRIMARY,
        )
        _cell_text(table.cell(row, column + 1), value, device_size)
    for row in table.rows:
        _set_row_min_height(row, 0.19 if compact else 0.23)


def _add_issue_action_table(
    doc: Document,
    valve: Any,
    *,
    blank_ratings: bool = False,
    blank_actions: bool = False,
    compact: bool = False,
) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(table, [2.7, 4.5])
    for index, (title, values, fill) in enumerate((
        ("主要问题", valve.issues or [valve.main_issue], "F1F7EE"),
        ("建议行动", [] if blank_actions else (valve.actions or [valve.action]), "E6F3F0"),
    )):
        cell = table.cell(0, index)
        _shade(cell, fill)
        _borders(cell)
        vertical_margin = 8 if compact else 60
        _margins(cell, vertical_margin, 95, vertical_margin, 95)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(2)
        _font(
            paragraph.add_run(title),
            8.8 if compact else 9.6,
            bold=True,
            color=BRAND_PRIMARY,
        )
        for value in values:
            paragraph = cell.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.08)
            paragraph.paragraph_format.first_line_indent = Inches(-0.08)
            paragraph.paragraph_format.space_after = Pt(0)
            _font(
                paragraph.add_run("•  "),
                DASHBOARD_BODY_SIZE if compact else BODY_SIZE,
                bold=True,
                color=BRAND_SECONDARY if blank_ratings else RATING_COLORS[int(valve.rating)],
            )
            _font(
                paragraph.add_run(value),
                DASHBOARD_BODY_SIZE if compact else BODY_SIZE,
            )
    _set_row_min_height(table.rows[0], 0.52 if compact else 0.63)


def _add_key_metrics(
    doc: Document,
    groups: dict[str, Any],
    *,
    compact: bool = False,
) -> None:
    table = doc.add_table(rows=6, cols=14)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(
        table,
        [
            0.70,
            0.70,
            0.46,
            0.46,
            0.10,
            0.76,
            0.52,
            0.52,
            0.52,
            0.10,
            0.525,
            0.525,
            0.635,
            0.635,
        ],
    )
    for row in table.rows:
        _keep_row_together(row)

    card_ranges = ((0, 3), (5, 8), (10, 13))
    card_titles = ("阶跃测试", "斜坡测试", "执行机构性能测试")
    for (start, end), title in zip(card_ranges, card_titles):
        cell = table.cell(0, start).merge(table.cell(0, end))
        _shade(cell, BRAND_PRIMARY)
        _borders(cell, BRAND_PRIMARY)
        _margins(cell, 20 if compact else 45, 70, 20 if compact else 45, 70)
        _cell_text(
            cell,
            title,
            9.5 if compact else 10.0,
            bold=True,
            color=WHITE,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for row_index, (label_text, value_text) in enumerate(groups["step"], start=1):
        label = table.cell(row_index, 0).merge(table.cell(row_index, 1))
        value = table.cell(row_index, 2).merge(table.cell(row_index, 3))
        fill = WHITE if row_index % 2 else PALE
        for cell in (label, value):
            _borders(cell)
            _margins(cell, 28 if compact else 42, 55, 28 if compact else 42, 55)
            _shade(cell, fill)
        _cell_text(
            label,
            label_text,
            8.6 if compact else COMPACT_BODY_SIZE,
            bold=True,
            color=BRAND_PRIMARY,
        )
        _cell_text(
            value,
            value_text,
            9.4 if compact else BODY_SIZE,
            bold=True,
            color=BRAND_PRIMARY,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for column, header in enumerate(groups["ramp_headers"], start=5):
        cell = table.cell(1, column)
        _borders(cell)
        _margins(cell, 27 if compact else 42, 35, 27 if compact else 42, 35)
        _shade(cell, MINT)
        _cell_text(
            cell,
            header,
            8.2 if compact else COMPACT_BODY_SIZE,
            bold=True,
            color=BRAND_PRIMARY,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for row_index, row_values in enumerate(groups["ramp"], start=2):
        fill = WHITE if row_index % 2 else PALE
        for column, value_text in enumerate(row_values, start=5):
            cell = table.cell(row_index, column)
            _borders(cell)
            _margins(cell, 34 if compact else 45, 35, 34 if compact else 45, 35)
            _shade(cell, fill)
            _cell_text(
                cell,
                value_text,
                (8.7 if column == 5 else 9.0)
                if compact
                else COMPACT_BODY_SIZE,
                bold=column > 5,
                color=TEXT if column == 5 else BRAND_PRIMARY,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
    for column in range(5, 9):
        cell = table.cell(5, column)
        _borders(cell)
        _margins(cell, 0, 0, 0, 0)
        _shade(cell, WHITE)
        _cell_text(cell, "", 8.2)

    for row_index, (label_text, value_text) in enumerate(
        groups["performance"], start=1
    ):
        label = table.cell(row_index, 10).merge(table.cell(row_index, 11))
        value = table.cell(row_index, 12).merge(table.cell(row_index, 13))
        warning = row_index == 1 and value_text != "有数据"
        fill = "FFF4E5" if warning else (WHITE if row_index % 2 else PALE)
        for cell in (label, value):
            _borders(cell, "D96C0B" if warning else LINE)
            _margins(cell, 24 if compact else 40, 50, 24 if compact else 40, 50)
            _shade(cell, fill)
        _cell_text(
            label,
            label_text,
            8.25 if compact else COMPACT_BODY_SIZE,
            bold=True,
            color=BRAND_PRIMARY,
        )
        _cell_text(
            value,
            value_text,
            (8.35 if row_index == 2 else 8.8)
            if compact
            else BODY_SIZE,
            bold=row_index != 2,
            color="8A4A10" if warning else BRAND_PRIMARY,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for spacer_column in (4, 9):
        for row_index in range(6):
            spacer = table.cell(row_index, spacer_column)
            _no_borders(spacer)
            _margins(spacer, 0, 0, 0, 0)
            _cell_text(spacer, "", 1.0)
    _set_row_min_height(table.rows[0], 0.22 if compact else 0.28)
    _repeat_table_header(table.rows[0])
    _repeat_table_header(table.rows[1])
    for row in table.rows[1:]:
        _set_row_min_height(row, 0.22 if compact else 0.30)


def _add_runtime_tables(
    doc: Document,
    remaining: dict[str, str],
    *,
    compact: bool = False,
) -> None:
    fault_key = next(
        (
            key
            for key in ("DTM活动故障", "当前故障", "活动故障")
            if key in remaining
        ),
        "",
    )
    _section_label(doc, "运行与定位器诊断信息", size=SECTION_TITLE_SIZE)
    runtime_keys = [
        "累计行程次数",
        "累计循环次数",
        "累计全开时间",
        "累计全关时间",
        "累计近关时间",
    ]
    if not any(key in remaining for key in runtime_keys):
        runtime_keys = ["累计行程 / 循环次数", "全开 / 全关 / 近关时间"]
    short_keys = runtime_keys
    short_rows = [(key, remaining.pop(key, "未提供")) for key in short_keys]
    long_rows: list[tuple[str, str]] = []
    summary_parts: list[str] = []
    if fault_key:
        summary_parts.append(f"活动故障：{remaining.pop(fault_key)}")
    snapshot_key = next(
        (
            key
            for key in ("DTM运行快照", "运行快照", "定位器运行快照")
            if key in remaining
        ),
        "",
    )
    if snapshot_key:
        summary_parts.append(f"快照：{remaining.pop(snapshot_key)}")
    if "历史故障" in remaining:
        summary_parts.append(f"历史故障：{remaining.pop('历史故障')}")
    if summary_parts:
        summary_text = re.sub(r"\bAuto\s+tune\b", "Autotune", "；".join(summary_parts))
        long_rows.append(("运行与故障摘要", summary_text))
    if "定位器使用时间 / 运行时间" in remaining:
        long_rows.append(
            (
                "定位器使用时间 / 运行时间",
                remaining.pop("定位器使用时间 / 运行时间"),
            )
        )
    long_rows.extend(remaining.items())

    short_table_rows = math.ceil(len(short_rows) / 2)
    table = doc.add_table(rows=short_table_rows + len(long_rows), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(table, [1.18, 2.40, 1.18, 2.40])
    for index, (key, value_text) in enumerate(short_rows):
        row = index // 2
        column = (index % 2) * 2
        label = table.cell(row, column)
        value = table.cell(row, column + 1)
        for cell in (label, value):
            _borders(cell)
            vertical_margin = 18 if compact else 30
            _margins(cell, vertical_margin, 45, vertical_margin, 45)
        _shade(label, MINT)
        _shade(value, WHITE)
        _cell_text(
            label,
            key,
            8.6 if compact else COMPACT_BODY_SIZE,
            bold=True,
            color=BRAND_PRIMARY,
        )
        _cell_text(value, value_text, 8.9 if compact else COMPACT_BODY_SIZE)
    for row in table.rows[:short_table_rows]:
        _set_row_min_height(row, 0.20 if compact else 0.24)

    for index, (label_text, value_text) in enumerate(long_rows):
        row_index = index + short_table_rows
        label = table.cell(row_index, 0)
        value = table.cell(row_index, 1).merge(table.cell(row_index, 3))
        for cell in (label, value):
            _borders(cell)
            vertical_margin = (
                0 if compact or label_text == "运行与故障摘要" else 30
            )
            _margins(cell, vertical_margin, 45, vertical_margin, 45)
        _shade(label, BRAND_SECONDARY if label_text == "运行与故障摘要" else MINT)
        _shade(value, PALE if index % 2 == 0 else WHITE)
        _cell_text(
            label,
            label_text,
            8.6 if compact else COMPACT_BODY_SIZE,
            bold=True,
            color=WHITE if label_text == "运行与故障摘要" else BRAND_PRIMARY,
        )
        _cell_text(
            value,
            value_text,
            (
                8.7
                if compact and label_text == "运行与故障摘要"
                else 8.6
                if compact
                else 8.1
                if label_text == "运行与故障摘要"
                else COMPACT_BODY_SIZE
            ),
        )
        _set_row_min_height(
            table.rows[row_index],
            (
                0.16
                if compact and label_text == "运行与故障摘要"
                else 0.18
                if compact
                else 0.18
                if label_text == "运行与故障摘要"
                else 0.24
            ),
        )


    # Keep the short runtime block with its heading, including long fault text.
    for index, row in enumerate(table.rows):
        _keep_row_together(row)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = index < len(table.rows) - 1


def _add_dashboard(
    doc: Document,
    valve: Any,
    unit_name: str,
    *,
    blank_ratings: bool = False,
    blank_conclusions: bool = False,
    blank_actions: bool = False,
    data_only: bool = False,
    rating_conclusion: str = "",
) -> None:
    compact = True

    title = doc.add_table(rows=1, cols=2)
    title.alignment = WD_TABLE_ALIGNMENT.CENTER
    title.autofit = False
    title.columns[0].width = Inches(5.85)
    title.columns[1].width = Inches(1.35)
    left, right = title.cell(0, 0), title.cell(0, 1)
    _shade(left, BRAND_PRIMARY)
    _shade(
        right,
        BRAND_PRIMARY
        if data_only
        else (BRAND_SECONDARY if blank_ratings else RATING_COLORS[int(valve.rating)]),
    )
    for cell in (left, right):
        _borders(cell, WHITE, "0")
        vertical_margin = 35 if compact else 95
        _margins(cell, vertical_margin, 150, vertical_margin, 150)
    left.text = ""
    paragraph = left.paragraphs[0]
    add_bookmark(paragraph, f"valve_{valve.folder}")
    _font(
        paragraph.add_run(valve.folder),
        DASHBOARD_TITLE_SIZE if compact else PAGE_TITLE_SIZE,
        bold=True,
        color=WHITE,
    )
    paragraph.add_run("\n")
    _font(
        paragraph.add_run(valve_model_spec(valve)),
        DASHBOARD_COMPACT_SIZE if compact else 8.2,
        color="D2ECE6",
    )
    right.text = ""
    paragraph = right.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not data_only:
        _font(
            paragraph.add_run("" if blank_ratings else f"{valve.rating}级"),
            DASHBOARD_TITLE_SIZE if compact else PAGE_TITLE_SIZE,
            bold=True,
            color=WHITE,
        )
        paragraph.add_run("\n")
        _font(paragraph.add_run("健康状态"), 7.2, color=WHITE)

        _section_label(doc, "诊断结论", size=DASHBOARD_SECTION_SIZE)
        has_assessment = (
            not blank_conclusions
            and bool(valve.overall_assessment)
            and valve.health_rating_detail_mode != "summary_only"
        )
        conclusion = doc.add_table(rows=2 if has_assessment else 1, cols=2)
        conclusion.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_column_widths(conclusion, [1.35, 5.85])
        conclusion_rows = [
            (
                "评级结论",
                "" if blank_conclusions else (
                    rating_conclusion
                ),
            )
        ]
        if has_assessment:
            conclusion_rows.append(("综合诊断结论", valve.overall_assessment))
        for row_index, (label_text, value_text) in enumerate(conclusion_rows):
            label = conclusion.cell(row_index, 0)
            value = conclusion.cell(row_index, 1)
            _shade(label, BRAND_PRIMARY if row_index else BRAND_SECONDARY)
            _shade(value, "E4F2EE" if row_index == 0 else PALE)
            _cell_text(
                label,
                label_text,
                DASHBOARD_COMPACT_SIZE,
                bold=True,
                color=WHITE,
            )
            _cell_text(
                value,
                value_text,
                DASHBOARD_BODY_SIZE,
                bold=row_index == 0,
                color=BRAND_PRIMARY if row_index == 0 else TEXT,
            )
            for conclusion_cell in (label, value):
                _borders(conclusion_cell, "9FC8BE")
                _margins(conclusion_cell, 12, 75, 12, 75)
            _set_row_min_height(conclusion.rows[row_index], 0.21)

    test_execution_issues = report_test_execution_issues(valve)
    if test_execution_issues:
        _section_label(
            doc,
            "测试执行异常（未完成/失败）",
            size=DASHBOARD_SECTION_SIZE,
        )
        issue_table = doc.add_table(rows=1, cols=1)
        issue_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        issue_cell = issue_table.cell(0, 0)
        _shade(issue_cell, "FFF4E5")
        _borders(issue_cell, "D96C0B")
        _margins(issue_cell, 8, 100, 8, 100)
        issue_cell.text = ""
        paragraph = issue_cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for index, item in enumerate(test_execution_issues):
            if index:
                _font(
                    paragraph.add_run("；  "),
                    DASHBOARD_COMPACT_SIZE,
                    color=MUTED,
                )
            _font(
                paragraph.add_run("•  "),
                DASHBOARD_BODY_SIZE,
                bold=True,
                color="D96C0B",
            )
            _font(
                paragraph.add_run(_compact_test_execution_issue_text(item)),
                DASHBOARD_BODY_SIZE,
            )
    elif not data_only and report_test_execution_evidence_present(valve):
        _section_label(doc, "测试执行状态", size=DASHBOARD_SECTION_SIZE)
        status_table = doc.add_table(rows=1, cols=1)
        status_cell = status_table.cell(0, 0)
        _shade(status_cell, PALE)
        _borders(status_cell)
        _margins(status_cell, 8, 100, 8, 100)
        _cell_text(status_cell, "未发现非成功任务", DASHBOARD_BODY_SIZE)

    _section_label(doc, "设备概览", size=DASHBOARD_SECTION_SIZE)
    _add_device_table(doc, valve, unit_name, compact=compact)
    if not data_only:
        _section_label(
            doc,
            "主要问题与建议行动",
            size=DASHBOARD_SECTION_SIZE,
        )
        _add_issue_action_table(
            doc,
            valve,
            blank_ratings=blank_ratings,
            blank_actions=blank_actions,
            compact=compact,
        )
    _section_label(
        doc,
        "关键测试数据" if data_only else "关键诊断数据",
        size=DASHBOARD_SECTION_SIZE,
    )
    metric_groups, remaining = _metric_groups(valve)
    _add_key_metrics(doc, metric_groups, compact=compact)
    _add_runtime_tables(doc, remaining, compact=compact)


def _curve_text(image: dict[str, Any]) -> tuple[str, str]:
    review = image.get("curve_review")
    if review is not None:
        return review.customer_text, review.image_key
    return "数据不足，本图需结合原始诊断记录复核。", ""


def _add_curve_pages(
    doc: Document,
    valve: Any,
    *,
    blank_ratings: bool = False,
    data_only: bool = False,
) -> None:
    images = sort_test_images(valve.images)
    if not images:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.keep_with_next = True
        heading = (
            f"{valve.folder} | 测试曲线与原始数据"
            if data_only
            else f"{valve.folder} | 曲线证据与图形判读"
        )
        _font(paragraph.add_run(heading), 13, bold=True, color=BRAND_PRIMARY)
        paragraph = doc.add_paragraph()
        message = (
            "未获取到可插入的测试曲线图。"
            if data_only
            else "未获取到可插入的诊断曲线图，数据不足，需复核诊断报告。"
        )
        _font(paragraph.add_run(message), 8.5)
        if not data_only:
            paragraph.paragraph_format.space_after = Pt(10)
            notice_rows = (
                ("数据状态", "未获取到可插入的诊断曲线图。"),
                (
                    "判读限制",
                    "当前无法进行逐图判读，相关结论按数据不足处理。",
                ),
                (
                    "后续处理",
                    "补充原始诊断曲线后，重新完成逐图复核。",
                ),
            )
            table = doc.add_table(rows=len(notice_rows), cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            _set_column_widths(table, [1.35, 5.85])
            for row, (label_text, value_text) in zip(
                table.rows,
                notice_rows,
                strict=True,
            ):
                label, value = row.cells
                for cell in (label, value):
                    _borders(cell)
                    _margins(cell, 100, 110, 100, 110)
                _shade(label, BRAND_PRIMARY)
                _shade(value, PALE)
                _cell_text(
                    label,
                    label_text,
                    9,
                    bold=True,
                    color=WHITE,
                )
                _cell_text(value, value_text, 9)
                _set_row_min_height(row, 0.55)
        return
    curves_per_page = 2
    test_key_counts = Counter(
        str(image.get("test_key") or "other") for image in images
    )
    for page_start in range(0, len(images), curves_per_page):
        if page_start:
            doc.add_page_break()
        paragraph = doc.add_paragraph()
        if page_start == 0:
            paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_after = Pt(2)
        _font(paragraph.add_run(valve.folder), 13, bold=True, color=BRAND_PRIMARY)
        _font(
            paragraph.add_run(
                " | 测试曲线与原始数据"
                if data_only
                else " | 曲线证据与图形判读"
            ),
            13,
            color=MUTED,
        )
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_after = Pt(4)
        _font(
            paragraph.add_run(
                "原始曲线与测试参数保留；右侧图形判读栏留空。"
                if data_only
                else "原始曲线保留；右侧文字为已校准的图形判读。"
            ),
            7.8,
            color=MUTED,
        )
        page_images = images[page_start:page_start + curves_per_page]
        for local_index, image in enumerate(page_images, 1):
            figure_number = page_start + local_index
            table = doc.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            _set_column_widths(table, [4.35, 2.85])
            left, right = table.cell(0, 0), table.cell(0, 1)
            for cell in (left, right):
                _borders(cell)
            _margins(left, 25, 25, 25, 25)
            _margins(right, 70, 85, 70, 85)
            _shade(right, MINT)
            left.text = ""
            paragraph = left.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                add_customer_picture(
                    paragraph.add_run(),
                    image["path"],
                    width=Inches(4.05),
                    picture_name=(
                        f"diagnostic-curve-{valve.folder}-"
                        f"{figure_number:02d}.png"
                    ),
                )
            except Exception as exc:  # noqa: BLE001 - customer output must retain a visible failure.
                _font(paragraph.add_run(f"图片插入失败：{exc}"), 7, color=RATING_COLORS[4])

            right.text = ""
            paragraph = right.paragraphs[0]
            _font(
                paragraph.add_run(f"{figure_number:02d}"),
                12.5,
                bold=True,
                color=(
                    BRAND_SECONDARY
                    if data_only or blank_ratings
                    else RATING_COLORS[int(valve.rating)]
                ),
            )
            paragraph.add_run("\n")
            test_key = image.get("test_key", "other")
            label = TEST_LABELS.get(test_key, TEST_LABELS["other"])
            _font(paragraph.add_run(label), 8.0, bold=True, color=BRAND_PRIMARY)
            if test_key_counts[str(test_key)] > 1:
                time_label = curve_execution_time_label(image)
                if time_label:
                    paragraph.add_run("\n")
                    _font(
                        paragraph.add_run(f"检测时间：{time_label}"),
                        6.7,
                        color=MUTED,
                    )
            paragraph = right.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(3)
            _font(paragraph.add_run("图形判读"), 7.3, bold=True, color=MUTED)
            interpretation, image_key = _curve_text(image)
            paragraph = right.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            if not data_only:
                _font(paragraph.add_run(interpretation), 7.4)
            if image_key:
                add_bookmark(paragraph, curve_review_marker(image_key))
            if data_only:
                paragraph = right.add_paragraph()
                _font(
                    paragraph.add_run(
                        f"原始报告第 {image.get('page', '未提供')} 页"
                    ),
                    6.5,
                    color=MUTED,
                )
            if local_index < len(page_images):
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(1)


def _page_field(page: Any, name: str, default: Any = "") -> Any:
    if isinstance(page, dict):
        return page.get(name, default)
    return getattr(page, name, default)


def _report_attachment_pages(
    valve: Any,
    *,
    include_specifications: bool,
    include_air_piping_pages: bool,
    include_accessory_pages: bool,
) -> list[dict[str, Any]]:
    """Merge attachment roles by source hash and one-based physical page."""
    candidates: list[tuple[Any, str, tuple[str, ...]]] = []
    for page in getattr(valve, "report_pages", ()) or ():
        raw_roles = _page_field(page, "roles", ()) or ()
        if isinstance(raw_roles, str):
            raw_roles = (raw_roles,)
        roles = tuple(str(role) for role in raw_roles)
        candidates.append((page, str(_page_field(page, "image_path", "")), roles))
    for role, page_name, image_name in (
        ("specification", "spec_page", "spec_image"),
        ("air_piping", "air_piping_page", "air_piping_image"),
        ("accessory", "accessory_page", "accessory_image"),
    ):
        page = getattr(valve, page_name, None)
        image_path = str(getattr(valve, image_name, "") or "")
        if page is not None and image_path:
            candidates.append((page, image_path, (role,)))

    enabled = {
        "specification": include_specifications,
        "air_piping": include_air_piping_pages,
        "accessory": include_accessory_pages,
    }
    merged: dict[tuple[str, int], dict[str, Any]] = {}
    for page, image_path, roles in candidates:
        visible_roles = {role for role in roles if enabled.get(role, False)}
        if not visible_roles or not image_path:
            continue
        page_number = int(_page_field(page, "page", 0) or 0)
        source_sha256 = str(_page_field(page, "source_sha256", "") or "").strip()
        source = str(_page_field(page, "source", "") or "").strip()
        identity = (source_sha256 or f"legacy:{source}", page_number)
        target = merged.setdefault(
            identity,
            {
                "source": source,
                "source_sha256": source_sha256,
                "page": page_number,
                "roles": set(),
                "image_path": image_path,
            },
        )
        target["roles"].update(visible_roles)
        if not target["image_path"]:
            target["image_path"] = image_path

    role_order = {"specification": 0, "air_piping": 1, "accessory": 2}
    pages = list(merged.values())
    pages.sort(
        key=lambda item: (
            min(role_order[role] for role in item["roles"]),
            item["source"],
            item["page"],
        )
    )
    return pages


def _add_report_attachment_page(
    doc: Document,
    valve: Any,
    page: dict[str, Any],
    *,
    data_only: bool = False,
) -> None:
    roles = set(page["roles"])
    role_labels = []
    if "specification" in roles:
        role_labels.append("Specification")
    if "air_piping" in roles:
        role_labels.append("Air Piping")
    if "accessory" in roles:
        role_labels.append("Accessory")
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    _font(paragraph.add_run("技术资料附件"), 9.5, bold=True, color=MUTED)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    _font(paragraph.add_run(valve.folder), 15.5, bold=True, color=BRAND_PRIMARY)
    _font(paragraph.add_run(f" | {' / '.join(role_labels)}"), 8.5, color=MUTED)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    description = "可靠匹配的设备规格与气路技术资料原页。"
    _font(paragraph.add_run(description), 7.8, color=MUTED)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_customer_picture(
        paragraph.add_run(),
        page["image_path"],
        width=Inches(6.55),
        picture_name=(
            f"technical-page-{valve.folder}-{int(page['page']):03d}.png"
        ),
    )
    for role in sorted(roles):
        add_bookmark(
            paragraph,
            spec_insertion_marker(valve.folder, role),
        )


def generate_business_dashboard_docx(
    valves: Iterable[Any],
    *,
    output_path: Path,
    project_name: str,
    report_name: str,
    report_date: str,
    unit_name: str,
    rating_conclusions: Mapping[str, str],
    overall_conclusion: str,
    recommendations: list[str],
    report_period_label: str = "",
    review_label: str = "",
    data_only: bool = False,
    blank_ratings: bool = False,
    blank_conclusions: bool = False,
    blank_actions: bool = False,
    include_field_context_overview: bool = False,
    include_overview_judgement: bool = False,
    include_overview_priority: bool = False,
    include_retest_recommendations: bool = False,
    include_specifications: bool = True,
    include_air_piping_pages: bool = True,
    include_accessory_pages: bool = False,
    term_replacements: dict[str, str] | None = None,
    recommendations_heading: str = "后续复测、维护和检修建议",
    omit_retest_recommendations: bool = False,
) -> None:
    """Render a complete business-dashboard DOCX for the supplied valves."""

    valves = list(valves)
    if not valves:
        raise ValueError("business dashboard report requires at least one valve")
    if not data_only and not blank_conclusions:
        missing = [valve.folder for valve in valves if not rating_conclusions.get(valve.folder)]
        if missing:
            raise ValueError(
                "prepared rating conclusions are missing for: " + ", ".join(missing)
            )
        if not overall_conclusion:
            raise ValueError("prepared overall conclusion is missing")
    doc = Document()
    _setup_document(doc, project_name=project_name, data_only=data_only)
    _add_cover(
        doc,
        project_name=project_name,
        report_name=report_name,
        report_date=report_date,
        report_period_label=report_period_label,
        review_label=review_label,
        data_only=data_only,
    )
    doc.add_page_break()
    if data_only:
        _add_data_overview(doc, valves, unit_name=unit_name)
    else:
        _add_overview(
            doc,
            valves,
            unit_name=unit_name,
            blank_ratings=blank_ratings,
            blank_actions=blank_actions,
        )
        doc.add_page_break()
        _add_overall_conclusion(
            doc,
            valves,
            blank_ratings=blank_ratings,
            blank_conclusions=blank_conclusions,
            blank_actions=blank_actions,
            show_blank_recommendations=bool(review_label.strip()),
            overall_conclusion=overall_conclusion,
            recommendations=recommendations,
            recommendations_heading=recommendations_heading,
        )
    _page_break_before_next(doc)
    for index, valve in enumerate(valves):
        _add_dashboard(
            doc,
            valve,
            unit_name,
            blank_ratings=blank_ratings,
            blank_conclusions=blank_conclusions,
            blank_actions=blank_actions,
            data_only=data_only,
            rating_conclusion=rating_conclusions.get(valve.folder, ""),
        )
        _add_curve_pages(
            doc,
            valve,
            blank_ratings=blank_ratings,
            data_only=data_only,
        )
        attachment_pages = _report_attachment_pages(
            valve,
            include_specifications=include_specifications,
            include_air_piping_pages=include_air_piping_pages,
            include_accessory_pages=include_accessory_pages,
        )
        for page in attachment_pages:
            doc.add_page_break()
            _add_report_attachment_page(doc, valve, page, data_only=data_only)
        if index != len(valves) - 1:
            doc.add_page_break()
    if not include_field_context_overview:
        _remove_field_context_overviews(doc)
    if not include_overview_judgement:
        _remove_heading_and_following_tables(doc, "整体判断", table_count=1)
    if not include_overview_priority:
        _remove_overview_priority(doc)
    if not include_retest_recommendations:
        _remove_recommendation_section(doc, recommendations_heading)
    _apply_term_replacements(doc, term_replacements)
    if omit_retest_recommendations:
        _remove_unavailable_retest_recommendations(doc)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    _optimize_docx_png_media(output_path)
