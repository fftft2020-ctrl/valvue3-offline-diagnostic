from __future__ import annotations

import json
import hashlib
import re
import shutil
import zipfile
from collections import Counter
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table
from docx.text.paragraph import Paragraph


@dataclass(frozen=True)
class ReportArtifact:
    kind: str
    version: int
    docx: Path
    pdf: Path
    work_dir: Path


@dataclass(frozen=True)
class ReviewImport:
    version: int
    source_docx: Path
    returned_docx: Path
    change_json: Path
    change_markdown: Path
    overrides_json: Path
    report_overrides_json: Path
    curve_overrides_json: Path | None = None


_LABELS = {
    "review": "待人工校验版",
    "business": "最终商务版",
    "data": "纯数据版",
}
_EMPTY_REVIEW_LABELS = {"主要问题", "建议行动", "诊断结论"}

_HUMAN_TEXT_CORRECTIONS = {
    "安装争取": "安装正确",
    "影响基本正常": "响应基本正常",
    "仅仅54%": "仅54%",
    "气缸可能漏气串气4级": "气缸可能漏气或串气，4级",
}


def _correct_human_text(value: object) -> object:
    if isinstance(value, str):
        for before, after in _HUMAN_TEXT_CORRECTIONS.items():
            value = value.replace(before, after)
        return value
    if isinstance(value, list):
        return [_correct_human_text(item) for item in value]
    if isinstance(value, dict):
        return {key: _correct_human_text(item) for key, item in value.items()}
    return value


def allocate_report_artifact(
    project_dir: Path,
    kind: str,
    report_name: str,
    *,
    period_label: str = "",
) -> ReportArtifact:
    """Return the next project-local report paths without touching existing versions."""
    if kind not in _LABELS:
        raise ValueError(f"unknown report kind: {kind}")
    project_dir = project_dir.resolve()
    output_dir = project_dir / "output" / kind
    label = _LABELS[kind]
    safe_period = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', "_", period_label).strip(" ._-")
    period_segment = f"-{safe_period}" if safe_period else ""
    if kind == "data":
        pattern = re.compile(
            rf"^{re.escape(report_name + period_segment)}-{label}-v(\d+)\.(?:docx|pdf)$"
        )
        work_root = project_dir / "work" / kind / (safe_period or "default")
    else:
        version_label = "(?:待人工校验版|人工校验版)" if kind == "review" else label
        pattern = re.compile(
            rf"^{re.escape(report_name)}(?:-.+)?-{version_label}-v(\d+)\.(?:docx|pdf)$"
        )
        work_root = project_dir / "work" / kind
    versions = {
        int(match.group(1))
        for path in output_dir.glob(f"{report_name}-*")
        if (match := pattern.match(path.name))
    }
    versions.update(
        int(match.group(1))
        for path in work_root.glob("v*")
        if path.is_dir() and (match := re.fullmatch(r"v(\d+)", path.name))
    )
    version = max(versions, default=0) + 1
    stem = f"{report_name}{period_segment}-{label}-v{version}"
    work_dir = work_root / f"v{version}"
    return ReportArtifact(
        kind=kind,
        version=version,
        docx=output_dir / f"{stem}.docx",
        pdf=output_dir / f"{stem}.pdf",
        work_dir=work_dir,
    )


def append_run_history(project_dir: Path, event: dict[str, object]) -> Path:
    path = project_dir.resolve() / "work" / "run_history.jsonl"
    path.parent.mkdir(parents=True, exist_ok=True)
    record = {
        "created_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        **event,
    }
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n")
    return path


def _file_sha256(path: Path) -> str:
    with path.open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def _atomic_write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def _document_records(path: Path) -> list[str]:
    document = Document(path)
    records = [paragraph.text.strip() for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            records.extend(cell.text.strip() for cell in row.cells)
    return records


def compare_review_documents(original: Path, returned: Path) -> list[dict[str, object]]:
    """Compare visible content only; formatting changes never appear in the result."""
    before = _document_records(original)
    after = _document_records(returned)
    changes: list[dict[str, object]] = []
    matcher = SequenceMatcher(a=before, b=after, autojunk=False)
    for operation, first_start, first_end, second_start, second_end in matcher.get_opcodes():
        if operation == "equal":
            continue
        old_values = before[first_start:first_end]
        new_values = after[second_start:second_end]
        count = max(len(old_values), len(new_values))
        for index in range(count):
            changes.append(
                {
                    "operation": operation,
                    "before": old_values[index] if index < len(old_values) else "",
                    "after": new_values[index] if index < len(new_values) else "",
                    "before_index": first_start + index if index < len(old_values) else None,
                    "after_index": second_start + index if index < len(new_values) else None,
                }
            )
    return changes


def _iter_blocks(document: DocumentObject) -> Iterator[Paragraph | Table]:
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def _tag_in_text(text: str, valve_tags: set[str]) -> str:
    for tag in sorted(valve_tags, key=len, reverse=True):
        if re.search(rf"(?<![A-Za-z0-9-]){re.escape(tag)}(?![A-Za-z0-9-])", text):
            return tag
    return ""


def _dashboard_cell_values(text: str, label: str) -> list[str]:
    """Return non-empty dashboard cell lines after its visible title."""
    values: list[str] = []
    seen_label = False
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if not seen_label and line == label:
            seen_label = True
            continue
        if seen_label:
            cleaned = re.sub(r"^[•·\-]\s*", "", line).strip()
            if cleaned:
                values.append(cleaned)
    return values


def extract_review_overrides(path: Path, *, valve_tags: set[str]) -> dict[str, dict[str, object]]:
    """Extract human-entered report values while retaining unrecognised edits in the diff."""
    document = Document(path)
    result: dict[str, dict[str, object]] = {}
    current_tag = ""
    section = ""
    for block in _iter_blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            found = _tag_in_text(text, valve_tags)
            if found:
                current_tag = found
                result.setdefault(current_tag, {})
            if re.match(r"^4\.\s*问题点", text):
                section = "issues"
                continue
            if re.match(r"^5\.\s*行动项", text):
                section = "actions"
                continue
            if text == "诊断结论":
                section = "dashboard_conclusion"
                continue
            if text == "主要问题与建议行动":
                section = "dashboard_issue_action"
                continue
            if block.style and block.style.name.startswith("Heading"):
                if not found:
                    section = ""
                continue
            if current_tag and section in {"issues", "actions"} and text:
                values = result[current_tag].setdefault(section, [])
                if isinstance(values, list):
                    values.append(text)
                if section == "actions" and "action" not in result[current_tag]:
                    result[current_tag]["action"] = text
            continue

        if block.rows and len(block.rows[0].cells) == 2:
            first_text = block.rows[0].cells[0].text.strip()
            second_text = block.rows[0].cells[1].text.strip()
            dashboard_tag = _tag_in_text(first_text, valve_tags)
            if dashboard_tag and "健康状态" in second_text:
                current_tag = dashboard_tag
                override = result.setdefault(current_tag, {})
                match = re.search(r"([1-5])\s*级?", second_text)
                if match:
                    override["rating"] = int(match.group(1))
                section = ""
                continue

        if not current_tag:
            continue
        override = result.setdefault(current_tag, {})
        if section == "dashboard_conclusion" and len(block.rows) == 1 and len(block.rows[0].cells) == 1:
            value = block.rows[0].cells[0].text.strip()
            if value:
                override["rating_conclusion"] = value
            section = ""
            continue
        if section == "dashboard_issue_action" and block.rows and len(block.rows[0].cells) >= 2:
            issues = _dashboard_cell_values(block.rows[0].cells[0].text, "主要问题")
            actions = _dashboard_cell_values(block.rows[0].cells[1].text, "建议行动")
            if issues:
                override["issues"] = issues
            if actions:
                override["actions"] = actions
                override["action"] = actions[0]
            section = ""
            continue
        for row in block.rows:
            if len(row.cells) < 2:
                continue
            key = row.cells[0].text.strip()
            value = row.cells[1].text.strip()
            if not value:
                continue
            if key == "健康状态评级":
                match = re.search(r"([1-5])\s*级?", value)
                if match:
                    override["rating"] = int(match.group(1))
            elif key == "评级结论":
                override["rating_conclusion"] = value
                if section == "dashboard_conclusion":
                    section = ""
            elif key == "综合诊断结论":
                override["overall_assessment"] = value
    return {tag: values for tag, values in result.items() if values}


def extract_report_overrides(path: Path) -> dict[str, object]:
    """Extract report-level human conclusion and recommendation paragraphs."""
    document = Document(path)
    in_overall = False
    section = ""
    dashboard_mode = False
    conclusions: list[str] = []
    recommendations: list[str] = []
    for block in _iter_blocks(document):
        if isinstance(block, Table):
            if dashboard_mode and in_overall and section == "conclusion":
                if len(block.rows) == 1 and len(block.rows[0].cells) == 1:
                    value = block.rows[0].cells[0].text.strip()
                    if value:
                        conclusions.append(value)
                    section = ""
                continue
            if dashboard_mode and in_overall and section == "recommendations":
                in_overall = False
            continue
        text = block.text.strip()
        if (
            text == "2. 总体结论与检修建议"
            and block.style
            and block.style.name == "Heading 1"
        ):
            in_overall = True
            section = "conclusion"
            dashboard_mode = False
            continue
        if text == "总体结论与检修建议":
            in_overall = True
            section = "conclusion"
            dashboard_mode = True
            continue
        if not in_overall:
            continue
        if text == "后续复测、维护和检修建议":
            section = "recommendations"
            continue
        if block.style and block.style.name.startswith("Heading"):
            if section == "recommendations":
                in_overall = False
            else:
                section = ""
            continue
        if not text or text.startswith("本次共分析"):
            continue
        if section == "conclusion":
            conclusions.append(text)
        elif section == "recommendations":
            recommendations.append(text)
    result: dict[str, object] = {}
    if conclusions:
        result["overall_conclusion"] = "\n".join(conclusions)
    if recommendations:
        result["recommendations"] = recommendations
    return result


def _retain_explicit_clears(
    base_valves: dict[str, dict[str, object]],
    valves: dict[str, dict[str, object]],
    base_report: dict[str, object],
    report: dict[str, object],
    changes: list[dict[str, object]],
) -> set[str]:
    """Keep fields the user cleared while ignoring fields that were blank already."""
    empty_changes = [
        change
        for change in changes
        if str(change.get("before") or "").strip()
        and not {
            line.strip()
            for line in str(change.get("after") or "").splitlines()
            if line.strip()
        } - _EMPTY_REVIEW_LABELS
    ]
    mapped_before: set[str] = set()

    def was_cleared(value: object) -> bool:
        texts = value if isinstance(value, list) else str(value or "").splitlines()
        matched = [
            change
            for change in empty_changes
            if any(
                str(text or "").strip()
                and str(text).strip() in str(change.get("before") or "")
                for text in texts
            )
        ]
        mapped_before.update(
            str(change.get("before") or "").strip() for change in matched
        )
        return bool(matched)

    for tag, base in base_valves.items():
        current = valves.get(tag, {})
        for field in ("rating_conclusion", "overall_assessment"):
            if field in base and field not in current and was_cleared(base[field]):
                current = valves.setdefault(tag, current)
                current[field] = ""
        for field in ("issues", "actions"):
            if field in base and field not in current and was_cleared(base[field]):
                current = valves.setdefault(tag, current)
                current[field] = []
                if field == "actions":
                    current["action"] = ""
    for field in ("overall_conclusion", "recommendations"):
        if field in base_report and field not in report and was_cleared(base_report[field]):
            report[field] = "" if field == "overall_conclusion" else []
    return mapped_before


def _curve_test_key(text: str) -> str:
    if "Pattern Step" in text:
        return "pattern_step"
    if "Run Step" in text:
        return "run_step"
    if "Run 斜坡" in text or "位置动态性能测试" in text:
        return "run_ramp"
    if "扩展执行机构动态性能测试" in text:
        return "extended_signature"
    if "标准动态性能测试" in text:
        return "standard_signature"
    return "other"


def extract_curve_review_overrides(
    path: Path,
    *,
    valve_tags: set[str],
    manifest_path: Path,
    base_interpretations_path: Path,
    valve_overrides: dict[str, dict[str, object]],
) -> tuple[dict[str, object], set[str]]:
    """Map human-edited chart text onto the existing immutable image identities."""
    if not manifest_path.is_file() or not base_interpretations_path.is_file():
        return {}, set()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    base_raw = json.loads(base_interpretations_path.read_text(encoding="utf-8"))
    images = manifest.get("images") if isinstance(manifest, dict) else None
    base_entries = (
        base_raw
        if isinstance(base_raw, list)
        else base_raw.get("interpretations")
        if isinstance(base_raw, dict)
        else None
    )
    if not isinstance(images, list) or not isinstance(base_entries, list):
        raise RuntimeError("curve review inputs have an invalid structure")
    if not images:
        return {}, set()
    base_by_key = {
        str(item.get("image_key") or ""): item
        for item in base_entries
        if isinstance(item, dict)
    }
    image_groups: dict[tuple[str, str], list[dict[str, object]]] = {}
    for image in images:
        if not isinstance(image, dict):
            continue
        key = (str(image.get("valve_tag") or ""), str(image.get("test_key") or ""))
        image_groups.setdefault(key, []).append(image)

    document = Document(path)
    current_tag = ""
    text_groups: dict[tuple[str, str], list[tuple[str, str]]] = {}
    for block in _iter_blocks(document):
        text = block.text.strip() if isinstance(block, Paragraph) else "\n".join(
            cell.text for row in block.rows for cell in row.cells
        )
        found = _tag_in_text(text, valve_tags)
        if found:
            current_tag = found
        if not isinstance(block, Table) or not current_tag or not block.rows:
            continue
        cells = block.rows[0].cells
        if len(cells) < 2 or "图形判读" not in cells[-1].text:
            continue
        full_text = cells[-1].text.strip()
        lines = [line.strip() for line in full_text.splitlines()]
        try:
            marker = lines.index("图形判读")
        except ValueError:
            continue
        customer_lines = [
            line for line in lines[marker + 1:]
            if line and not re.match(r"^诊断报告第\s*\d+\s*页$", line)
        ]
        if not customer_lines:
            raise RuntimeError(f"human curve interpretation is empty: {current_tag}")
        key = (current_tag, _curve_test_key("\n".join(lines[:marker])))
        text_groups.setdefault(key, []).append(("\n".join(customer_lines), full_text))

    entries: list[dict[str, object]] = []
    mapped_blocks: set[str] = set()
    for key, human_items in text_groups.items():
        group_images = image_groups.get(key, [])
        if len(group_images) != len(human_items):
            raise RuntimeError(
                f"human curve count does not match manifest for {key[0]}/{key[1]}: "
                f"{len(human_items)} != {len(group_images)}"
            )
        for image, (customer_text, full_text) in zip(group_images, human_items):
            image_key = str(image.get("image_key") or "")
            base = base_by_key.get(image_key)
            if not isinstance(base, dict):
                raise RuntimeError(f"base curve interpretation is missing: {image_key}")
            rating = int((valve_overrides.get(key[0]) or {}).get("rating") or base.get("rating") or 2)
            action = str((valve_overrides.get(key[0]) or {}).get("action") or base.get("action") or "按人工确认结论处理。")
            entries.append(
                {
                    **base,
                    "image_key": image_key,
                    "test_key": key[1],
                    "observation": customer_text,
                    "diagnosis": customer_text,
                    "rating": rating,
                    "action": action,
                    "customer_text": customer_text,
                    "confidence": "high",
                    "experience_rule_ids": list(base.get("experience_rule_ids") or []),
                }
            )
            mapped_blocks.add(full_text)
    if len(entries) != len(images):
        raise RuntimeError(f"human review covers {len(entries)} curves; expected {len(images)}")
    return {
        "schema_version": 1,
        "authoring_method": "confirmed_human_review_return",
        "interpretations": entries,
    }, mapped_blocks


def _latest_review_source(project_dir: Path) -> tuple[int, Path]:
    candidates = list((project_dir / "output/review").glob("*.docx"))
    if len(candidates) != 1:
        raise RuntimeError("review base is ambiguous or missing; specify --review-base")
    match = re.search(r"-v(\d+)\.docx$", candidates[0].name)
    return int(match.group(1)) if match else 0, candidates[0]


def review_base_binding(project_dir: Path, source_docx: Path, *, _seen=None, _expected_sha=None) -> dict[str, str]:
    """Resolve a registered report or a verified format-only descendant."""
    from .format_revision import compare_format_only_content, format_revision_manifest_path

    project_dir, source_docx = project_dir.resolve(), source_docx.resolve()
    if not source_docx.is_relative_to(project_dir / "output") or not source_docx.is_file():
        raise RuntimeError("review base must be an existing report inside project output")
    seen = set() if _seen is None else set(_seen)
    if source_docx in seen:
        raise RuntimeError("cyclic format-only report ancestry")
    seen.add(source_docx)
    sha = _file_sha256(source_docx)
    bindings = []
    for path in (project_dir / "work").rglob("run_manifest.json"):
        manifest = json.loads(path.read_text(encoding="utf-8"))
        outputs = manifest.get("outputs") or {}
        registered = outputs.get("docx") or manifest.get("output_docx")
        if not registered or (project_dir / registered).resolve() != source_docx:
            continue
        if (outputs.get("docx_sha256") or manifest.get("output_sha256")) != sha:
            raise RuntimeError("review source DOCX does not match its run manifest")
        formal = manifest.get("formal_diagnosis") or {}
        binding = {key: str(formal.get(key) or "") for key in ("tree_sha256", "project_manifest_sha256")}
        if not all(binding.values()):
            raise RuntimeError("review source run manifest has no formal diagnosis binding")
        bindings.append(binding)
    revision_path = format_revision_manifest_path(project_dir, source_docx)
    if revision_path.is_file():
        revision = json.loads(revision_path.read_text(encoding="utf-8"))
        parent = (project_dir / revision["parent_docx"]).resolve()
        if not parent.is_relative_to(project_dir / "output") or not parent.is_file():
            raise RuntimeError("invalid format-only report parent")
        if _file_sha256(parent) != revision.get("parent_sha256"):
            raise RuntimeError("format-only parent changed")
        if not compare_format_only_content(parent, source_docx)["ok"]:
            raise RuntimeError("format-only report content changed")
        # The prepared candidate hash predates formatting; delivery binds its final bytes.
        delivery_path = project_dir / "work/wps-validation" / source_docx.stem / "delivery_manifest.json"
        delivery = json.loads(delivery_path.read_text(encoding="utf-8")) if delivery_path.is_file() else {}
        registered_sha = (delivery.get("delivery") or {}).get("docx_sha256") or revision.get("candidate_sha256")
        if registered_sha != sha and _expected_sha != sha:
            raise RuntimeError("format-only report bytes are not registered")
        bindings.append(review_base_binding(project_dir, parent, _seen=seen, _expected_sha=revision["parent_sha256"]))
    if not bindings or any(binding != bindings[0] for binding in bindings):
        raise RuntimeError("review source run manifest is missing or has ambiguous binding")
    return bindings[0]


def _review_formal_diagnosis_binding(project_dir: Path, version: int, source_docx: Path) -> dict[str, str]:
    return review_base_binding(project_dir, source_docx)


def _review_import_from_path(path: Path) -> ReviewImport:
    data = json.loads(path.read_text(encoding="utf-8"))
    # Both legacy vN and returns/<id> live below project/work/review.
    project = next(parent.parent for parent in path.parents if parent.name == "work")
    curve = path.parent / "manual_curve_interpretations.json"
    return ReviewImport(
        version=int(data["review_version"]), source_docx=project / data["source_docx"],
        returned_docx=project / data["returned_docx"], change_json=path,
        change_markdown=path.with_suffix(".md"), overrides_json=path.parent / "diagnostic_overrides.json",
        report_overrides_json=path.parent / "report_overrides.json",
        curve_overrides_json=curve if curve.is_file() else None,
    )


def _copy_without_overwrite(source: Path, destination_dir: Path) -> Path:
    destination_dir.mkdir(parents=True, exist_ok=True)
    destination = destination_dir / source.name
    if not destination.exists():
        shutil.copy2(source, destination)
        return destination
    if _file_sha256(destination) == _file_sha256(source):
        return destination
    counter = 2
    while True:
        candidate = destination.with_name(f"{destination.stem}-{counter}{destination.suffix}")
        if not candidate.exists():
            shutil.copy2(source, candidate)
            return candidate
        counter += 1


def _docx_comments(path: Path) -> list[str]:
    try:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/comments.xml").decode("utf-8")
    except (KeyError, OSError, zipfile.BadZipFile):
        return []
    comments: list[str] = []
    for body in re.findall(r"<w:comment\b.*?</w:comment>", xml, re.S):
        text = "".join(re.findall(r"<w:t(?:\s[^>]*)?>(.*?)</w:t>", body, re.S))
        text = re.sub(r"\s+", " ", text).strip()
        if text:
            comments.append(text)
    return comments


def _review_document_structure(path: Path) -> dict[str, object]:
    """Changes outside mapped body text must not disappear during regeneration."""
    from xml.etree import ElementTree as ET
    from .format_revision import _relationship_map, _relationship_token
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as package:
        body = ET.fromstring(package.read("word/document.xml"))
        relationships = _relationship_map(package, "word/document.xml")
        relationship_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        return {
            "tables": [[len(row.findall(f"{w}tc")) for row in table.findall(f"{w}tr")]
                       for table in body.iter(f"{w}tbl")],
            "bookmarks": [node.attrib for node in body.iter() if node.tag in {f"{w}bookmarkStart", f"{w}bookmarkEnd"}],
            "links": [{key: (_relationship_token(package, "word/document.xml", relationships.get(value))
                              if key == relationship_id else value) for key, value in node.attrib.items()}
                      for node in body.iter(f"{w}hyperlink")],
            "headers_footers_notes": {name: [node.text for node in ET.fromstring(package.read(name)).iter(f"{w}t")]
                for name in package.namelist() if re.fullmatch(r"word/(?:header\d+|footer\d+|footnotes|endnotes)\.xml", name)},
        }


def ingest_review_return(
    project_dir: Path,
    returned_docx: Path,
    *,
    valve_tags: set[str],
    review_base: Path | None = None,
    expected_binding: dict[str, str] | None = None,
    report_period: str = "",
    curve_manifest_path: Path | None = None,
    base_interpretations_path: Path | None = None,
) -> ReviewImport:
    """Archive a human-edited DOCX and persist its content-only change list."""
    project_dir = project_dir.resolve()
    returned_docx = returned_docx.resolve()
    if not returned_docx.is_file():
        raise RuntimeError(f"human review DOCX is missing: {returned_docx}")
    if review_base is None:
        version, source_docx = _latest_review_source(project_dir)
    else:
        source_docx = (project_dir / review_base).resolve()
        match = re.search(r"-v(\d+)\.docx$", source_docx.name)
        version = int(match.group(1)) if match else 0
    binding = review_base_binding(project_dir, source_docx)
    if expected_binding is not None and binding != expected_binding:
        raise RuntimeError("review base does not match current formal diagnosis binding")
    source_sha, returned_sha = _file_sha256(source_docx), _file_sha256(returned_docx)
    return_id = hashlib.sha256(f"{source_sha}:{returned_sha}".encode()).hexdigest()
    review_dir = project_dir / "work/review/returns" / return_id
    if (review_dir / "change_list.json").is_file():
        existing = _review_import_from_path(review_dir / "change_list.json")
        if _file_sha256(existing.returned_docx) != returned_sha:
            raise RuntimeError("archived review return changed after ingestion")
        return existing
    archived = _copy_without_overwrite(returned_docx, project_dir / "experience" / "feedback_inbox")
    changes = compare_review_documents(source_docx, archived)
    overrides = extract_review_overrides(archived, valve_tags=valve_tags)
    report_overrides = extract_report_overrides(archived)
    mapped_empty_before = _retain_explicit_clears(
        extract_review_overrides(source_docx, valve_tags=valve_tags),
        overrides,
        extract_report_overrides(source_docx),
        report_overrides,
        changes,
    )
    curve_overrides, mapped_curve_blocks = extract_curve_review_overrides(
        archived,
        valve_tags=valve_tags,
        manifest_path=curve_manifest_path or Path(),
        base_interpretations_path=base_interpretations_path or Path(),
        valve_overrides=overrides,
    )
    mapped_values: set[str] = set()
    for override in overrides.values():
        for key, value in override.items():
            if key == "rating" and isinstance(value, int):
                mapped_values.update({str(value), f"{value}级"})
            elif isinstance(value, str):
                mapped_values.add(value)
            elif isinstance(value, list):
                mapped_values.update(str(item) for item in value)
    for value in report_overrides.values():
        if isinstance(value, str):
            mapped_values.add(value)
        elif isinstance(value, list):
            mapped_values.update(str(item) for item in value)
    mapped_values.update(mapped_curve_blocks)

    def mapped_change(change: dict[str, object]) -> bool:
        after = str(change.get("after") or "").strip()
        before = str(change.get("before") or "").strip()
        after_lines = {
            line.strip() for line in after.splitlines() if line.strip()
        }
        if before in mapped_empty_before and not after_lines - _EMPTY_REVIEW_LABELS:
            return True
        if not after:
            return before in mapped_empty_before
        if after in mapped_values:
            return True
        lines = [
            re.sub(r"^[•·\-]\s*", "", line).strip()
            for line in after.splitlines()
            if line.strip() not in {"主要问题", "建议行动", "健康状态"}
        ]
        if lines and all(line in mapped_values for line in lines):
            return True
        return before.replace("（当前身份待确认）", "") == after

    unresolved_changes = [
        change for change in changes
        if not mapped_change(change)
    ]
    overrides = _correct_human_text(overrides)
    report_overrides = _correct_human_text(report_overrides)
    curve_overrides = _correct_human_text(curve_overrides)
    change_json = review_dir / "change_list.json"
    change_markdown = review_dir / "change_list.md"
    overrides_json = review_dir / "diagnostic_overrides.json"
    report_overrides_json = review_dir / "report_overrides.json"
    curve_overrides_json = review_dir / "manual_curve_interpretations.json"
    payload = {
        "schema_version": 2,
        "return_id": return_id,
        "report_period": report_period,
        "review_version": version,
        "source_docx": str(source_docx.relative_to(project_dir)),
        "source_sha256": _file_sha256(source_docx),
        "returned_docx": str(archived.relative_to(project_dir)),
        "returned_sha256": _file_sha256(archived),
        "formal_diagnosis": binding,
        "changes": changes,
        "comments": _docx_comments(archived),
        "overrides": overrides,
        "report_overrides": report_overrides,
        "unresolved_changes": unresolved_changes,
    }
    from .format_revision import docx_content_identity
    before_identity, after_identity = docx_content_identity(source_docx), docx_content_identity(archived)
    payload["package_changes"] = [key for key in ("media_sha256", "embedded_sha256")
                                  if before_identity[key] != after_identity[key]]
    before_structure, after_structure = _review_document_structure(source_docx), _review_document_structure(archived)
    payload["structure_changes"] = [key for key in before_structure if before_structure[key] != after_structure[key]]
    if payload["structure_changes"]:
        payload["unresolved_changes"].append({"before": "", "after": "表格结构、书签、链接或页眉页脚脚注已更改，需核对保留"})
    if payload["package_changes"]:
        payload["unresolved_changes"].append({"before": "", "after": "图片或嵌入对象已更改，需重新绑定原始证据"})
    _atomic_write_json(change_json, payload)
    _atomic_write_json(overrides_json, overrides)
    _atomic_write_json(report_overrides_json, report_overrides)
    if curve_overrides:
        _atomic_write_json(curve_overrides_json, curve_overrides)
    lines = [f"# 人工回传 v{version} 修改清单", ""]
    for change in changes:
        lines.extend(
            [
                f"- 原文：{change['before'] or '（空）'}",
                f"  修改为：{change['after'] or '（空）'}",
            ]
        )
    for comment in payload["comments"]:
        lines.append(f"- 批注：{comment}")
    if unresolved_changes:
        lines.extend(["", "## 待确认的未映射修改", ""])
        for change in unresolved_changes:
            lines.append(f"- {change['after']}")
    change_markdown.parent.mkdir(parents=True, exist_ok=True)
    change_markdown.write_text("\n".join(lines) + "\n", encoding="utf-8")
    append_run_history(
        project_dir,
        {
            "event": "review_return_ingested",
            "version": version,
            "returned_docx": str(archived.relative_to(project_dir)),
            "returned_sha256": payload["returned_sha256"],
        },
    )
    return ReviewImport(
        version=version,
        source_docx=source_docx,
        returned_docx=archived,
        change_json=change_json,
        change_markdown=change_markdown,
        overrides_json=overrides_json,
        report_overrides_json=report_overrides_json,
        curve_overrides_json=curve_overrides_json if curve_overrides else None,
    )


def latest_review_import(project_dir: Path, *, binding: dict[str, str] | None = None,
                         review_return: Path | None = None, report_period: str | None = None) -> Path:
    root = project_dir.resolve() / "work/review"
    paths = [review_return.resolve()] if review_return is not None else [
        *root.glob("v*/change_list.json"), *root.glob("returns/*/change_list.json")]
    candidates = []
    for path in paths:
        if not path.is_relative_to(root) or not path.is_file():
            raise RuntimeError("review return must be a registered project change list")
        data = json.loads(path.read_text(encoding="utf-8"))
        if binding is not None and data.get("formal_diagnosis") != binding:
            continue
        if report_period is not None and data.get("report_period", report_period) != report_period:
            continue
        candidates.append(path)
    if len(candidates) != 1:
        raise RuntimeError("no unique matching human review return; specify --review-return")
    return candidates[0]


def _set_visible_text(target: Paragraph | object, value: str) -> None:
    paragraphs = [target] if isinstance(target, Paragraph) else list(target.paragraphs)
    first = paragraphs[0]
    run_properties = None
    if first.runs and first.runs[0]._r.rPr is not None:
        run_properties = deepcopy(first.runs[0]._r.rPr)
    for paragraph in paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)
    for run in list(first.runs):
        first._element.remove(run._element)
    run = first.add_run(value)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def _remove_paragraph(paragraph: Paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def finalize_human_review_document(
    source_docx: Path,
    destination_docx: Path,
    *,
    report_title: str,
    report_date: str,
    overrides: dict[str, dict[str, object]],
) -> dict[str, object]:
    """Preserve a human-edited report while synchronising its derived dashboard."""
    document = Document(source_docx)
    typo_fixes = {
        **_HUMAN_TEXT_CORRECTIONS,
        "需调整至数据表要求的供气压力\n": "需调整至数据表要求的供气压力。\n",
    }

    def clean(value: object) -> str:
        text = str(value or "")
        for before, after in typo_fixes.items():
            text = text.replace(before, after)
        return text.replace("（内部评审版，待用户确认）", "").replace(
            "（阀门无法全开，待用户确认）", "（阀门无法全开）"
        )

    overrides = deepcopy(overrides)
    for values in overrides.values():
        for key in ("rating_conclusion", "overall_assessment", "action"):
            if key in values:
                values[key] = clean(values[key])
        for key in ("issues", "actions"):
            if isinstance(values.get(key), list):
                values[key] = [clean(item) for item in values[key]]
    ratings = {
        tag: int(values["rating"])
        for tag, values in overrides.items()
        if isinstance(values.get("rating"), int)
    }
    if len(ratings) != len(overrides):
        raise RuntimeError("every human-reviewed valve requires a rating")
    counts = Counter(ratings.values())

    for paragraph in list(document.paragraphs):
        if paragraph.text.startswith("人工校验状态："):
            _remove_paragraph(paragraph)
            continue
        value = paragraph.text
        if value == "报告日期\n2026-08-25":
            value = f"报告日期\n{report_date}"
        if value == "原始曲线保留；右侧文字为已校准的图形判读。":
            value = "原始曲线保留；右侧为图形判读。"
        for before, after in typo_fixes.items():
            value = value.replace(before, after)
        if value != paragraph.text:
            _set_visible_text(paragraph, value)

    seen_cells: set[object] = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                identity = cell._tc
                if identity in seen_cells:
                    continue
                seen_cells.add(identity)
                for paragraph in cell.paragraphs:
                    value = paragraph.text
                    for before, after in typo_fixes.items():
                        value = value.replace(before, after)
                    value = value.replace("（内部评审版，待用户确认）", "")
                    value = value.replace("（阀门无法全开，待用户确认）", "（阀门无法全开）")
                    if value != paragraph.text:
                        _set_visible_text(paragraph, value)

    if not document.tables:
        raise RuntimeError("human review report has no dashboard tables")
    _set_visible_text(
        document.tables[0].cell(0, 0),
        f"BAKER HUGHES | MASONEILAN\n{report_title}",
    )

    count_table = next(
        table for table in document.tables
        if any("诊断阀门" in cell.text for cell in table.rows[0].cells)
        and any("1级" in cell.text for cell in table.rows[0].cells)
    )
    _set_visible_text(count_table.cell(0, 0), f"{len(ratings)}\n诊断阀门")
    for level in range(1, 6):
        _set_visible_text(count_table.cell(0, level), f"{counts[level]}\n{level}级")

    rating_list = next(
        table for table in document.tables
        if len(table.rows[0].cells) >= 3
        and [cell.text.strip() for cell in table.rows[0].cells[:3]] == ["序号", "位号", "评级"]
    )
    listed_tags: set[str] = set()
    for row in rating_list.rows[1:]:
        for offset in range(0, len(row.cells), 3):
            if offset + 2 >= len(row.cells):
                continue
            tag = row.cells[offset + 1].text.strip()
            if tag in ratings:
                _set_visible_text(row.cells[offset + 2], f"{ratings[tag]}级")
                listed_tags.add(tag)
    if listed_tags != set(ratings):
        raise RuntimeError("rating list does not cover every human-reviewed valve")

    judgement = next(
        table for table in document.tables
        if table.rows and all("级 |" in cell.text for cell in table.rows[0].cells)
    )
    descriptions = {
        1: "无明显影响，保留趋势基线",
        2: "轻微缺陷，建议计划复核",
        3: "有缺陷，建议计划复核或维护",
        4: "严重缺陷，建议优先复核和检修",
        5: "无数据或影响未知",
    }
    levels = [level for level in range(1, 6) if counts[level]]
    cell_count = len(judgement.rows[0].cells)
    groups = [[level] for level in levels]
    if len(groups) > cell_count:
        groups = [levels[: len(levels) - cell_count + 1]] + [
            [level] for level in levels[len(levels) - cell_count + 1 :]
        ]
    for index, group in enumerate(groups):
        _set_visible_text(
            judgement.cell(0, index),
            "\n".join(
                f"{level}级 | {counts[level]}台\n{descriptions[level]}"
                for level in group
            ),
        )
    for index in range(len(groups), cell_count):
        _set_visible_text(judgement.cell(0, index), "")

    priority_tags = [tag for tag, rating in ratings.items() if rating == max(ratings.values())]
    first_attention = next(
        table for table in document.tables
        if len(table.rows) == 1 and table.cell(0, 0).text.strip() == "首要关注"
    )
    attention_lines = []
    for tag in priority_tags:
        action = str(overrides[tag].get("action") or "").strip()
        attention_lines.append(f"{tag}：{action}")
    _set_visible_text(first_attention.cell(0, 1), "\n".join(attention_lines))

    conclusion = next(
        table for table in document.tables
        if len(table.rows) == 1 and len(table.rows[0].cells) == 1
        and table.cell(0, 0).text.strip().startswith("总体判断：")
    )
    distribution = "，".join(f"{level}级 {counts[level]}台" for level in range(1, 6))
    _set_visible_text(
        conclusion.cell(0, 0),
        f"总体判断：本批共 {len(ratings)} 台阀门，{distribution}。"
        "4级阀门应列入优先复核与检修清单，3级阀门建议在窗口期复测并结合现场工况确认。",
    )

    priority_table = next(
        table for table in document.tables
        if len(table.rows[0].cells) == 5
        and [cell.text.strip() for cell in table.rows[0].cells[:2]] == ["评级", "位号"]
    )
    priority_rows: set[str] = set()
    for row in priority_table.rows[1:]:
        tag = row.cells[1].text.strip()
        if tag not in overrides:
            continue
        values = overrides[tag]
        _set_visible_text(row.cells[0], f"{ratings[tag]}级")
        _set_visible_text(row.cells[2], "人工校准")
        _set_visible_text(row.cells[3], "\n".join(f"•  {item}" for item in values.get("issues", [])))
        _set_visible_text(row.cells[4], "\n".join(f"•  {item}" for item in values.get("actions", [])))
        priority_rows.add(tag)
    if priority_rows != set(overrides):
        raise RuntimeError("priority action list does not cover every human-reviewed valve")

    document.core_properties.title = report_title
    document.core_properties.subject = "ValVue3 离线诊断报告"
    destination_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination_docx)
    return {
        "valve_count": len(ratings),
        "rating_counts": {str(level): counts[level] for level in range(1, 6)},
        "priority_valves": priority_tags,
    }
