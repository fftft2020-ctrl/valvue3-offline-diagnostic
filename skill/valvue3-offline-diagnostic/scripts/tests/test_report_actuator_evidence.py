"""Report rendering must not invent actuator specifications from valve models."""

from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document

from offline_engine import business_report, generator, pipeline
from offline_engine.config import load_project


ACTUATOR_LABEL = "执行机构型号 / 尺寸"
MISSING_ACTUATOR_VALUES = (
    None,
    "",
    "   ",
    "-",
    "—",
    "–",
    "一",
    "未提供",
    "规格书未提供",
    "规格书栏位为空",
)
# Cover both the observed valve model and an unknown prefix. Neither establishes
# the actuator model, even when the valve model itself is present and valid.
SPECIFICATION_CASES = [
    pytest.param({}, "- / -", id="both-models-absent"),
    *[
        pytest.param({"model_no": model}, "- / -", id=f"{model}-actuator-absent")
        for model in ("88-21124", "99-ABC12")
    ],
    *[
        pytest.param(
            {"model_no": model, "actuator_model": value},
            "- / -",
            id=f"{model}-missing-actuator-{index}",
        )
        for model in ("88-21124", "99-ABC12")
        for index, value in enumerate(MISSING_ACTUATOR_VALUES)
    ],
    pytest.param(
        {"model_no": "88-21124", "actuator_size": "6"},
        "- / 6",
        id="actuator-size-only",
    ),
    *[
        pytest.param(
            {"model_no": model, "actuator_model": "87", "actuator_size": "6"},
            "87 / 6",
            id=f"{model}-explicit-actuator",
        )
        for model in ("88-21124", "99-ABC12")
    ],
]


@pytest.mark.parametrize("spec_fields,expected", SPECIFICATION_CASES)
def test_specification_supplement_uses_only_explicit_actuator_values(
    spec_fields, expected
) -> None:
    original = deepcopy(spec_fields)
    valve = generator.ValveReport(folder="FV-101", spec_fields=spec_fields)

    rows = dict(business_report.specification_supplement_rows(valve))

    assert rows[ACTUATOR_LABEL] == expected
    assert spec_fields == original


@pytest.mark.parametrize("data_only", [True, False], ids=["data", "diagnostic"])
def test_prepared_report_docx_keeps_actuator_evidence_without_model_fallback(
    tmp_path: Path, data_only: bool
) -> None:
    """Exercise the real shared renderer in each mode, without mocking it.

    This is prepared-report rendering coverage, not source extraction, diagnosis
    submission, approval, freezing, or formal PDF delivery acceptance.
    """
    (tmp_path / "project.toml").write_text(
        '[project]\ncustomer="Anonymous"\nreport_title="Actuator Evidence"\n'
        'report_date="2026-10-09"\n[paths]\noutput_docx="report.docx"\n'
        f'[presentation]\ndata_only={str(data_only).lower()}\n',
        encoding="utf-8",
    )
    config = load_project(tmp_path)
    valves = [
        generator.ValveReport(
            folder=f"FV-{index:03d}",
            valve_model=case.values[0].get("model_no", "未提供"),
            spec_fields=deepcopy(case.values[0]),
            rating_conclusion_override="Synthetic prepared diagnosis.",
            overall_assessment="Synthetic prepared assessment.",
        )
        for index, case in enumerate(SPECIFICATION_CASES)
    ]
    originals = [deepcopy(valve.spec_fields) for valve in valves]

    pipeline._render_prepared_report(config, valves)

    document = Document(config.output_docx)
    rendered_values = []
    device_tags = []
    for table in document.tables:
        for row in table.rows:
            cells = row.cells
            for index, cell in enumerate(cells[:-1]):
                if cell.text == ACTUATOR_LABEL:
                    rendered_values.append(cells[index + 1].text)
                elif cell.text == "位号" and len(cells) == 4:
                    device_tags.append(cells[index + 1].text)

    assert device_tags == [valve.folder for valve in valves]
    assert rendered_values == [case.values[1] for case in SPECIFICATION_CASES]
    assert [valve.spec_fields for valve in valves] == originals
    headings = [paragraph.text for paragraph in document.paragraphs]
    assert ("诊断结论" in headings) is not data_only
