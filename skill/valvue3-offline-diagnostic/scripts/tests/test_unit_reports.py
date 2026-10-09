from __future__ import annotations

import json
import hashlib
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from docx import Document
import fitz
import pytest

from offline_engine.unit_reports import assemble_unit_reports, validate_unit_report_pdf


def _source(path: Path, chapters: list[tuple[str, int, str]]) -> None:
    document = Document()
    for tag, rating, color in chapters:
        if document.element.body.xpath("./w:tbl"):
            document.add_page_break()
        header = document.add_table(rows=1, cols=2)
        header.cell(0, 0).text = tag
        header.cell(0, 1).text = f"{rating}级"
        document.add_paragraph(f"{tag} 的人工判断")
        picture = BytesIO()
        Image.new("RGB", (8, 8), color).save(picture, format="PNG")
        picture.seek(0)
        document.add_picture(picture)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def test_assemble_units_preserves_human_chapters_without_other_valve_media(tmp_path: Path) -> None:
    project = tmp_path / "project"
    (project / "work").mkdir(parents=True)
    (project / "experience").mkdir()
    (project / "project.toml").write_text(
        '[report_period_overrides.first]\nonly_valves = ["A", "C"]\n'
        '[report_period_overrides.second]\nonly_valves = ["B"]\n',
        encoding="utf-8",
    )
    (project / "experience/valve_context.json").write_text(
        json.dumps({tag: {"unit": unit} for tag, unit in
                    (("A", "甲装置"), ("B", "甲装置"), ("C", "乙装置"))}, ensure_ascii=False),
        encoding="utf-8",
    )
    _source(project / "experience/feedback_inbox/first.docx", [("A", 1, "red"), ("C", 3, "blue")])
    _source(project / "experience/agent_synced/second.docx", [("B", 2, "green")])
    first = project / "experience/feedback_inbox/first.docx"
    original = Document(first)
    original.sections[0].footer.paragraphs[0].text = "Repeated footer"
    original.sections[0].even_page_footer.paragraphs[0].text = "Repeated footer"
    original.sections[0].first_page_footer.paragraphs[0].text = "Repeated footer"
    original.save(first)
    second = project / "experience/agent_synced/second.docx"
    clean = Document(second)
    clean.sections[0].footer.paragraphs[0].text = "Clean footer"
    clean.save(second)
    selection = project / "work/selection.json"
    selection.write_text(json.dumps({
        "title": "测试报告", "unit_context": "experience/valve_context.json",
        "sources": [
            {"period": "first", "docx": "experience/feedback_inbox/first.docx"},
            {"period": "second", "docx": "experience/agent_synced/second.docx",
             "provenance": {"kind": "agent_synced_from_human_return"}},
        ],
    }, ensure_ascii=False), encoding="utf-8")

    (project / "output/business/按装置分册-v2").mkdir(parents=True)
    manifest_path = assemble_unit_reports(project, selection)
    assert manifest_path.parent.name == "按装置分册-v3"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert (manifest["valve_count"], manifest["unit_count"]) == (3, 2)
    assert manifest["sources"][1]["provenance"]["kind"] == "agent_synced_from_human_return"
    by_unit = {item["unit"]: item for item in manifest["reports"]}
    assert by_unit["甲装置"]["ratings"] == {"1": 1, "2": 1}
    for unit, expected, unwanted in (("甲装置", {"A", "B"}, "C"), ("乙装置", {"C"}, "A")):
        path = project / by_unit[unit]["docx"]
        document = Document(path)
        assert len(document.sections[0]._sectPr.xpath("./w:footerReference")) == 1
        assert document.sections[0].footer.paragraphs[0].text == "Clean footer"
        body = " ".join(document.element.body.xpath(".//w:t/text()"))
        assert all(f"{tag} 的人工判断" in body for tag in expected)
        assert f"{unwanted} 的人工判断" not in body
        with ZipFile(path) as package:
            media = [name for name in package.namelist() if name.startswith("word/media/")]
        assert len(media) == len(expected)


@pytest.mark.parametrize("omitted", [
    "Actuator friction is elevated and needs review.", "必须停机检修", "不要继续运行",
])
def test_unit_pdf_validation_checks_bound_docx_text_and_rating(tmp_path: Path, omitted: str) -> None:
    project = tmp_path / "project"
    folder = project / "output/business/按装置分册-v1"
    folder.mkdir(parents=True)
    docx = folder / "甲装置-阀门离线诊断报告.docx"
    pdf = docx.with_suffix(".pdf")
    document = Document()
    header = document.add_table(rows=1, cols=2)
    header.cell(0, 0).text = "FV-1"
    header.cell(0, 1).text = "2级"
    paragraph = "Actuator friction is elevated and needs review."
    document.add_paragraph(paragraph)
    document.add_paragraph("必须停机检修")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "不要继续运行"
    document.add_paragraph("ＡＢＣ：１２")
    document.sections[0].header.paragraphs[0].text = "Header not part of body"
    document.add_paragraph().add_run("Hidden text").font.hidden = True
    document.save(docx)
    (folder / "分册清单.json").write_text(json.dumps({"reports": [{
        "unit": "甲装置", "docx": str(docx.relative_to(project)),
        "sha256": hashlib.sha256(docx.read_bytes()).hexdigest(),
        "valves": ["FV-1"], "ratings": {"2": 1},
    }]}, ensure_ascii=False), encoding="utf-8")
    rendered = fitz.open()
    page = rendered.new_page()
    page.insert_text((72, 72), "FV-1 2级", fontname="china-s")
    page.insert_text((72, 100), paragraph)
    page.insert_text((72, 130), "必须停机\n检修", fontname="china-s")
    page.insert_text((72, 170), "不要继续运行", fontname="china-s")
    page.insert_text((72, 200), "ABC:12")
    rendered.save(pdf)
    rendered.close()

    assert validate_unit_report_pdf(project, docx, pdf)["ok"] is True
    rendered = fitz.open()
    page = rendered.new_page()
    page.insert_text((72, 72), "FV-1 2级", fontname="china-s")
    for index, text in enumerate([paragraph, "必须停机检修", "不要继续运行", "ABC:12"]):
        if text != omitted:
            page.insert_text((72, 110 + index * 30), text, fontname="china-s")
    rendered.save(pdf, incremental=False)
    rendered.close()
    result = validate_unit_report_pdf(project, docx, pdf)
    assert result["checked_paragraphs"] == 6
    assert "DOCX body paragraphs" in result["failures"][0]
