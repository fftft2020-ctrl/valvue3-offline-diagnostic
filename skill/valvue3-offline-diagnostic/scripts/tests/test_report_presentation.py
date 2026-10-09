import fitz

from docx import Document

from offline_engine import business_report, generator, presentation as presentation_module


def test_effective_spec_content_rect_prefers_data_image_over_report_furniture() -> None:
    class Page:
        rect = fitz.Rect(0, 0, 595, 842)

        def get_text(self, mode):
            assert mode == "dict"
            return {
                "blocks": [
                    {"type": 1, "bbox": (56, 31, 243, 85)},  # report header image
                    {"type": 1, "bbox": (57, 96, 514, 713)},  # specification body
                ]
            }

        def get_drawings(self):
            return []

    rect = generator.effective_spec_content_rect(Page())
    assert rect.x0 == 55
    assert rect.y0 == 94
    assert rect.x1 == 516
    assert rect.y1 == 715


def test_effective_spec_content_rect_keeps_vector_fields_outside_raster_table() -> None:
    class Page:
        rect = fitz.Rect(0, 0, 612, 792)

        def get_text(self, mode):
            assert mode == "dict"
            return {
                "blocks": [
                    {"type": 1, "bbox": (54, 148, 498, 599)},
                    {
                        "type": 0,
                        "bbox": (26, 29, 559, 83),
                        "lines": [{"spans": [{"text": "technical header"}]}],
                    },
                    {
                        "type": 0,
                        "bbox": (21, 733, 406, 753),
                        "lines": [{"spans": [{"text": "technical footer"}]}],
                    },
                ]
            }

        def get_drawings(self):
            return []

    rect = generator.effective_spec_content_rect(Page())

    assert rect.y0 == 27
    assert rect.y1 == 755


def test_effective_spec_content_rect_falls_back_to_body_drawings() -> None:
    class Page:
        rect = fitz.Rect(0, 0, 595, 842)

        def get_text(self, mode):
            return {"blocks": []}

        def get_drawings(self):
            return [
                {"rect": fitz.Rect(20, 30, 200, 60)},  # header
                {"rect": fitz.Rect(55, 100, 500, 720)},  # form body
                {"rect": fitz.Rect(50, 795, 250, 805)},  # footer
            ]

    rect = generator.effective_spec_content_rect(Page())
    assert rect.x0 == 53
    assert rect.y0 == 98
    assert rect.x1 == 502
    assert rect.y1 == 722


def presentation():
    return presentation_module


def test_customer_facing_display_helpers_normalize_review_text() -> None:
    module = presentation()

    assert module.display_manufacturer("Masoneilan / Baker Hughes") == "Masoneilan"
    assert module.display_manufacturer("Masoneilan阀门") == "Masoneilan"
    assert module.display_manufacturer("MASONEILAN（Baker Hughes）") == "Masoneilan"
    assert module.display_positioner_model("SVI II AP (Firmware 3.1)") == "SVI II AP"
    assert module.display_actuator_style("Single Acting") == "单作用"
    assert module.display_actuator_style("双作用") == "双作用"
    assert module.display_air_action("Air To Open") == "气开"
    assert module.display_air_action("气动关闭") == "气关"
    assert module.display_valve_type("Globe（待复核）") == "Globe"
    assert module.display_valve_type("Rotary (to be confirmed)") == "Rotary"
    assert module.display_valve_type("流量控制阀（依据位号推断，待现场复核）") == "流量控制阀"
    assert (
        module.display_valve_type(
            "流量控制阀（依据现场长期运行状态、设备编号、工艺介质、设计条件综合判断，待现场复核）"
        )
        == "流量控制阀"
    )
    assert module.display_valve_type("流量控制阀（依据（位号推断），待现场复核）") == "流量控制阀"
    assert module.display_valve_type("Rotary (inferred from (tag), pending confirmation)") == "Rotary"
    assert module.display_valve_type("Globe (body (待复核)) rotary") == "Globe rotary"
    assert module.display_valve_type("Globe (angle body)") == "Globe (angle body)"


def test_simplify_action_items_splits_and_deduplicates_without_losing_safety_actions() -> None:
    module = presentation()

    actions = module.simplify_action_items(
        [
            "优先安排现场复核和窗口期检修；检查填料、阀杆及执行机构。",
            "建议检查填料、阀杆及执行机构。",
            "处理后复测扩展签名和Ramp，确认回差改善。",
        ]
    )

    assert actions == [
        "优先安排现场复核和窗口期检修。",
        "检查填料、阀杆及执行机构。",
        "处理后复测扩展签名和Ramp，确认回差改善。",
    ]


def test_simplify_action_items_can_bound_display_bullets_without_dropping_tail() -> None:
    module = presentation()

    actions = module.simplify_action_items(
        "优先现场复核；窗口期检修；检查气源；处理后复测。",
        max_items=3,
    )

    assert actions == [
        "优先现场复核。",
        "窗口期检修。",
        "检查气源。；处理后复测。",
    ]


def test_simplify_action_items_merges_related_check_and_retest_bullets() -> None:
    module = presentation()

    actions = module.simplify_action_items(
        [
            "检查气源压力、过滤减压阀和气管泄漏。",
            "检查仪表气源、过滤减压阀容量和压力表。",
            "处理后复测Ramp。",
            "复测阶跃并确认响应改善。",
        ]
    )

    assert len(actions) == 2
    assert "气源压力" in actions[0]
    assert "压力表" in actions[0]
    assert "复测Ramp" in actions[1]
    assert "复测阶跃" in actions[1]


def test_customer_paragraph_hides_internal_curve_path_but_keeps_review_evidence() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()

    paragraph.add_run(presentation_module.sanitize_customer_text(
        "待复核：VALVE-001/anonymous_pattern_step_p04_01.png "
        "尚未完成基于该裁剪曲线的人工图形判读，不能使用通用模板替代。"
    ))

    rendered = paragraph.text
    assert "VALVE-001/anonymous_pattern_step_p04_01.png" not in rendered
    assert "anonymous_pattern_step_p04_01.png" not in rendered
    assert rendered == "待复核：尚未完成基于该裁剪曲线的人工图形判读，不能使用通用模板替代。"


def test_customer_visible_sinks_hide_generic_internal_image_tokens() -> None:
    absolute_path = "/workspace/project/raw/FV-1/source_page.png"
    relative_path = "work/crops/FV-1/foo.png"
    sequencer_image = "Seq_FV-1_other_p4_1.png"
    doc = Document()

    paragraph = doc.add_paragraph()
    paragraph.add_run(presentation_module.sanitize_customer_text(f"待复核：{absolute_path} 曲线显示响应迟滞，需现场复核。"))
    cell = doc.add_table(rows=1, cols=2).cell(0, 1)
    business_report._cell_text(cell, f"{relative_path} 显示压力回环特征。")
    hyperlink_paragraph = doc.add_paragraph()
    business_report.add_internal_hyperlink(
        hyperlink_paragraph,
        f"图形判读：{sequencer_image} 提示执行机构响应需复核。",
        "curve-review",
    )

    visible_text = "\n".join(
        [paragraph.text, doc.tables[0].cell(0, 1).text, "".join(hyperlink_paragraph._p.itertext())]
    )
    for internal_token in (absolute_path, relative_path, sequencer_image):
        assert internal_token not in visible_text
    assert "待复核" in visible_text
    assert "曲线显示响应迟滞，需现场复核。" in visible_text
    assert "显示压力回环特征。" in visible_text
    assert "提示执行机构响应需复核。" in visible_text


def test_size_pressure_keeps_missing_values_explicit() -> None:
    module = presentation()

    assert module.display_size_pressure("2 in", "CL600") == "2 in / CL600"
    assert module.display_size_pressure("", "CL600") == "- / CL600"
    assert module.display_size_pressure("2 in", "") == "2 in / -"


def test_fail_position_only_translates_explicit_open_close_values() -> None:
    module = presentation()

    assert module.display_fail_position("OPEN") == "Fail Open"
    assert module.display_fail_position("close") == "Fail Close"
    assert module.display_fail_position("Closed") == "Fail Close"
    assert module.display_fail_position("air to open") == "-"
    assert module.display_fail_position("") == "-"


def test_basic_information_rows_use_the_keyword_presentation_interface() -> None:
    rows = presentation().basic_information_rows(
        manufacturer="Masoneilan / Baker Hughes",
        positioner_model="SVI II AP (Firmware 3.1)",
        valve_model="21000",
        valve_type="Globe（待复核）",
        spring_range="153.3–306.4 kPa",
        spec_fields={
            "valve_size": "2 in",
            "pressure_class": "CL600",
            "fail_position": "OPEN",
            "rated_cv": "50",
            "flow_action": "flow to open",
            "actuator_model": "88",
        },
    )

    assert rows == [
        ("厂家", "Masoneilan"),
        ("定位器型号", "SVI II AP"),
        ("阀门型号 / 类型", "21000 / Globe"),
        ("阀门尺寸 / 压力等级", "2 in / CL600"),
        ("失效位置", "Fail Open"),
        ("弹簧范围", "153.3–306.4 kPa"),
    ]
    rendered = " ".join(value for _, value in rows)
    assert "Cv" not in rendered
    assert "flow to open" not in rendered
    assert "88" not in rendered


def test_customer_model_spec_hides_cv_but_keeps_model_size_and_pressure() -> None:
    valve = generator.ValveReport(
        folder="FV-1004",
        valve_model="88-21124",
        valve_spec="4 in / ASME CLASS 900 / Rated Cv 49",
        spec_fields={"valve_size": "4 in", "pressure_class": "ASME CLASS 900", "rated_cv": "49"},
    )

    rendered = business_report.valve_model_spec(valve)

    assert rendered == "88-21124 / 4 in / ASME CLASS 900"
    assert "Cv" not in rendered


def test_svi_ii_ap_dtm_fields_are_integrated_into_existing_tables() -> None:
    valve = generator.ValveReport(
        folder="FV-DTM",
        dtm_info={
            "dtm_report_type": "svi_ii_ap",
            "date_time": "2026-07-13 14:05",
            "firmware": "3.2.3",
            "dtm_revision": "3.30.0 BuildID: 092325",
            "actuator_style": "Single Acting",
            "air_action": "Air To Open",
            "strokes": "580",
            "cycles": "94441",
            "time_open": "70201",
            "time_closed": "0",
            "time_near_closed": "1634",
            "device_mode": "Normal",
            "valve_position": "117.20 %",
            "input_signal": "25.24 mA",
            "supply_pressure": "400.6 kPa",
            "min_temperature": "-6.99 degC",
            "max_temperature": "60.55 degC",
            "faults": "Low Power；Actuator Error",
        },
        active_faults=[],
    )

    assert business_report.dtm_basic_information_rows(valve) == [
        ("DTM报告时间", "2026-07-13 14:05"),
        ("固件版本 / DTM版本", "3.2.3 / 3.30.0 BuildID: 092325"),
        ("执行机构形式 / 气路动作", "单作用 / 气开"),
    ]

    rows = dict(business_report.dtm_metric_rows(valve))
    assert "累计行程 / 循环次数" in rows
    assert "全开 / 全关 / 近关时间" in rows
    assert "DTM运行快照" in rows
    assert rows["DTM活动故障"] == "Low Power；Actuator Error"


def test_execution_report_fields_replace_dtm_labels_in_customer_tables() -> None:
    valve = generator.ValveReport(
        folder="FV-EXECUTION",
        positioner_type="SVI2AP",
        positioner_report_info={
            "firmware": "3.2.8 / 58728",
            "diagnostic_software_version": "3.40.0 BuildID: 061126",
            "diagnostics_level": "Advanced Diagnostics",
            "option_config": "Advanced Diagnostics; Single Acting;",
            "actuator_style": "Single Acting",
            "air_action": "ATC",
            "strokes": "1633",
            "cycles": "378432",
            "time_open": "34644",
            "time_closed": "294",
            "time_near_closed": "186",
            "device_mode": "Setup",
            "valve_position": "61.7 %",
            "manual_setpoint": "0.2 %",
            "input_signal": "20.18 mA",
            "supply_pressure": "47.22 psi",
            "min_temperature": "0.3 degC",
            "max_temperature": "55.7 degC",
            "active_faults": ["Bias Out Of Range"],
            "historical_faults": ["Actuator Error", "Position Error"],
            "fault_data_status": "present",
        },
    )

    assert business_report.positioner_report_basic_information_rows(valve) == [
        ("固件版本 / 诊断组件版本", "3.2.8 / 58728 / 3.40.0 BuildID: 061126"),
        ("诊断等级 / 配置", "Advanced Diagnostics / Advanced Diagnostics; Single Acting;"),
        ("执行机构形式 / 气路动作", "单作用 / 气关"),
    ]

    rows = business_report.positioner_report_metric_rows(valve)
    assert [row[0] for row in rows] == [
        "累计行程 / 循环次数",
        "全开 / 全关 / 近关时间",
        "运行快照",
        "当前故障",
        "历史故障",
    ]
    assert all("DTM" not in row[0] for row in rows)


def test_positioner_report_does_not_claim_no_fault_when_source_has_no_fault_fields() -> None:
    valve = generator.ValveReport(
        folder="FV-NO-FAULT-FIELDS",
        positioner_report_info={"strokes": "100", "fault_data_status": "source_missing"},
    )

    rows = dict(business_report.positioner_report_metric_rows(valve))

    assert rows["当前故障"] == "源文件未提供当前故障字段"
    assert rows["历史故障"] == "源文件未提供历史故障字段"


def test_specification_supplement_and_historical_runtime_rows_are_customer_facing() -> None:
    valve = generator.ValveReport(
        folder="FV-HISTORY",
        spec_fields={
            "end_connection": "RF Flanges",
            "rated_cv": "3650",
            "flow_action": "Flow to Open",
            "leakage_class": "IV",
            "actuator_model": "5A",
            "actuator_size": "12",
            "supply_pressure": "2.82 bar g",
            "actuator_travel": "50 deg",
            "positioner_diagnostics_level": "Standard",
        },
        positioner_report_info={
            "historical_runtime": {
                "historical_period": "2024年7月",
                "note": "2024年7月更换定位器，前后计数累计。",
                "historical": {
                    "strokes": 10,
                    "cycles": 20,
                    "time_open": 30,
                    "time_closed": 40,
                    "time_near_closed": 50,
                },
                "total": {
                    "strokes": 110,
                    "cycles": 220,
                    "time_open": 330,
                    "time_closed": 440,
                    "time_near_closed": 550,
                },
            }
        },
    )

    assert business_report.specification_supplement_rows(valve) == [
        ("端连接 / 额定 Cv", "RF Flanges / 3650"),
        ("流向 / 泄漏等级", "Flow to Open / IV"),
        ("执行机构型号 / 尺寸", "5A / 12"),
        ("执行机构供气 / 行程", "2.82 bar g / 50 deg"),
    ]
    rows = dict(business_report.positioner_report_metric_rows(valve))
    assert "2024年7月更换定位器" in rows["历史数据备注"]
    assert "累计：行程 110；循环 220" in rows["历史数据备注"]


def test_specification_supplement_keeps_required_template_rows_when_values_are_missing() -> None:
    valve = generator.ValveReport(folder="FV-MISSING-SPEC")

    assert business_report.specification_supplement_rows(valve) == [
        ("端连接 / 额定 Cv", "- / -"),
        ("流向 / 泄漏等级", "- / -"),
        ("执行机构型号 / 尺寸", "- / -"),
        ("执行机构供气 / 行程", "- / -"),
    ]


def test_specification_supplement_uses_explicit_missing_model_note() -> None:
    valve = generator.ValveReport(
        folder="FV-SPEC-NOTES",
        spec_fields={
            "actuator_model_note": "规格书未提供",
            "actuator_size": "#12 (15)",
        },
    )

    assert dict(business_report.specification_supplement_rows(valve))[
        "执行机构型号 / 尺寸"
    ] == "- / #12 (15)"


def test_specification_supplement_normalizes_leakage_class_and_actuator_prefix() -> None:
    valve = generator.ValveReport(
        folder="FV-COMPOSITE-MODEL",
        spec_fields={
            "model_no": "51-21129",
            "leakage_class": "0.001%/Max.Diff.Pr.",
            "actuator_model_note": "规格书未提供",
            "actuator_size": "#20",
        },
    )

    rows = dict(business_report.specification_supplement_rows(valve))
    assert rows["流向 / 泄漏等级"] == "- / V"
    assert rows["执行机构型号 / 尺寸"] == "51 / #20"


def test_positioner_identity_prefers_accessory_model_and_falls_back_to_device_type() -> None:
    assert business_report.positioner_information_row("SVI 2 AP - 31123111", "SVI2AP") == (
        "定位器型号",
        "SVI 2 AP - 31123111",
    )
    assert business_report.positioner_information_row("未提供", "SVI2AP") == (
        "定位器类型",
        "SVI2AP",
    )
