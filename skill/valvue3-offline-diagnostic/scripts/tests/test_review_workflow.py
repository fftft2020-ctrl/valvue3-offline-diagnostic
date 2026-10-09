import hashlib
import json
from pathlib import Path

import pytest
from docx import Document

from offline_engine.review_workflow import (
    allocate_report_artifact,
    append_run_history,
    compare_review_documents,
    extract_report_overrides,
    extract_review_overrides,
    finalize_human_review_document,
    ingest_review_return,
    latest_review_import,
)


def test_finalize_human_review_synchronizes_dashboard_and_preserves_body(tmp_path: Path) -> None:
    source = tmp_path / "returned.docx"
    output = tmp_path / "formal.docx"
    document = Document()
    document.add_paragraph("报告日期\n2026-08-25")
    document.add_paragraph("人工校验状态：内部评审建议，待用户确认")
    document.add_paragraph("人工曲线判读保留")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "旧封面"
    count = document.add_table(rows=1, cols=6)
    for index, value in enumerate(("1\n诊断阀门", "0\n1级", "1\n2级", "0\n3级", "0\n4级", "0\n5级")):
        count.cell(0, index).text = value
    ratings = document.add_table(rows=2, cols=3)
    for index, value in enumerate(("序号", "位号", "评级")):
        ratings.cell(0, index).text = value
    for index, value in enumerate(("1", "FV-101", "2级")):
        ratings.cell(1, index).text = value
    judgement = document.add_table(rows=1, cols=3)
    for index, value in enumerate(("2级 | 1台", "3级 | 0台", "4级 | 0台")):
        judgement.cell(0, index).text = value
    attention = document.add_table(rows=1, cols=2)
    attention.cell(0, 0).text = "首要关注"
    conclusion = document.add_table(rows=1, cols=1)
    conclusion.cell(0, 0).text = "总体判断：旧结论"
    priority = document.add_table(rows=2, cols=5)
    for index, value in enumerate(("评级", "位号", "校验来源", "主要问题", "建议行动")):
        priority.cell(0, index).text = value
    priority.cell(1, 1).text = "FV-101"
    document.save(source)

    result = finalize_human_review_document(
        source,
        output,
        report_title="正式报告",
        report_date="2026-09-04",
        overrides={
            "FV-101": {
                "rating": 4,
                "issues": ["阀门无法全开。"],
                "actions": ["现场检查。"],
                "action": "现场检查。",
            }
        },
    )

    finalized = Document(output)
    assert result["rating_counts"] == {"1": 0, "2": 0, "3": 0, "4": 1, "5": 0}
    assert finalized.tables[0].cell(0, 0).text.endswith("正式报告")
    assert finalized.tables[1].cell(0, 4).text == "1\n4级"
    assert finalized.tables[2].cell(1, 2).text == "4级"
    assert finalized.tables[4].cell(0, 1).text == "FV-101：现场检查。"
    assert finalized.tables[6].cell(1, 0).text == "4级"
    assert finalized.tables[6].cell(1, 3).text == "•  阀门无法全开。"
    assert all("待用户确认" not in paragraph.text for paragraph in finalized.paragraphs)
    assert any(paragraph.text == "人工曲线判读保留" for paragraph in finalized.paragraphs)


def test_report_versions_are_project_local_and_never_overwrite(tmp_path: Path) -> None:
    existing = tmp_path / "output/review/Demo-人工校验版-v1.docx"
    existing.parent.mkdir(parents=True)
    existing.write_bytes(b"existing")
    (tmp_path / "work/review/v3").mkdir(parents=True)

    review = allocate_report_artifact(tmp_path, "review", "Demo")
    business = allocate_report_artifact(tmp_path, "business", "Demo")

    assert review.version == 4
    assert review.docx == tmp_path / "output/review/Demo-待人工校验版-v4.docx"
    assert review.pdf == tmp_path / "output/review/Demo-待人工校验版-v4.pdf"
    assert review.work_dir == tmp_path / "work/review/v4"
    assert business.version == 1
    assert business.docx == tmp_path / "output/business/Demo-最终商务版-v1.docx"
    assert existing.read_bytes() == b"existing"


def test_report_filename_includes_configured_period_label(tmp_path: Path) -> None:
    (tmp_path / "work/review/v16").mkdir(parents=True)

    artifact = allocate_report_artifact(
        tmp_path,
        "review",
        "Demo",
        period_label="2026年7月18日至7月21日",
    )

    assert artifact.version == 17
    assert artifact.docx.name == "Demo-2026年7月18日至7月21日-待人工校验版-v17.docx"
    assert artifact.pdf.name == "Demo-2026年7月18日至7月21日-待人工校验版-v17.pdf"


def test_review_versions_count_both_names_without_work_history(tmp_path: Path) -> None:
    output = tmp_path / "output/review"
    output.mkdir(parents=True)
    old = output / "Demo-往期-人工校验版-v8.pdf"
    new = output / "Demo-待人工校验版-v9.docx"
    old.write_bytes(b"old pdf")
    new.write_bytes(b"current docx")
    artifact = allocate_report_artifact(tmp_path, "review", "Demo")
    assert artifact.version == 10
    assert artifact.docx.name == "Demo-待人工校验版-v10.docx"
    assert old.read_bytes() == b"old pdf"
    assert new.read_bytes() == b"current docx"


def test_data_report_versions_are_independent_for_each_period(tmp_path: Path) -> None:
    first = allocate_report_artifact(
        tmp_path,
        "data",
        "Demo",
        period_label="2026年7月18日至7月21日",
    )
    first.work_dir.mkdir(parents=True)
    second = allocate_report_artifact(
        tmp_path,
        "data",
        "Demo",
        period_label="2026年7月22日至今",
    )

    assert first.version == 1
    assert second.version == 1
    assert first.docx == (
        tmp_path
        / "output/data/Demo-2026年7月18日至7月21日-纯数据版-v1.docx"
    )
    assert second.docx == (
        tmp_path
        / "output/data/Demo-2026年7月22日至今-纯数据版-v1.docx"
    )
    assert first.work_dir == (
        tmp_path / "work/data/2026年7月18日至7月21日/v1"
    )
    assert second.work_dir == (
        tmp_path / "work/data/2026年7月22日至今/v1"
    )


def test_run_history_is_append_only_json_lines(tmp_path: Path) -> None:
    first = append_run_history(tmp_path, {"event": "review_generated", "version": 1})
    second = append_run_history(tmp_path, {"event": "business_generated", "version": 1})

    assert first == second == tmp_path / "work/run_history.jsonl"
    records = [json.loads(line) for line in first.read_text(encoding="utf-8").splitlines()]
    assert [record["event"] for record in records] == ["review_generated", "business_generated"]
    assert all(record["created_at"] for record in records)


def _write_review(path: Path, *, rating: str, conclusion: str, action: str) -> None:
    document = Document()
    document.add_heading("1. FV-101", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "健康状态评级"
    table.cell(0, 1).text = rating
    table.cell(1, 0).text = "评级结论"
    table.cell(1, 1).text = conclusion
    document.add_heading("5. 行动项", level=3)
    document.add_paragraph(action, style="List Bullet")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def _write_review_run_manifest(project: Path, source: Path) -> None:
    manifest = project / "work/review/v1/run_manifest.json"
    manifest.parent.mkdir(parents=True, exist_ok=True)
    manifest.write_text(
        json.dumps(
            {
                "outputs": {
                    "docx": str(source),
                    "docx_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                },
                "formal_diagnosis": {
                    "tree_sha256": "a" * 64,
                    "project_manifest_sha256": "b" * 64,
                },
            }
        ),
        encoding="utf-8",
    )


def test_review_comparison_ignores_formatting_and_extracts_human_content(tmp_path: Path) -> None:
    original = tmp_path / "original.docx"
    returned = tmp_path / "returned.docx"
    _write_review(original, rating="", conclusion="", action="")
    _write_review(returned, rating="3级", conclusion="存在迟滞，建议计划复核。", action="窗口期复测。")

    changes = compare_review_documents(original, returned)
    overrides = extract_review_overrides(returned, valve_tags={"FV-101"})

    assert any(change["after"] == "3级" for change in changes)
    assert overrides == {
        "FV-101": {
            "rating": 3,
            "rating_conclusion": "存在迟滞，建议计划复核。",
            "actions": ["窗口期复测。"],
            "action": "窗口期复测。",
        }
    }


@pytest.mark.parametrize("label", ["人工校验版", "待人工校验版"])
def test_ingest_review_return_keeps_original_and_writes_project_change_list(tmp_path: Path, label: str) -> None:
    source = tmp_path / f"output/review/Demo-{label}-v1.docx"
    returned = tmp_path / "incoming/人工填写.docx"
    _write_review(source, rating="", conclusion="", action="")
    _write_review(returned, rating="2级", conclusion="轻微迟滞，建议持续观察。", action="下次窗口期复测。")
    _write_review_run_manifest(tmp_path, source)

    imported = ingest_review_return(tmp_path, returned, valve_tags={"FV-101"})

    assert imported.version == 1
    assert imported.returned_docx == tmp_path / "experience/feedback_inbox/人工填写.docx"
    assert imported.returned_docx.is_file()
    assert imported.change_json.parent.parent == tmp_path / "work/review/returns"
    payload = json.loads(imported.change_json.read_text(encoding="utf-8"))
    assert payload["overrides"]["FV-101"]["rating"] == 2
    assert payload["overrides"]["FV-101"]["action"] == "下次窗口期复测。"
    assert payload["formal_diagnosis"] == {
        "tree_sha256": "a" * 64,
        "project_manifest_sha256": "b" * 64,
    }
    assert payload["unresolved_changes"] == []
    assert latest_review_import(tmp_path) == imported.change_json


def test_ingest_marks_unmapped_human_content_for_confirmation(tmp_path: Path) -> None:
    source = tmp_path / "output/review/Demo-人工校验版-v1.docx"
    returned = tmp_path / "incoming/人工填写.docx"
    _write_review(source, rating="", conclusion="", action="")
    _write_review(returned, rating="2级", conclusion="轻微迟滞。", action="复测。")
    _write_review_run_manifest(tmp_path, source)
    document = Document(returned)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "供气压力"
    table.cell(0, 1).text = "420 kPa（人工修改）"
    document.save(returned)

    imported = ingest_review_return(tmp_path, returned, valve_tags={"FV-101"})
    payload = json.loads(imported.change_json.read_text(encoding="utf-8"))

    assert any(change["after"] == "420 kPa（人工修改）" for change in payload["unresolved_changes"])


def test_ingest_preserves_explicitly_cleared_required_valve_fields(
    tmp_path: Path,
) -> None:
    source = tmp_path / "output/review/Demo-人工校验版-v1.docx"
    returned = tmp_path / "incoming/人工清空.docx"
    _write_review(
        source,
        rating="2级",
        conclusion="原评级结论。",
        action="原建议行动。",
    )
    _write_review(returned, rating="2级", conclusion="", action="")
    _write_review_run_manifest(tmp_path, source)

    imported = ingest_review_return(
        tmp_path, returned, valve_tags={"FV-101"}
    )
    payload = json.loads(imported.change_json.read_text(encoding="utf-8"))

    assert payload["overrides"]["FV-101"]["rating_conclusion"] == ""
    assert payload["overrides"]["FV-101"]["actions"] == []
    assert payload["overrides"]["FV-101"]["action"] == ""
    assert payload["unresolved_changes"] == []


def test_ingest_preserves_explicitly_cleared_issue_list(tmp_path: Path) -> None:
    source = tmp_path / "output/review/Demo-人工校验版-v1.docx"
    returned = tmp_path / "incoming/人工清空问题.docx"

    def write(path: Path, issue: str) -> None:
        document = Document()
        title = document.add_table(rows=1, cols=2)
        title.cell(0, 0).text = "FV-101\nDEVICE-1"
        title.cell(0, 1).text = "2级\n健康状态"
        document.add_paragraph("诊断结论")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "原评级结论。"
        document.add_paragraph("主要问题与建议行动")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "主要问题" + (f"\n•  {issue}" if issue else "")
        table.cell(0, 1).text = "建议行动\n•  原建议。"
        path.parent.mkdir(parents=True, exist_ok=True)
        document.save(path)

    write(source, "原问题。")
    write(returned, "")
    _write_review_run_manifest(tmp_path, source)
    imported = ingest_review_return(
        tmp_path, returned, valve_tags={"FV-101"}
    )
    payload = json.loads(imported.change_json.read_text(encoding="utf-8"))

    assert payload["overrides"]["FV-101"]["issues"] == []
    assert payload["unresolved_changes"] == []


def test_ingest_preserves_explicitly_cleared_report_fields(tmp_path: Path) -> None:
    source = tmp_path / "output/review/Demo-人工校验版-v1.docx"
    returned = tmp_path / "incoming/人工清空总体.docx"

    def write(path: Path, conclusion: str, recommendation: str) -> None:
        document = Document()
        document.add_heading("2. 总体结论与检修建议", level=1)
        document.add_paragraph(conclusion)
        document.add_heading("后续复测、维护和检修建议", level=2)
        document.add_paragraph(recommendation)
        document.add_heading("3. 逐台阀门诊断章节", level=1)
        path.parent.mkdir(parents=True, exist_ok=True)
        document.save(path)

    write(source, "原总体结论。", "原总体建议。")
    write(returned, "", "")
    _write_review_run_manifest(tmp_path, source)
    imported = ingest_review_return(tmp_path, returned, valve_tags=set())
    payload = json.loads(imported.change_json.read_text(encoding="utf-8"))

    assert payload["report_overrides"] == {
        "overall_conclusion": "",
        "recommendations": [],
    }
    assert payload["unresolved_changes"] == []


def test_ingest_rejects_review_source_changed_after_run_manifest(
    tmp_path: Path,
) -> None:
    source = tmp_path / "output/review/Demo-人工校验版-v1.docx"
    returned = tmp_path / "incoming/人工填写.docx"
    _write_review(source, rating="2级", conclusion="原结论。", action="原建议。")
    _write_review(returned, rating="2级", conclusion="人工结论。", action="人工建议。")
    _write_review_run_manifest(tmp_path, source)
    document = Document(source)
    document.add_paragraph("母版被改写。")
    document.save(source)

    with pytest.raises(RuntimeError, match="does not match its run manifest"):
        ingest_review_return(tmp_path, returned, valve_tags={"FV-101"})


def test_extract_report_level_human_conclusion_and_recommendations(tmp_path: Path) -> None:
    path = tmp_path / "returned.docx"
    document = Document()
    document.add_heading("2. 总体结论与检修建议", level=1)
    document.add_paragraph("本次共分析 2 台控制阀。")
    document.add_paragraph("人工总体结论原文。")
    document.add_heading("后续复测、维护和检修建议", level=2)
    document.add_paragraph("人工建议一。", style="List Bullet")
    document.add_paragraph("人工建议二。", style="List Bullet")
    document.save(path)

    assert extract_report_overrides(path) == {
        "overall_conclusion": "人工总体结论原文。",
        "recommendations": ["人工建议一。", "人工建议二。"],
    }


def test_extract_report_overrides_ignores_matching_toc_entry(tmp_path: Path) -> None:
    path = tmp_path / "returned-with-toc.docx"
    document = Document()
    document.add_heading("目录", level=1)
    document.add_paragraph("2. 总体结论与检修建议")
    document.add_paragraph("3. 逐台阀门诊断章节")
    document.add_paragraph("目录说明不属于总体结论。")
    document.add_heading("1. 离线诊断阀门健康状态评级表", level=1)
    document.add_heading("2. 总体结论与检修建议", level=1)
    document.add_paragraph("人工总体结论原文。")
    document.add_heading("后续复测、维护和检修建议", level=2)
    document.add_paragraph("人工建议一。", style="List Bullet")
    document.add_heading("3. 逐台阀门诊断章节", level=1)
    document.save(path)

    assert extract_report_overrides(path) == {
        "overall_conclusion": "人工总体结论原文。",
        "recommendations": ["人工建议一。"],
    }


def test_extracts_human_fields_from_business_dashboard_review_tables(tmp_path: Path) -> None:
    path = tmp_path / "business-dashboard-returned.docx"
    document = Document()
    document.add_paragraph("总体结论与检修建议")
    overall = document.add_table(rows=1, cols=1)
    overall.cell(0, 0).text = "人工总体结论原文。"
    document.add_paragraph("后续复测、维护和检修建议")
    document.add_paragraph("人工建议一。")
    document.add_paragraph("人工建议二。")

    title = document.add_table(rows=1, cols=2)
    title.cell(0, 0).text = "FV-101\n88-21124"
    title.cell(0, 1).text = "3级\n健康状态"
    document.add_paragraph("诊断结论")
    conclusion = document.add_table(rows=1, cols=1)
    conclusion.cell(0, 0).text = "存在迟滞，建议计划复核。"
    document.add_paragraph("主要问题与建议行动")
    issue_action = document.add_table(rows=1, cols=2)
    issue_action.cell(0, 0).text = "主要问题\n•  开阀方向存在迟滞。"
    issue_action.cell(0, 1).text = "建议行动\n•  窗口期复测。"
    document.save(path)

    assert extract_report_overrides(path) == {
        "overall_conclusion": "人工总体结论原文。",
        "recommendations": ["人工建议一。", "人工建议二。"],
    }
    assert extract_review_overrides(path, valve_tags={"FV-101"}) == {
        "FV-101": {
            "rating": 3,
            "rating_conclusion": "存在迟滞，建议计划复核。",
            "issues": ["开阀方向存在迟滞。"],
            "actions": ["窗口期复测。"],
            "action": "窗口期复测。",
        }
    }


@pytest.mark.parametrize(
    "test_status",
    ["未发现非成功任务", "•  Run Step Test：失败 [Failed]；  Run Ramp Test：已取消 [Cancelled]"],
)
def test_dashboard_key_value_conclusion_is_not_overwritten_by_test_status(
    tmp_path: Path, test_status: str,
) -> None:
    source = tmp_path / "output/review/Demo-人工校验版-v1.docx"
    returned = tmp_path / "incoming/returned.docx"
    _write_review(source, rating="2级", conclusion="原结论。", action="原行动。")
    _write_review_run_manifest(tmp_path, source)
    document = Document()
    for tag, rating, conclusion in (
        ("FV-101", "2级", "存在迟滞，安排复测。"),
        ("FV-102", "1级", "未见需要处理的性能异常。"),
    ):
        title = document.add_table(rows=1, cols=2)
        title.cell(0, 0).text = tag
        title.cell(0, 1).text = f"{rating}\n健康状态"
        document.add_paragraph("诊断结论")
        keyed = document.add_table(rows=1, cols=2)
        keyed.cell(0, 0).text = "评级结论"
        keyed.cell(0, 1).text = conclusion
        status = document.add_table(rows=1, cols=1)
        status.cell(0, 0).text = test_status
    returned.parent.mkdir(parents=True, exist_ok=True)
    document.save(returned)

    expected = {
        "FV-101": {"rating": 2, "rating_conclusion": "存在迟滞，安排复测。"},
        "FV-102": {"rating": 1, "rating_conclusion": "未见需要处理的性能异常。"},
    }
    assert extract_review_overrides(returned, valve_tags=set(expected)) == expected
    imported = ingest_review_return(tmp_path, returned, valve_tags=set(expected), review_base=source)
    saved = json.loads(imported.overrides_json.read_text(encoding="utf-8"))
    assert saved == expected
    assert json.loads(imported.change_json.read_text(encoding="utf-8"))["overrides"] == expected


def test_explicit_base_multiple_returns_are_immutable_and_ambiguous(tmp_path: Path) -> None:
    source = tmp_path / "output/review/Demo-人工校验版-v1.docx"
    returned = tmp_path / "incoming/returned.docx"
    _write_review(source, rating="", conclusion="", action="")
    _write_review_run_manifest(tmp_path, source)
    _write_review(returned, rating="2级", conclusion="轻微迟滞。", action="复测。")
    first = ingest_review_return(tmp_path, returned, valve_tags={"FV-101"}, review_base=source)
    original = first.change_json.read_bytes()
    repeat = ingest_review_return(tmp_path, returned, valve_tags={"FV-101"}, review_base=source)
    assert repeat == first
    _write_review(returned, rating="3级", conclusion="迟滞影响定位。", action="检查。")
    second = ingest_review_return(tmp_path, returned, valve_tags={"FV-101"}, review_base=source)
    assert first.change_json != second.change_json
    assert first.change_json.read_bytes() == original
    with pytest.raises(RuntimeError, match="unique"):
        latest_review_import(tmp_path)
    assert latest_review_import(tmp_path, review_return=second.change_json) == second.change_json
    with pytest.raises(RuntimeError, match="current formal"):
        ingest_review_return(tmp_path, returned, valve_tags=set(), review_base=source,
                             expected_binding={"tree_sha256": "wrong"})
    _write_review(source, rating="4级", conclusion="changed", action="changed")
    with pytest.raises(RuntimeError, match="does not match"):
        ingest_review_return(tmp_path, returned, valve_tags=set(), review_base=source)


def test_explicit_format_base_and_unmapped_package_changes(tmp_path: Path) -> None:
    import shutil
    import zipfile
    from offline_engine.review_workflow import review_base_binding
    from offline_engine.format_revision import format_revision_manifest_path
    source = tmp_path / "output/review/Demo-人工校验版-v1.docx"
    _write_review(source, rating="2级", conclusion="轻微迟滞。", action="复测。")
    _write_review_run_manifest(tmp_path, source)
    child = tmp_path / "output/business/Demo-v2.docx"
    child.parent.mkdir()
    shutil.copy2(source, child)
    manifest = format_revision_manifest_path(tmp_path, child)
    manifest.parent.mkdir(parents=True)
    manifest.write_text(json.dumps({"parent_docx": str(source), "candidate_docx": str(child),
        "parent_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "candidate_sha256": hashlib.sha256(child.read_bytes()).hexdigest()}))
    assert review_base_binding(tmp_path, child) == review_base_binding(tmp_path, source)
    returned = tmp_path / "returned.docx"
    document = Document(child)
    document.sections[0].header.paragraphs[0].text = "人工新增说明"
    document.save(returned)
    with zipfile.ZipFile(returned, "a") as package:
        package.writestr("word/embeddings/added.bin", b"changed object")
    result = ingest_review_return(tmp_path, returned, valve_tags={"FV-101"}, review_base=child)
    data = json.loads(result.change_json.read_text())
    assert data["package_changes"] == ["embedded_sha256"]
    assert data["structure_changes"] == ["headers_footers_notes"]
    assert data["unresolved_changes"]
    document = Document(child)
    document.add_paragraph("unregistered content drift")
    document.save(child)
    with pytest.raises(RuntimeError, match="content changed"):
        review_base_binding(tmp_path, child)
