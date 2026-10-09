import importlib.util
import base64
import json
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from offline_engine import business_report, generator
from offline_engine.curve_identity import spec_insertion_marker


VALIDATOR_PATH = Path(__file__).resolve().parents[1] / "validate_valvue3_report.py"
spec = importlib.util.spec_from_file_location("report_validator", VALIDATOR_PATH)
assert spec and spec.loader
validator = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = validator
spec.loader.exec_module(validator)


def _docx(path: Path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(path)


def _png(path: Path) -> None:
    path.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    ))


def _valid_project_toml(
    project: Path,
    *,
    formal_diagnosis_run: str | None = None,
    presentation: str = "",
    features: str = "",
) -> None:
    project.mkdir(parents=True, exist_ok=True)
    formal_line = (
        f'formal_diagnosis_run = "{formal_diagnosis_run}"\n'
        if formal_diagnosis_run
        else ""
    )
    project.joinpath("project.toml").write_text(
        "\n".join(
            [
                "[project]",
                'customer = "Demo"',
                'report_title = "ValVue3 Offline Diagnostic Report"',
                'report_date = "2026-08-28"',
                "",
                "[paths]",
                'raw = "raw"',
                'specs = "specs"',
                'inbound = "inbound"',
                'work = "work"',
                'output = "output"',
                'output_docx = "report.docx"',
                'output_pdf = "report.pdf"',
                formal_line.rstrip("\n"),
                features,
                presentation,
            ]
        )
        + "\n",
        encoding="utf-8",
    )


def _complete_contract_evidence() -> dict[str, object]:
    return {
        "positioner": {
            "option_config": "Standard Diagnostics",
            "device_mode": "Normal",
            "valve_position": 10.0,
            "manual_setpoint": 0.0,
            "input_signal": 4.0,
            "board_temperature": 31.2,
            "min_temperature": -3.9,
            "max_temperature": 48.5,
        },
        "runtime": {
            "strokes": 10,
            "cycles": 20,
            "time_open": 30,
            "time_closed": 40,
            "time_near_closed": 5,
        },
        "file": {
            "report_start_time": "2026-08-28 09:00:00",
            "report_end_time": "2026-08-28 10:00:00",
            "report_time": "2026-08-28 09:00:00 - 10:00:00",
        },
        "status": {
            "active_faults": [],
            "historical_faults": [],
            "task_execution": [
                {
                    "task_name": "Run Ramp",
                    "start_time": "2026-08-28 09:00:00",
                    "end_time": "2026-08-28 09:01:00",
                    "result": "Succeeded",
                }
            ],
        },
        "metrics": {
            "run_ramp": {
                metric: {bound: float(index) for bound in ("max", "average", "min")}
                for index, metric in enumerate(("hysteresis", "dead_band", "linearity"), 1)
            }
        },
    }


def _nested_material_evidence() -> tuple[dict[str, object], dict[str, object]]:
    """Build the current frozen-bundle shape without customer identifiers."""
    values: dict[str, object] = {
        "positioner.option_config": "Standard Diagnostics",
        "positioner.device_mode": "Normal",
        "positioner.valve_position": 0,
        "positioner.manual_setpoint": 12.5,
        "positioner.input_signal": 4.0,
        "positioner.board_temperature": 31.2,
        "positioner.min_temperature": -3.9,
        "positioner.max_temperature": 48.5,
        "runtime.strokes": 0,
        "runtime.cycles": 20,
        "runtime.time_open": 30,
        "runtime.time_closed": 40,
        "runtime.time_near_closed": 5,
        "file.report_start_time": "2026-08-28 09:00:00",
        "file.report_end_time": "2026-08-28 10:00:00",
        "file.report_time": "2026-08-28 09:00:00 - 10:00:00",
        "status.active_faults": [],
        "status.historical_faults": [],
        "status.task_execution": [
            {
                "task_name": "Run Ramp",
                "start_time": "2026-08-28 09:00:00",
                "end_time": "2026-08-28 09:01:00",
                "result": "Succeeded",
            }
        ],
    }
    for metric in ("hysteresis", "dead_band", "linearity"):
        for statistic in ("max", "average", "min"):
            values[f"metrics.run_ramp.{metric}.{statistic}"] = 0
    evidence: dict[str, object] = {
        "material_extraction": {
            "fields": {
                path: {"selected": {"normalized_value": value}}
                for path, value in values.items()
            }
        },
        # This intentionally contains a conflicting legacy view. The current
        # material-extraction selection must win over compatibility fallbacks.
        "source_record": {
            "positioner_report_info": {
                "positioner": {"device_mode": "Wrong legacy value"}
            }
        },
    }
    return evidence, values


def _report_page(
    *,
    source_sha256: str,
    page: int,
    roles: list[str],
) -> dict[str, object]:
    return {
        "source": "specs/FV-1/spec.pdf",
        "source_sha256": source_sha256,
        "page": page,
        "roles": roles,
        "image_path": f"work/spec-page-{page}.png",
        "match_basis": ["directory_tag", "page_label"],
    }


def _marked_picture(document: Document, image: Path, *markers: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run().add_picture(str(image))
    for marker in markers:
        business_report.add_bookmark(paragraph, marker)


def test_validator_rejects_unreviewed_curves_and_internal_review_metadata(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    _docx(
        docx,
        "图形判读：待复核。 confidence=high experience_rule_ids=rule-a "
        "/internal/work/crop.png Codex 中间评分",
    )
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps({"valves": [{"folder": "FV-1", "rating": 3}]}), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx), "--summary", str(summary), "--json",
    ])

    assert validator.main() == 1


def test_validator_rejects_zero_spec_matches_when_project_has_spec_pdfs(
    tmp_path: Path, monkeypatch
) -> None:
    project = tmp_path / "project"
    (project / "specs/FV-1").mkdir(parents=True)
    (project / "specs/FV-1/spec.pdf").write_bytes(b"candidate")
    docx = tmp_path / "report.docx"
    _docx(docx, "Report")
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(
        json.dumps([{"folder": "FV-1", "rating": 3, "spec_page": None}]),
        encoding="utf-8",
    )
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.main() == 1


def test_validator_rejects_expected_spec_page_without_actual_docx_insertion_marker(
    tmp_path: Path, monkeypatch
) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    _docx(docx, "Report")
    marker = spec_insertion_marker("FV-1", "spec")
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "spec_page": {"source": "FV-1/spec.pdf", "page": 1},
        "spec_insertion_marker": marker,
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.main() == 1


def test_validator_accepts_matching_actual_spec_insertion_marker(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    marker = spec_insertion_marker("FV-1", "spec")
    docx = tmp_path / "report.docx"
    document = Document()
    image_path = tmp_path / "spec.png"
    _png(image_path)
    paragraph = document.add_paragraph()
    paragraph.add_run().add_picture(str(image_path))
    business_report.add_bookmark(paragraph, marker)
    document.save(docx)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "spec_page": {"source": "FV-1/spec.pdf", "page": 1},
        "spec_insertion_marker": marker,
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.main() == 0


def test_validator_rejects_spec_marker_without_picture_entity(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    marker = spec_insertion_marker("FV-1", "spec")
    docx = tmp_path / "report.docx"
    document = Document()
    paragraph = document.add_paragraph("marker only")
    business_report.add_bookmark(paragraph, marker)
    document.save(docx)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "spec_page": {"source": "FV-1/spec.pdf", "page": 1},
        "spec_insertion_marker": marker,
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.main() == 1


def test_validator_rejects_blip_relationship_targeting_non_media_part(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    marker = spec_insertion_marker("FV-1", "spec")
    original = tmp_path / "original.docx"
    image_path = tmp_path / "spec.png"
    _png(image_path)
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run().add_picture(str(image_path))
    business_report.add_bookmark(paragraph, marker)
    document.save(original)
    docx = tmp_path / "report.docx"
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(docx, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                content = content.replace(b'Target="media/', b'Target="')
                content = content.replace(b'.png"', b'document.xml"')
            target.writestr(item, content)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "spec_page": {"source": "FV-1/spec.pdf", "page": 1},
        "spec_insertion_marker": marker,
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.main() == 1


def test_validator_rejects_spec_page_declaration_with_missing_marker_field(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    _docx(docx, "Report")
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "spec_page": {"source": "FV-1/spec.pdf", "page": 1},
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.main() == 1


def test_finalized_validator_requires_each_curve_marker_to_match_a_valid_summary_interpretation(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    _docx(docx, "图形判读：曲线响应平稳。")
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "images": [{"review_marker": "curve-review-a", "interpretation_validated": True}],
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx), "--summary", str(summary),
        "--require-finalized-curve-review", "--json",
    ])

    assert validator.main() == 1


def test_finalized_validator_accepts_actual_docx_curve_marker_matching_valid_interpretation(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    document = Document()
    paragraph = document.add_paragraph("图形判读：曲线响应平稳。")
    business_report.add_bookmark(paragraph, "curve-review-a")
    document.save(docx)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "images": [{"review_marker": "curve-review-a", "interpretation_validated": True}],
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx), "--summary", str(summary),
        "--require-finalized-curve-review", "--json",
    ])

    assert validator.main() == 0


def test_finalized_validator_accepts_human_confirmed_pending_review_text_without_display_label(
    tmp_path: Path, monkeypatch,
) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    document = Document()
    paragraph = document.add_paragraph("待复核：该性能签名图保留原始曲线，暂不作确定判断。")
    business_report.add_bookmark(paragraph, "curve-review-a")
    document.save(docx)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "images": [{"review_marker": "curve-review-a", "interpretation_validated": True}],
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx), "--summary", str(summary),
        "--require-finalized-curve-review", "--json",
    ])

    assert validator.main() == 0


@pytest.mark.parametrize("legacy_curve_text", ["本图评级依据：", "图形判读为", "建议行动："])
def test_finalized_validator_rejects_legacy_single_curve_output_text(
    tmp_path: Path, monkeypatch, legacy_curve_text: str,
) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("图形判读：客户可见说明。")
    paragraph = document.add_paragraph(f"{legacy_curve_text}旧版单图输出。")
    business_report.add_bookmark(paragraph, "curve-review-a")
    document.save(docx)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "images": [{
            "review_marker": "curve-review-a",
            "interpretation_validated": True,
            "visual_rating": 3,
        }],
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx), "--summary", str(summary),
        "--require-finalized-curve-review", "--json",
    ])

    assert validator.main() == 1


def test_finalized_validator_allows_comprehensive_diagnosis_text(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    document = Document()
    paragraph = document.add_paragraph("综合诊断：客户可见说明。图形判读：客户可见说明。")
    business_report.add_bookmark(paragraph, "curve-review-a")
    document.save(docx)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "images": [{
            "review_marker": "curve-review-a",
            "interpretation_validated": True,
            "visual_rating": 3,
        }],
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx), "--summary", str(summary),
        "--require-finalized-curve-review", "--json",
    ])

    assert validator.main() == 0


def test_finalized_validator_rejects_standalone_diagnosis_label(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    document = Document()
    paragraph = document.add_paragraph("诊断：")
    business_report.add_bookmark(paragraph, "curve-review-a")
    document.save(docx)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "images": [{
            "review_marker": "curve-review-a",
            "interpretation_validated": True,
            "visual_rating": 3,
        }],
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx), "--summary", str(summary),
        "--require-finalized-curve-review", "--json",
    ])

    assert validator.main() == 1


def test_finalized_validator_allows_customer_text_while_summary_keeps_visual_rating(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    document = Document()
    paragraph = document.add_paragraph("图形判读：客户可见说明。")
    business_report.add_bookmark(paragraph, "curve-review-a")
    document.save(docx)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "images": [{
            "review_marker": "curve-review-a",
            "interpretation_validated": True,
            "visual_rating": 3,
            "customer_text": "客户可见说明。",
        }],
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx), "--summary", str(summary),
        "--require-finalized-curve-review", "--json",
    ])

    assert validator.main() == 0


@pytest.mark.parametrize("legacy_curve_text", ["本图评级依据：", "图形判读为", "建议行动："])
def test_finalized_validator_rejects_legacy_single_curve_output_text_in_pdf(
    tmp_path: Path, monkeypatch, legacy_curve_text: str,
) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    document = Document()
    paragraph = document.add_paragraph("图形判读：客户可见说明。")
    business_report.add_bookmark(paragraph, "curve-review-a")
    document.save(docx)
    pdf = tmp_path / "report.pdf"
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "images": [{
            "review_marker": "curve-review-a",
            "interpretation_validated": True,
            "visual_rating": 3,
        }],
    }]), encoding="utf-8")
    monkeypatch.setattr(validator, "pdf_text", lambda _: (True, f"图形判读：客户可见说明。\n{legacy_curve_text}旧版单图输出。", ""))
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx), "--pdf", str(pdf),
        "--summary", str(summary), "--require-finalized-curve-review", "--json",
    ])

    assert validator.main() == 1


def test_finalized_validator_rejects_standalone_diagnosis_label_in_pdf(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    document = Document()
    paragraph = document.add_paragraph("图形判读：客户可见说明。")
    business_report.add_bookmark(paragraph, "curve-review-a")
    document.save(docx)
    pdf = tmp_path / "report.pdf"
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "images": [{
            "review_marker": "curve-review-a",
            "interpretation_validated": True,
            "visual_rating": 3,
        }],
    }]), encoding="utf-8")
    monkeypatch.setattr(validator, "pdf_text", lambda _: (True, "图形判读：客户可见说明。\n诊断：\n", ""))
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx), "--pdf", str(pdf),
        "--summary", str(summary), "--require-finalized-curve-review", "--json",
    ])

    assert validator.main() == 1


@pytest.mark.parametrize("output_kind", ["docx", "pdf"])
@pytest.mark.parametrize("legacy_curve_text", [
    "图形判读（客观观察）：旧版单图输出。",
    "客观观察：旧版单图输出。",
    "诊断：旧版单图输出。",
])
def test_finalized_validator_rejects_legacy_single_curve_prefixes_with_body(
    tmp_path: Path, monkeypatch, output_kind: str, legacy_curve_text: str,
) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("图形判读：客户可见说明。")
    paragraph = document.add_paragraph(legacy_curve_text if output_kind == "docx" else "客户可见说明。")
    business_report.add_bookmark(paragraph, "curve-review-a")
    document.save(docx)
    pdf = tmp_path / "report.pdf"
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "images": [{
            "review_marker": "curve-review-a",
            "interpretation_validated": True,
            "visual_rating": 3,
        }],
    }]), encoding="utf-8")
    if output_kind == "pdf":
        monkeypatch.setattr(validator, "pdf_text", lambda _: (True, f"图形判读：客户可见说明。\n{legacy_curve_text}", ""))
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        *( ["--pdf", str(pdf)] if output_kind == "pdf" else [] ),
        "--summary", str(summary), "--require-finalized-curve-review", "--json",
    ])

    assert validator.main() == 1


@pytest.mark.parametrize("output_kind", ["docx", "pdf"])
@pytest.mark.parametrize("allowed_text", ["综合诊断：客户可见说明。", "诊断数据摘要：客户可见说明。"])
def test_finalized_validator_allows_nonlegacy_diagnosis_prefixes(
    tmp_path: Path, monkeypatch, output_kind: str, allowed_text: str,
) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("图形判读：客户可见说明。")
    paragraph = document.add_paragraph(allowed_text)
    business_report.add_bookmark(paragraph, "curve-review-a")
    document.save(docx)
    pdf = tmp_path / "report.pdf"
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "images": [{
            "review_marker": "curve-review-a",
            "interpretation_validated": True,
            "visual_rating": 3,
        }],
    }]), encoding="utf-8")
    if output_kind == "pdf":
        monkeypatch.setattr(validator, "pdf_text", lambda _: (True, f"图形判读：客户可见说明。\n{allowed_text}", ""))
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        *( ["--pdf", str(pdf)] if output_kind == "pdf" else [] ),
        "--summary", str(summary), "--require-finalized-curve-review", "--json",
    ])

    assert validator.main() == 0


def test_finalized_validator_allows_diagnostic_data_summary_in_docx_and_pdf(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    document = Document()
    paragraph = document.add_paragraph("诊断数据摘要：客户可见说明。图形判读：客户可见说明。")
    business_report.add_bookmark(paragraph, "curve-review-a")
    document.save(docx)
    pdf = tmp_path / "report.pdf"
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "images": [{
            "review_marker": "curve-review-a",
            "interpretation_validated": True,
            "visual_rating": 3,
            "customer_text": "客户可见说明。",
        }],
    }]), encoding="utf-8")
    monkeypatch.setattr(validator, "pdf_text", lambda _: (True, "诊断数据摘要：客户可见说明。图形判读：客户可见说明。", ""))
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx), "--pdf", str(pdf),
        "--summary", str(summary), "--require-finalized-curve-review", "--json",
    ])

    assert validator.main() == 0


def test_finalized_validator_rejects_empty_pdf_text(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    document = Document()
    paragraph = document.add_paragraph("图形判读：客户可见说明。")
    business_report.add_bookmark(paragraph, "curve-review-a")
    document.save(docx)
    pdf = tmp_path / "report.pdf"
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "images": [{
            "review_marker": "curve-review-a",
            "interpretation_validated": True,
            "visual_rating": 3,
        }],
    }]), encoding="utf-8")
    monkeypatch.setattr(validator, "pdf_text", lambda _: (True, "", ""))
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx), "--pdf", str(pdf),
        "--summary", str(summary), "--require-finalized-curve-review", "--json",
    ])

    assert validator.main() == 1


def test_validator_rejects_general_image_paths_and_internal_score_wording(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    _docx(docx, "图形判读：crop_01.PNG；内部综合评分为3级；data_rating=2。")
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx), "--summary", str(summary), "--json",
    ])

    assert validator.main() == 1


def test_validator_allows_customer_action_text_with_a_slash_when_not_a_path(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    _docx(docx, "图形判读：建议计划复核/维护，并在窗口期确认异常重复性。")
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx), "--summary", str(summary), "--json",
    ])

    assert validator.main() == 0


def test_validator_accepts_specification_and_air_piping_on_independent_physical_pages(
    tmp_path: Path, monkeypatch
) -> None:
    project = tmp_path / "project"
    project.mkdir()
    (project / "project.toml").write_text(
        "[features]\ninclude_accessory_pages = false\n", encoding="utf-8"
    )
    first = tmp_path / "spec.png"
    second = tmp_path / "air.png"
    _png(first)
    second.write_bytes(first.read_bytes() + b"distinct-package-part")
    document = Document()
    _marked_picture(document, first, validator.deterministic_spec_marker("FV-1", "specification"))
    _marked_picture(document, second, validator.deterministic_spec_marker("FV-1", "air_piping"))
    docx = tmp_path / "report.docx"
    document.save(docx)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "report_pages": [
            _report_page(source_sha256="a" * 64, page=2, roles=["specification"]),
            _report_page(source_sha256="a" * 64, page=5, roles=["air_piping"]),
        ],
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.main() == 0


def test_validator_accepts_one_media_with_two_role_bookmarks_for_same_physical_page(
    tmp_path: Path, monkeypatch
) -> None:
    project = tmp_path / "project"
    project.mkdir()
    image = tmp_path / "combined.png"
    _png(image)
    document = Document()
    _marked_picture(
        document,
        image,
        validator.deterministic_spec_marker("FV-1", "specification"),
        validator.deterministic_spec_marker("FV-1", "air_piping"),
    )
    docx = tmp_path / "report.docx"
    document.save(docx)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "report_pages": [
            _report_page(
                source_sha256="b" * 64,
                page=2,
                roles=["specification", "air_piping"],
            ),
        ],
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.main() == 0


def test_validator_rejects_two_media_items_for_same_physical_page_roles(
    tmp_path: Path, monkeypatch
) -> None:
    project = tmp_path / "project"
    project.mkdir()
    first = tmp_path / "spec.png"
    second = tmp_path / "air.png"
    _png(first)
    second.write_bytes(first.read_bytes() + b"distinct-package-part")
    document = Document()
    _marked_picture(document, first, validator.deterministic_spec_marker("FV-1", "specification"))
    _marked_picture(document, second, validator.deterministic_spec_marker("FV-1", "air_piping"))
    docx = tmp_path / "report.docx"
    document.save(docx)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "report_pages": [
            _report_page(
                source_sha256="c" * 64,
                page=2,
                roles=["specification", "air_piping"],
            ),
        ],
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.main() == 1


def test_validator_rejects_accessory_role_and_disabled_overview_sections_by_default(
    tmp_path: Path, monkeypatch
) -> None:
    project = tmp_path / "project"
    project.mkdir()
    image = tmp_path / "accessory.png"
    _png(image)
    document = Document()
    document.add_paragraph("现场信息清单")
    _marked_picture(document, image, validator.deterministic_spec_marker("FV-1", "accessory"))
    docx = tmp_path / "report.docx"
    document.save(docx)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "report_pages": [
            _report_page(source_sha256="d" * 64, page=3, roles=["accessory"]),
        ],
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.main() == 1


def test_explicit_presentation_switch_reenables_legacy_overview_section(
    tmp_path: Path,
) -> None:
    project = tmp_path / "project"
    _valid_project_toml(
        project,
        presentation="[presentation]\ninclude_overview_judgement = true",
    )

    assert "整体判断" not in validator.disabled_overview_sections(project)
    assert "首要关注" in validator.disabled_overview_sections(project)


def test_validator_rejects_incomplete_structured_evidence(tmp_path: Path, monkeypatch) -> None:
    project = tmp_path / "project"
    project.mkdir()
    docx = tmp_path / "report.docx"
    _docx(docx, "运行与定位器诊断信息")
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "structured_evidence": {"runtime": {"strokes": 10}},
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.main() == 1


def test_validator_requires_structured_evidence_for_configured_frozen_run(
    tmp_path: Path, monkeypatch
) -> None:
    project = tmp_path / "project"
    _valid_project_toml(
        project,
        formal_diagnosis_run="work/formal_diagnosis/runs/v3",
    )
    docx = tmp_path / "report.docx"
    _docx(docx, "运行与定位器诊断信息")
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{"folder": "FV-1"}]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.main() == 1


def test_contract_evidence_view_reads_frozen_source_record_shape() -> None:
    view = validator._contract_evidence_view({
        "source_record": {
            "positioner_report_info": {
                "positioner": {"device_mode": "Normal"},
                "runtime": {"cycles": {"value": 42, "unit": "counts"}},
                "file": {"report_start_time": "2026-08-26 09:00:00"},
                "status": {"historical_faults": []},
            },
            "metrics": {"run_ramp": {"hysteresis": {"max": 1.2}}},
        }
    })

    assert view["positioner"]["device_mode"] == "Normal"
    assert view["runtime"]["cycles"]["value"] == 42
    assert view["metrics"]["run_ramp"]["hysteresis"]["max"] == 1.2


def test_structured_evidence_validator_reads_current_nested_material_extraction_shape(
    tmp_path: Path,
) -> None:
    evidence, _values = _nested_material_evidence()
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(
        json.dumps([{"folder": "FV-1", "structured_evidence": evidence}]),
        encoding="utf-8",
    )

    checked, failures = validator.structured_evidence_failures(summary)

    assert checked == 1
    assert failures == []


def test_report_field_manifest_validator_uses_same_nested_evidence_contract(
    tmp_path: Path,
) -> None:
    evidence, values = _nested_material_evidence()
    manifest = [
        {
            "path": path,
            "evidence_value": values[path],
            "display_text": f"visible-{path}",
        }
        for path in validator.REQUIRED_STRUCTURED_EVIDENCE_PATHS
    ]
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(
        json.dumps([{
            "folder": "FV-1",
            "structured_evidence": evidence,
            "report_field_manifest": manifest,
        }]),
        encoding="utf-8",
    )
    docx_body = "\n".join(item["display_text"] for item in manifest)

    checked, failures = validator.report_field_manifest_failures(summary, docx_body)

    assert checked == 1
    assert failures == []


def test_report_field_manifest_rejects_value_different_from_nested_evidence(
    tmp_path: Path,
) -> None:
    evidence, values = _nested_material_evidence()
    manifest = [
        {
            "path": path,
            "evidence_value": values[path],
            "display_text": f"visible-{path}",
        }
        for path in validator.REQUIRED_STRUCTURED_EVIDENCE_PATHS
    ]
    manifest[2]["evidence_value"] = 99
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(
        json.dumps([{
            "folder": "FV-1",
            "structured_evidence": evidence,
            "report_field_manifest": manifest,
        }]),
        encoding="utf-8",
    )
    docx_body = "\n".join(item["display_text"] for item in manifest)

    _checked, failures = validator.report_field_manifest_failures(summary, docx_body)

    assert any("positioner.valve_position differs" in failure for failure in failures)


def test_validator_accepts_generator_specification_role_marker_from_summary(
    tmp_path: Path, monkeypatch
) -> None:
    project = tmp_path / "project"
    project.mkdir()
    image = tmp_path / "spec.png"
    _png(image)
    document = Document()
    marker = spec_insertion_marker("FV-1", "specification")
    _marked_picture(document, image, marker)
    docx = tmp_path / "report.docx"
    document.save(docx)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{
        "folder": "FV-1",
        "report_pages": [
            _report_page(source_sha256="f" * 64, page=2, roles=["specification"]),
        ],
    }]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.main() == 0


def test_summary_resolution_prefers_explicit_then_current_work_then_legacy(
    tmp_path: Path,
) -> None:
    project = tmp_path / "project"
    _valid_project_toml(project)
    current = project / "work/analysis_summary.json"
    legacy = project / "work/shared_engine/analysis_summary.json"
    explicit = tmp_path / "explicit.json"
    current.parent.mkdir(parents=True)
    current.write_text("[]", encoding="utf-8")
    legacy.parent.mkdir(parents=True)
    legacy.write_text("[]", encoding="utf-8")

    assert validator.resolve_summary_path(project) == current.resolve()
    assert validator.resolve_summary_path(project, explicit) == explicit.resolve()
    current.unlink()
    assert validator.resolve_summary_path(project) == legacy.resolve()


def test_summary_resolution_uses_report_period_work_override(
    tmp_path: Path,
) -> None:
    project = tmp_path / "project"
    _valid_project_toml(project)
    with (project / "project.toml").open("a", encoding="utf-8") as stream:
        stream.write(
            "\n[report_periods]\n"
            'active = "base"\n'
            'base_label = "base"\n'
            'special_label = "special"\n'
            "\n[report_period_overrides.special]\n"
            'work = "work/special"\n'
        )
    current = project / "work/special/analysis_summary.json"
    current.parent.mkdir(parents=True)
    current.write_text("[]", encoding="utf-8")

    assert validator.resolve_summary_path(project, report_period="special") == current.resolve()


def test_validator_applies_old_accessory_only_config_migration_to_air_piping(
    tmp_path: Path, monkeypatch
) -> None:
    project = tmp_path / "project"
    _valid_project_toml(
        project,
        features="[features]\ninclude_accessory_pages = true",
    )
    image = tmp_path / "air.png"
    _png(image)
    document = Document()
    _marked_picture(document, image, validator.deterministic_spec_marker("FV-1", "air_piping"))
    docx = tmp_path / "report.docx"
    document.save(docx)
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(
        json.dumps([{
            "folder": "FV-1",
            "report_pages": [_report_page(source_sha256="e" * 64, page=2, roles=["air_piping"])],
        }]),
        encoding="utf-8",
    )
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.project_report_page_features(project) == {
        "specification": True,
        "air_piping": False,
        "accessory": True,
    }
    assert validator.main() == 1


def test_formal_validator_requires_report_field_manifest(
    tmp_path: Path, monkeypatch
) -> None:
    project = tmp_path / "project"
    _valid_project_toml(
        project,
        formal_diagnosis_run="work/formal_diagnosis/runs/v3",
    )
    docx = tmp_path / "report.docx"
    _docx(docx, "图形判读：客户可见说明。")
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(
        json.dumps([{
            "folder": "FV-1",
            "structured_evidence": _complete_contract_evidence(),
        }]),
        encoding="utf-8",
    )
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.main() == 1


def test_nonformal_data_only_report_does_not_require_formal_bindings(
    tmp_path: Path, monkeypatch
) -> None:
    project = tmp_path / "project"
    _valid_project_toml(
        project,
        formal_diagnosis_run="work/formal_diagnosis/runs/v3",
        presentation="[presentation]\ndata_only = true",
    )
    docx = tmp_path / "report.docx"
    _docx(docx, "图形判读")
    summary = tmp_path / "analysis_summary.json"
    summary.write_text(json.dumps([{"folder": "FV-1"}]), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", [
        "validate_valvue3_report.py", "--project", str(project), "--docx", str(docx),
        "--summary", str(summary), "--min-figure-interpretations", "0", "--json",
    ])

    assert validator.project_uses_formal_diagnosis(project) is False
    assert validator.main() == 0


def _required_information_fixture(
    tmp_path: Path,
    *,
    omitted_display_paths: set[str] | None = None,
) -> tuple[Path, Path, Path]:
    project = tmp_path / "information-project"
    _valid_project_toml(
        project,
        formal_diagnosis_run="work/formal_diagnosis/runs/v1",
    )
    for name in ("raw", "specs", "inbound", "work"):
        (project / name).mkdir(parents=True, exist_ok=True)
    evidence = _complete_contract_evidence()
    manifest = []
    display_by_path: dict[str, str] = {}
    for index, path in enumerate(validator.REQUIRED_STRUCTURED_EVIDENCE_PATHS, 1):
        present, value = validator.resolve_report_evidence_value(evidence, path)
        assert present
        display = f"客户字段值-{index:02d}"
        display_by_path[path] = display
        manifest.append(
            {
                "path": path,
                "evidence_value": value,
                "display_text": display,
            }
        )
    summary = project / "work/analysis_summary.json"
    summary.write_text(
        json.dumps(
            [
                {
                    "folder": "FV-101",
                    "included_in_report": True,
                    "rating": 2,
                    "issues": ["客户可见问题。"],
                    "actions": ["客户可见建议。"],
                    "images": [],
                    "report_pages": [],
                    "structured_evidence": evidence,
                    "report_field_manifest": manifest,
                }
            ],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    omitted_display_paths = omitted_display_paths or set()
    document = Document()
    document.add_paragraph("总体结论与检修建议")
    document.add_paragraph("FV-101")
    document.add_paragraph("客户可见问题。")
    document.add_paragraph("客户可见建议。")
    for path, display in display_by_path.items():
        if path not in omitted_display_paths:
            document.add_paragraph(display)
    docx = project / "report.docx"
    document.save(docx)
    return project, docx, summary


def test_information_only_accepts_complete_required_customer_information(
    tmp_path: Path,
) -> None:
    project, docx, summary = _required_information_fixture(tmp_path)

    result = validator.check_required_information(project, docx, summary)

    assert result["ok"] is True
    assert result["status"] == "information_complete"
    assert result["issues"] == []
    assert result["counts"] == {
        "valves": {"expected": 1, "confirmed": 1},
        "fields": {"expected": 28, "confirmed": 28},
        "curves": {"expected": 0, "confirmed": 0},
        "technical_pages": {"expected": 0, "confirmed": 0},
    }


@pytest.mark.parametrize("omitted", ["客户可见问题。", "客户可见建议。", "不要继续运行", "2级", "客户字段值-01"])
def test_information_pdf_preserves_short_body_and_table_text(tmp_path: Path, omitted: str) -> None:
    import fitz

    project, docx, summary = _required_information_fixture(tmp_path)
    document = Document(docx)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "2级"
    # Human final wording that is not in the original summary is still required.
    table.cell(0, 1).text = "不要继续运行"
    document.add_paragraph("ＡＢＣ：１２")
    document.save(docx)
    parts = ["".join(p.xpath(".//w:t/text()")) for p in document.element.body.xpath(".//w:p")]
    pdf = docx.with_suffix(".pdf")

    def render(skip: str = "") -> None:
        with fitz.open() as output:
            page = output.new_page()
            for index, part in enumerate(parts):
                if part != skip:
                    text = part.replace("ＡＢＣ：１２", "ABC:12").replace("不要继续运行", "不要继续\n运行")
                    page.insert_text((40, 30 + index * 20), text, fontname="china-s", fontsize=9)
            output.save(pdf)

    render()
    assert validator.check_required_information(project, docx, summary, pdf=pdf)["ok"]
    render(omitted)
    result = validator.check_required_information(project, docx, summary, pdf=pdf)
    assert not result["ok"]
    assert any(item["category"] == "pdf" and omitted in item["items"] for item in result["issue_details"])


def test_information_acceptance_uses_rendered_project_terminology(tmp_path: Path) -> None:
    from dataclasses import replace
    from offline_engine.business_report import _apply_term_replacements

    project, docx, summary = _required_information_fixture(tmp_path)
    records = json.loads(summary.read_text(encoding="utf-8"))
    raw_issue = "Ramp 与扩展签名需要检查。"
    records[0]["issues"] = [raw_issue]
    summary.write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")
    document = Document(docx)
    for paragraph in document.paragraphs:
        if paragraph.text == "客户可见问题。":
            paragraph.text = raw_issue
    replacements = {"Ramp": "斜坡"}
    _apply_term_replacements(document, replacements)
    document.save(docx)
    config = validator.load_project(project)
    config = replace(config, presentation={**config.presentation, "term_replacements": replacements})
    assert validator.check_required_information(project, docx, summary, config=config)["ok"]

    for paragraph in document.paragraphs:
        if "斜坡" in paragraph.text:
            paragraph.text = "缺少该诊断内容"
    document.save(docx)
    assert not validator.check_required_information(project, docx, summary, config=config)["ok"]


def test_information_only_accepts_declared_missing_field_when_placeholder_is_visible(
    tmp_path: Path,
) -> None:
    project, docx, summary = _required_information_fixture(tmp_path)
    payload = json.loads(summary.read_text(encoding="utf-8"))
    path = "metrics.run_ramp.hysteresis.max"
    evidence = payload[0]["structured_evidence"]
    evidence["metrics"]["run_ramp"]["hysteresis"].pop("max")
    evidence["material_extraction"] = {"fields": {}}
    evidence["material_extraction"]["fields"][path] = {
        "selected": None,
        "status": "missing",
        "candidates": [],
    }
    entry = next(
        item
        for item in payload[0]["report_field_manifest"]
        if item["path"] == path
    )
    entry["evidence_value"] = None
    entry["display_text"] = "未提供"
    summary.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    document = Document(docx)
    document.add_paragraph("未提供")
    document.save(docx)

    result = validator.check_required_information(project, docx, summary)

    assert result["ok"] is True
    assert result["counts"]["fields"] == {"expected": 28, "confirmed": 28}


def test_information_only_ignores_disabled_accessory_page_declaration(
    tmp_path: Path,
) -> None:
    project, docx, summary = _required_information_fixture(tmp_path)
    with (project / "project.toml").open("a", encoding="utf-8") as stream:
        stream.write("\n[features]\ninclude_accessory_pages = false\n")
    payload = json.loads(summary.read_text(encoding="utf-8"))
    payload[0]["report_pages"] = [
        _report_page(
            source_sha256="f" * 64,
            page=4,
            roles=["accessory"],
        )
    ]
    summary.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    result = validator.check_required_information(project, docx, summary)

    assert result["ok"] is True
    assert result["counts"]["technical_pages"] == {
        "expected": 0,
        "confirmed": 0,
    }


def test_information_only_compares_diagnosis_after_customer_term_normalization(
    tmp_path: Path,
) -> None:
    project, docx, summary = _required_information_fixture(tmp_path)
    payload = json.loads(summary.read_text(encoding="utf-8"))
    payload[0]["issues"] = ["排查Extended Actuator Signature失败原因。"]
    payload[0]["actions"] = ["处理后复测扩展签名。"]
    summary.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    document = Document(docx)
    document.add_paragraph("排查执行机构性能测试失败原因。")
    document.add_paragraph("处理后复测执行机构性能测试。")
    document.save(docx)

    result = validator.check_required_information(project, docx, summary)

    assert result["ok"] is True
    assert result["issues"] == []


def test_information_only_groups_missing_fields_into_one_compact_issue(
    tmp_path: Path,
) -> None:
    missing = {
        "positioner.option_config",
        "metrics.run_ramp.linearity.min",
    }
    project, docx, summary = _required_information_fixture(
        tmp_path,
        omitted_display_paths=missing,
    )

    result = validator.check_required_information(project, docx, summary)

    assert result["ok"] is False
    assert result["status"] == "information_missing"
    field_issues = [
        item
        for item in result["issue_details"]
        if item["category"] == "required_fields"
    ]
    assert len(field_issues) == 1
    assert field_issues[0]["valve"] == "FV-101"
    assert field_issues[0]["items"] == [
        "positioner.option_config",
        "metrics.run_ramp.linearity.min",
    ]
    assert len([item for item in result["issues"] if item.startswith("FV-101：")]) == 1
    assert result["summary_issues"] == [
        "1台阀门：所需字段未完整进入报告，共2项"
    ]
    assert result["counts"]["fields"] == {"expected": 28, "confirmed": 26}
