import base64
import importlib
import random
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from offline_engine import business_report, generator
from offline_engine.curve_identity import curve_review_marker, spec_insertion_marker
from offline_engine.visual_review import CurveInterpretation


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _prepared_report_text(valves):
    return {
        "rating_conclusions": {
            valve.folder: (
                valve.rating_conclusion_override or "测试冻结评级结论。"
            )
            for valve in valves
        },
        "overall_conclusion": "\n".join(
            valve.overall_assessment for valve in valves if valve.overall_assessment
        ) or "测试冻结综合结论。",
        "recommendations": list(
            dict.fromkeys(action for valve in valves for action in valve.actions)
        ) or ["测试冻结建议。"],
    }


def test_business_dashboard_optimizes_embedded_png_without_changing_dimensions(
    tmp_path: Path,
) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    chart = tmp_path / "rgb-gradient.png"
    image = Image.frombytes(
        "RGB",
        (640, 480),
        random.Random(42).randbytes(640 * 480 * 3),
    )
    image.save(chart)
    valve = generator.ValveReport(
        folder="FV-SLIM",
        rating=2,
        final_rating=2,
        images=[{"path": str(chart), "test_key": "run_step", "page": 1}],
    )
    output = tmp_path / "slim.docx"

    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Valve Diagnostic Project",
        report_name="Masoneilan ValVue3 Offline Diagnostic Report",
        report_date="",
        unit_name="Valve Diagnostic Project",
    )

    with ZipFile(output) as archive:
        media = [
            name
            for name in archive.namelist()
            if name.startswith("word/media/") and name.endswith(".png")
        ]
        assert len(media) == 1
        embedded = archive.read(media[0])
    with Image.open(BytesIO(embedded)) as optimized:
        assert optimized.size == (640, 480)
        assert optimized.mode == "P"
    assert len(embedded) < chart.stat().st_size * 0.60


def test_business_dashboard_report_keeps_complete_evidence_and_spec_after_curves(tmp_path: Path) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    chart = tmp_path / "Seq_20260711_FV-101_INTERNAL_run_step.png"
    spec = tmp_path / "raw-customer-spec-scan.png"
    accessory = tmp_path / "raw-customer-accessory-scan.png"
    chart.write_bytes(PNG_1X1)
    spec.write_bytes(PNG_1X1)
    accessory.write_bytes(PNG_1X1)
    review = CurveInterpretation(
        image_key="FV-101-pattern-step",
        test_key="pattern_step",
        observation="阶跃图显示开阀方向滞后。",
        diagnosis="存在轻微响应滞后。",
        rating=2,
        action="复测并检查填料。",
        customer_text="开阀方向存在轻微滞后，建议复测并检查填料。",
        confidence="high",
        experience_rule_ids=(),
    )
    valve = generator.ValveReport(
        folder="FV-101",
        device_id="1001",
        serial_number="S-1001",
        detection_date="2026-07-11 10:30",
        manufacturer="Masoneilan",
        positioner_model="SVI II AP",
        valve_model="88-21124",
        valve_type="控制阀",
        valve_spec="4 in / ASME CLASS 900",
        spec_fields={"valve_size": "4 in", "pressure_class": "ASME CLASS 900", "fail_position": "CLOSE"},
        spring_range="150–300 kPa",
        dtm_info={
            "dtm_report_type": "svi_ii_ap",
            "date_time": "2026-07-11 10:35",
            "firmware": "3.2.3",
            "dtm_revision": "3.30.0",
            "actuator_style": "Single Acting",
            "air_action": "ATO",
            "strokes": "401",
            "cycles": "37605",
            "time_open": "32886",
            "time_closed": "37362",
            "time_near_closed": "0",
            "device_mode": "Normal",
            "valve_position": "110.70 %",
            "input_signal": "25.04 mA",
            "supply_pressure": "381.9 kPa",
            "min_temperature": "-6.91 degC",
            "max_temperature": "43.86 degC",
            "faults": "Actuator Error",
        },
        metrics={
            "Average Error": 1.0,
            "T86": 6.44,
            "T63": 2.99,
            "Td": 0.3,
            "Overshoot": 0.31,
            "Maximum Hysteresis + Dead Band": 1.87,
            "Average Hysteresis + Dead Band": 1.62,
            "Max Friction": None,
            "Mean Friction": None,
            "Supply Pressure": {"min": 53.93, "max": 54.9, "avg": 54.31},
        },
        rating=3,
        final_rating=3,
        rating_conclusion_override="阀门可完成基本动作，存在轻微滞后，建议复核。",
        main_issue="开阀方向存在滞后。",
        action="复测并检查填料。",
        issues=["开阀方向存在滞后。"],
        actions=["复测并检查填料。"],
        spec_page=generator.SpecPage(valve="FV-101", tag="FV-101", source="spec.pdf", page=1),
        spec_image=str(spec),
        accessory_page=generator.SpecPage(
            valve="FV-101",
            tag="FV-101",
            source="spec.pdf",
            page=2,
        ),
        accessory_image=str(accessory),
        images=[{
            "path": str(chart),
            "test_key": "pattern_step",
            "page": 4,
            "curve_review": review,
        }],
    )
    output = tmp_path / "business.docx"

    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Demo Customer",
        report_name="ValVue3 Offline Diagnostic Report",
        report_date="2026-07-11",
        report_period_label="2026年7月18日至7月21日",
        unit_name="Demo Unit",
        include_field_context_overview=True,
        include_overview_judgement=True,
        include_overview_priority=True,
        include_retest_recommendations=True,
        include_accessory_pages=True,
    )

    rendered = Document(output)
    section = rendered.sections[0]
    assert 8.26 <= section.page_width.inches <= 8.28
    assert 11.68 <= section.page_height.inches <= 11.70
    assert section.top_margin.inches <= 0.43
    assert section.bottom_margin.inches <= 0.43
    assert section.left_margin.inches <= 0.51
    assert section.right_margin.inches <= 0.51
    visible = "\n".join(
        [paragraph.text for paragraph in rendered.paragraphs]
        + [cell.text for table in rendered.tables for row in table.rows for cell in row.cells]
    )
    for required in (
        "诊断总览",
        "总体结论与检修建议",
        "诊断周期",
        "2026年7月18日至7月21日",
        "设备 ID / 序列号",
        "2026-07-11 10:35",
        "运行与定位器诊断信息",
        "401 / 37605",
        "32886 h / 37362 h / 0 h",
        "Actuator Error",
        "开阀方向存在轻微滞后，建议复测并检查填料。",
        "技术资料附件",
        "FV-101 | Accessory",
    ):
        assert required in visible
    assert "诊断报告第" not in visible

    body_text = ["".join(node.xpath(".//w:t/text()")) for node in rendered._element.body.iterchildren()]
    assert body_text.index("FV-101 | 曲线证据与图形判读") < body_text.index("FV-101 | Specification")
    assert body_text.index("FV-101 | Specification") < body_text.index("FV-101 | Accessory")
    curve_heading = next(
        paragraph for paragraph in rendered.paragraphs
        if paragraph.text == "FV-101 | 曲线证据与图形判读"
    )
    assert curve_heading.paragraph_format.page_break_before
    assert curve_heading.paragraph_format.keep_with_next

    document_xml = rendered._element.xml
    for source_name in (chart.name, spec.name, accessory.name):
        assert source_name not in document_xml
    for customer_name in (
        "diagnostic-curve-FV-101-01.png",
        "technical-page-FV-101-001.png",
        "technical-page-FV-101-002.png",
    ):
        assert customer_name in document_xml
    for color in ("004B46", "2E8B57", "C79200", "D96C0B", "C9342C", "6B7280"):
        assert color in document_xml
    assert curve_review_marker(review.image_key) in document_xml
    assert spec_insertion_marker(valve.folder, "specification") in document_xml
    assert spec_insertion_marker(valve.folder, "accessory") in document_xml

    def table_text(table) -> str:
        return "\n".join(cell.text for row in table.rows for cell in row.cells)

    def top_margin(cell) -> int:
        margins = cell._tc.get_or_add_tcPr().find(qn("w:tcMar"))
        assert margins is not None
        top = margins.find(qn("w:top"))
        assert top is not None
        return int(top.get(qn("w:w")))

    curve_table = next(table for table in rendered.tables if table._tbl.xpath(".//w:drawing"))
    curve_grid_widths = [
        int(column.get(qn("w:w"))) for column in curve_table._tbl.tblGrid.gridCol_lst
    ]
    assert curve_grid_widths[0] > curve_grid_widths[1]
    picture_extent = curve_table.cell(0, 0)._tc.xpath(".//wp:extent")[0]
    assert int(picture_extent.get("cx")) >= int(4.0 * 914400)

    rating_cards = next(table for table in rendered.tables if "诊断阀门" in table_text(table))
    overview_matrix = next(
        table for table in rendered.tables
        if len(table.columns) == 9 and "序号" in table.rows[0].cells[0].text
    )
    priority_table = next(
        table for table in rendered.tables
        if len(table.columns) == 5
        and [cell.text for cell in table.rows[0].cells]
        == ["评级", "位号", "校验来源", "主要问题", "建议行动"]
    )
    assert top_margin(rating_cards.cell(0, 0)) <= 160
    assert top_margin(overview_matrix.cell(1, 0)) <= 150
    assert top_margin(priority_table.cell(1, 1)) <= 120
    assert 0.43 <= overview_matrix.rows[1].height.inches <= 0.46
    assert 0.41 <= priority_table.rows[1].height.inches <= 0.44
    assert all(
        row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
        for row in priority_table.rows
    )
    assert sum(paragraph.text == "\u00a0" for paragraph in rendered.paragraphs) == 0

    bookmark_names = rendered._element.body.xpath(".//w:bookmarkStart/@w:name")
    hyperlink_anchors = rendered._element.body.xpath(".//w:hyperlink/@w:anchor")
    assert bookmark_names.count("valve_FV-101") == 1
    assert hyperlink_anchors.count("valve_FV-101") >= 3

    key_metrics = next(
        table for table in rendered.tables
        if "平均误差" in table_text(table)
        and "供气压力" in table_text(table)
    )
    assert len(key_metrics.columns) == 14
    assert key_metrics.rows[0].cells[0].text == "阶跃测试"
    assert key_metrics.rows[0].cells[5].text == "斜坡测试"
    assert key_metrics.rows[0].cells[10].text == "执行机构性能测试"
    assert all(row._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None for row in key_metrics.rows[:2])
    key_metric_xml_text = "\n".join(key_metrics._tbl.xpath(".//w:t/text()"))
    for title in ("阶跃测试", "斜坡测试", "执行机构性能测试"):
        assert key_metric_xml_text.count(title) == 1
    for friction_label in ("最大摩擦力", "平均摩擦力", "最小摩擦力"):
        assert key_metric_xml_text.count(friction_label) == 1
    layout = key_metrics._tbl.tblPr.find(qn("w:tblLayout"))
    assert layout is not None
    assert layout.get(qn("w:type")) == "fixed"
    grid_widths = [int(column.get(qn("w:w"))) for column in key_metrics._tbl.tblGrid.gridCol_lst]
    assert grid_widths == [
        1008, 1008, 662, 662, 144,
        1094, 749, 749, 749, 144,
        756, 756, 914, 914,
    ]
    assert all(row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None for row in key_metrics.rows)
    average_error_run = next(
        run
        for row in key_metrics.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        for run in paragraph.runs
        if run.text == "平均误差"
    )
    fonts = average_error_run._r.get_or_add_rPr().rFonts
    assert fonts.get(qn("w:ascii")) == "Arial"
    assert fonts.get(qn("w:eastAsia")) == "Microsoft YaHei"
    assert 8.5 <= average_error_run.font.size.pt <= 9.0
    assert key_metrics.rows[0].height.inches <= 0.23

    settings = rendered.settings._element
    compatibility_mode = next(
        node
        for node in settings.xpath(".//w:compatSetting")
        if node.get(qn("w:name")) == "compatibilityMode"
    )
    assert compatibility_mode.get(qn("w:val")) == "15"
    theme_language = settings.find(qn("w:themeFontLang"))
    assert theme_language is not None
    assert theme_language.get(qn("w:val")) == "en-US"
    assert theme_language.get(qn("w:eastAsia")) == "zh-CN"

    default_fonts = rendered.styles.element.xpath(
        "./w:docDefaults/w:rPrDefault/w:rPr/w:rFonts"
    )[0]
    assert default_fonts.get(qn("w:ascii")) == "Arial"
    assert default_fonts.get(qn("w:hAnsi")) == "Arial"
    assert default_fonts.get(qn("w:cs")) == "Arial"
    assert default_fonts.get(qn("w:eastAsia")) == "Microsoft YaHei"

    key_metrics_heading = next(
        paragraph for paragraph in rendered.paragraphs
        if paragraph.text == "关键诊断数据"
    )
    assert key_metrics_heading.runs[0].font.size.pt <= 10.5

    priority_issue_run = next(
        run
        for cell in priority_table.rows[1].cells[2:]
        for paragraph in cell.paragraphs
        for run in paragraph.runs
        if run.text.strip()
    )
    assert priority_issue_run.font.size.pt >= 8.5

    device_table = next(
        table for table in rendered.tables
        if len(table.columns) == 4 and "设备 ID / 序列号" in table_text(table)
    )
    assert device_table.rows[0].height.inches <= 0.20
    runtime_table = next(
        table for table in rendered.tables
        if len(table.columns) == 4
        and "累计行程 / 循环次数" in table_text(table)
    )
    runtime_text = table_text(runtime_table)
    for friction_label in ("最大摩擦力", "平均摩擦力", "最小摩擦力"):
        assert friction_label not in runtime_text
    runtime_layout = runtime_table._tbl.tblPr.find(qn("w:tblLayout"))
    assert runtime_layout is not None
    assert runtime_layout.get(qn("w:type")) == "fixed"
    runtime_widths = [
        int(column.get(qn("w:w")))
        for column in runtime_table._tbl.tblGrid.gridCol_lst
    ]
    assert runtime_widths == [1699, 3456, 1699, 3456]

    issue_action_table = next(
        table for table in rendered.tables
        if len(table.columns) == 2
        and table.rows[0].cells[0].text.startswith("主要问题")
        and table.rows[0].cells[1].text.startswith("建议行动")
    )
    assert top_margin(issue_action_table.cell(0, 0)) <= 8
    assert issue_action_table.rows[0].height.inches <= 0.55
    issue_action_widths = [
        int(column.get(qn("w:w")))
        for column in issue_action_table._tbl.tblGrid.gridCol_lst
    ]
    assert issue_action_widths == [3888, 6480]

    recommendation = next(
        paragraph for paragraph in rendered.paragraphs
        if "复测并检查填料" in paragraph.text
    )
    assert recommendation.paragraph_format.space_after.pt <= 4
    assert recommendation.paragraph_format.line_spacing <= 1.05
    recommendation_paragraphs = [
        paragraph
        for paragraph in rendered.paragraphs
        if paragraph.text.startswith("•  ")
    ]
    assert recommendation_paragraphs
    assert all(
        paragraph.paragraph_format.keep_with_next
        for paragraph in recommendation_paragraphs[:-1]
    )
    for heading in (
        "设备概览",
        "主要问题与建议行动",
        "关键诊断数据",
        "运行与定位器诊断信息",
        "整体判断",
        "后续复测、维护和检修建议",
    ):
        heading_paragraph = next(
            paragraph
            for paragraph in rendered.paragraphs
            if paragraph.text == heading
        )
        assert heading_paragraph.paragraph_format.keep_with_next


def test_shared_business_renderer_contains_no_customer_specific_defaults() -> None:
    module = importlib.import_module("offline_engine.business_report")
    source = Path(module.__file__).read_text(encoding="utf-8")

    for forbidden in ("离线诊断/projects/", "/home/openre/"):
        assert forbidden not in source


def test_dashboard_uses_wps_safe_dense_layout() -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    normal = generator.ValveReport(
        folder="FV-NORMAL",
        rating=2,
        final_rating=2,
        issues=[f"问题 {index}" for index in range(3)],
        actions=[f"行动 {index}" for index in range(3)],
    )
    dense = generator.ValveReport(
        folder="FV-DENSE",
        rating=3,
        final_rating=3,
        issues=[f"问题 {index}" for index in range(3)],
        actions=[f"行动 {index}" for index in range(4)],
    )
    doc = Document()
    business_report._setup_document(doc, project_name="Demo")
    business_report._add_dashboard(doc, normal, "Demo")
    business_report._add_dashboard(doc, dense, "Demo")

    def table_text(table) -> str:
        return "\n".join(cell.text for row in table.rows for cell in row.cells)

    def top_margin(cell) -> int:
        margins = cell._tc.get_or_add_tcPr().find(qn("w:tcMar"))
        assert margins is not None
        top = margins.find(qn("w:top"))
        assert top is not None
        return int(top.get(qn("w:w")))

    def tables_for(tag: str):
        start = next(
            index
            for index, table in enumerate(doc.tables)
            if table.cell(0, 0).text.startswith(tag)
        )
        end = next(
            (
                index
                for index, table in enumerate(doc.tables[start + 1 :], start + 1)
                if table.cell(0, 0).text.startswith("FV-")
            ),
            len(doc.tables),
        )
        return doc.tables[start:end]

    normal_tables = tables_for("FV-NORMAL")
    dense_tables = tables_for("FV-DENSE")
    normal_device = next(table for table in normal_tables if "设备 ID / 序列号" in table_text(table))
    dense_device = next(table for table in dense_tables if "设备 ID / 序列号" in table_text(table))
    normal_issues = next(table for table in normal_tables if table.cell(0, 0).text.startswith("主要问题"))
    dense_issues = next(table for table in dense_tables if table.cell(0, 0).text.startswith("主要问题"))
    normal_metrics = next(table for table in normal_tables if "平均误差" in table_text(table))
    dense_metrics = next(table for table in dense_tables if "平均误差" in table_text(table))
    normal_runtime = next(
        table for table in normal_tables
        if "累计行程" in table_text(table)
    )
    dense_runtime = next(
        table for table in dense_tables
        if "累计行程" in table_text(table)
    )

    assert top_margin(normal_tables[0].cell(0, 0)) <= 35
    assert top_margin(dense_tables[0].cell(0, 0)) <= 35
    assert top_margin(normal_device.cell(0, 0)) <= 8
    assert top_margin(dense_device.cell(0, 0)) <= 8
    assert top_margin(normal_issues.cell(0, 0)) <= 8
    assert top_margin(dense_issues.cell(0, 0)) <= 8
    assert top_margin(normal_metrics.cell(0, 0)) <= 20
    assert top_margin(dense_metrics.cell(0, 0)) <= 20
    assert normal_metrics.rows[0].height.inches <= 0.23
    assert dense_metrics.rows[0].height.inches <= 0.23
    assert top_margin(normal_runtime.cell(0, 0)) <= 18
    assert top_margin(dense_runtime.cell(0, 0)) <= 18
    assert normal_runtime.rows[0].height.inches <= 0.21
    assert dense_runtime.rows[0].height.inches <= 0.21

    def bullet_font_sizes(table) -> set[float]:
        return {
            run.font.size.pt
            for cell in table.rows[0].cells
            for paragraph in cell.paragraphs[1:]
            for run in paragraph.runs
            if run.text.strip() and run.font.size is not None
        }

    assert max(bullet_font_sizes(normal_issues)) <= 8.1
    assert min(bullet_font_sizes(normal_issues)) >= 7.8
    assert bullet_font_sizes(dense_issues) == bullet_font_sizes(normal_issues)


def test_business_dashboard_can_leave_conclusions_and_actions_blank(tmp_path: Path) -> None:
    valve = generator.ValveReport(
        folder="FV-BLANK",
        rating=3,
        final_rating=3,
        main_issue="保留的主要问题",
        action="CUSTOM-ACTION-MUST-BE-BLANK",
        issues=["保留的主要问题"],
        actions=["CUSTOM-ACTION-MUST-BE-BLANK"],
        rating_conclusion_override="CUSTOM-CONCLUSION-MUST-BE-BLANK",
    )
    output = tmp_path / "blank-fields.docx"
    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Valve Diagnostic Project",
        report_name="Masoneilan ValVue3 Offline Diagnostic Report",
        report_date="",
        unit_name="Valve Diagnostic Project",
        blank_conclusions=True,
        blank_actions=True,
    )

    rendered = Document(output)
    visible = "\n".join(
        [paragraph.text for paragraph in rendered.paragraphs]
        + [cell.text for table in rendered.tables for row in table.rows for cell in row.cells]
    )
    assert "诊断结论" in visible
    assert "建议行动" in visible
    assert "总体结论与检修建议" in visible
    assert "保留的主要问题" in visible
    assert "CUSTOM-CONCLUSION-MUST-BE-BLANK" not in visible
    assert "CUSTOM-ACTION-MUST-BE-BLANK" not in visible
    assert "测试冻结综合结论。" not in visible
    assert "后续复测、维护和检修建议" not in visible


def test_business_dashboard_review_mode_shows_ratings_actions_and_field_context(
    tmp_path: Path,
) -> None:
    valve = generator.ValveReport(
        folder="FV-REVIEW",
        description="反应进料流量调节阀",
        unit="加氢装置",
        failure_consequence="失效将导致装置停车",
        bypass="无副线",
        site_note="高级诊断未完成",
        rating=4,
        final_rating=4,
        main_issue="保留的主要问题",
        action="ACTION-MUST-BE-VISIBLE",
        issues=["保留的主要问题"],
        actions=["ACTION-MUST-BE-VISIBLE"],
        rating_conclusion_override="CONCLUSION-MUST-BE-BLANK",
    )
    output = tmp_path / "business-review.docx"
    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Valve Diagnostic Project",
        report_name="Masoneilan ValVue3 Offline Diagnostic Report",
        report_date="",
        unit_name="Valve Diagnostic Project",
        review_label="人工校验版",
        blank_conclusions=True,
        include_field_context_overview=True,
    )

    rendered = Document(output)
    visible = "\n".join(
        [paragraph.text for paragraph in rendered.paragraphs]
        + [cell.text for table in rendered.tables for row in table.rows for cell in row.cells]
    )
    assert "人工校验版" in visible
    assert "诊断总览" in visible
    assert "总体结论与检修建议" in visible
    assert "后续复测、维护和检修建议" not in visible
    assert "保留的主要问题" in visible
    assert "反应进料流量调节阀" in visible
    assert "加氢装置" in visible
    assert "失效将导致装置停车" in visible
    assert "无副线" in visible
    assert "高级诊断未完成" in visible
    assert "ACTION-MUST-BE-VISIBLE" in visible
    assert "CONCLUSION-MUST-BE-BLANK" not in visible

    rating_cards = next(
        table for table in rendered.tables
        if any("诊断阀门" in cell.text for cell in table.rows[0].cells)
    )
    assert rating_cards.rows[0].cells[0].text.splitlines()[0] == "1"
    assert rating_cards.rows[0].cells[4].text.splitlines()[0] == "1"

    overview = next(
        table for table in rendered.tables
        if len(table.columns) == 9
        and [cell.text for cell in table.rows[0].cells] == ["序号", "位号", "评级"] * 3
    )
    assert overview.rows[1].cells[1].text == "FV-REVIEW"
    assert overview.rows[1].cells[2].text == "4级"
    assert overview.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None

    field_context = next(
        table for table in rendered.tables
        if len(table.columns) == 6
        and [cell.text for cell in table.rows[0].cells]
        == ["位号", "描述", "装置", "失效后果", "有无副线", "现场备注"]
    )
    assert [cell.text for cell in field_context.rows[1].cells] == [
        "FV-REVIEW",
        "反应进料流量调节阀",
        "加氢装置",
        "失效将导致装置停车",
        "无副线",
        "高级诊断未完成",
    ]

    device_overview = next(
        table for table in rendered.tables
        if len(table.columns) == 4
        and "设备 ID / 序列号" in "\n".join(
            cell.text for row in table.rows for cell in row.cells
        )
    )
    device_text = "\n".join(
        cell.text for row in device_overview.rows for cell in row.cells
    )
    assert "描述" in device_text
    assert "反应进料流量调节阀" in device_text

    priority = next(
        table for table in rendered.tables
        if len(table.columns) == 5
        and [cell.text for cell in table.rows[0].cells]
        == ["评级", "位号", "校验来源", "主要问题", "建议行动"]
    )
    assert len(priority.rows) == 2
    assert [cell.text for cell in priority.rows[1].cells] == [
        "4级",
        "FV-REVIEW",
        "诊断评估",
        "保留的主要问题",
        "ACTION-MUST-BE-VISIBLE",
    ]

    dashboard = next(
        table for table in rendered.tables
        if len(table.columns) == 2 and table.rows[0].cells[0].text.startswith("FV-REVIEW")
    )
    assert dashboard.rows[0].cells[1].text.splitlines()[0] == "4级"
    assert "健康状态" in dashboard.rows[0].cells[1].text


def test_business_dashboard_lists_and_labels_human_verified_low_rating_valve(
    tmp_path: Path,
) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    human = generator.ValveReport(
        folder="FV-HUMAN",
        rating=2,
        final_rating=2,
        main_issue="小开度运行时间较长。",
        action="结合DCS趋势核对开度并检查冲刷。",
        human_verification_status="已纳入人工核对意见",
    )
    automatic = generator.ValveReport(
        folder="FV-AUTO",
        rating=2,
        final_rating=2,
        main_issue="轻微动态偏差。",
        action="计划复核。",
    )
    output = tmp_path / "human-verification.docx"

    business_report.generate_business_dashboard_docx(
        [human, automatic],
        **_prepared_report_text([human, automatic]),
        output_path=output,
        project_name="Demo",
        report_name="Report",
        report_date="2026-07-27",
        unit_name="Demo",
    )

    rendered = Document(output)
    priority = next(
        table for table in rendered.tables
        if [cell.text for cell in table.rows[0].cells]
        == ["评级", "位号", "校验来源", "主要问题", "建议行动"]
    )
    assert len(priority.rows) == 2
    assert [cell.text for cell in priority.rows[1].cells] == [
        "2级",
        "FV-HUMAN",
        "人工校验",
        "小开度运行时间较长。",
        "结合DCS趋势核对开度并检查冲刷。",
    ]
    visible = "\n".join(
        [paragraph.text for paragraph in rendered.paragraphs]
        + [cell.text for table in rendered.tables for row in table.rows for cell in row.cells]
    )
    assert "人工校验状态" not in visible


def test_business_dashboard_can_hide_review_status_and_field_context_and_use_short_bullets(
    tmp_path: Path,
) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    valve = generator.ValveReport(
        folder="FV-PRESENTATION",
        rating=3,
        final_rating=3,
        main_issue="完整主要问题第一句。完整主要问题第二句。",
        action="完整建议行动第一步。完整建议行动第二步。",
        issues=["完整主要问题第一句。", "完整主要问题第二句。"],
        actions=["完整建议行动第一步。", "完整建议行动第二步。"],
        summary_issues=["保留斜坡存在迟滞。", "主因指向性能曲线异常。"],
        summary_actions=["复核填料压紧。", "完成小斜坡复测。"],
        human_verification_status="待完整DOCX统一审核",
        images=[],
    )
    output = tmp_path / "presentation-options.docx"

    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Demo",
        report_name="Ramp 签名 Report",
        report_date="2026-08-03",
        unit_name="Demo",
        include_field_context_overview=False,
        term_replacements={"Ramp": "斜坡", "签名": "性能曲线"},
    )

    rendered = Document(output)
    visible = "\n".join(
        [paragraph.text for paragraph in rendered.paragraphs]
        + [cell.text for table in rendered.tables for row in table.rows for cell in row.cells]
    )
    priority = next(
        table for table in rendered.tables
        if [cell.text for cell in table.rows[0].cells]
        == ["评级", "位号", "校验来源", "主要问题", "建议行动"]
    )

    assert "人工校验状态" not in visible
    assert "现场信息清单" not in visible
    assert "Ramp" not in visible
    assert "签名" not in visible
    assert "斜坡 性能曲线 Report" in visible
    assert priority.rows[1].cells[3].text.splitlines() == [
        "• 保留斜坡存在迟滞。",
        "• 主因指向性能曲线异常。",
    ]
    assert priority.rows[1].cells[4].text.splitlines() == [
        "• 复核填料压紧。",
        "• 完成小斜坡复测。",
    ]
    assert "完整主要问题第一句。" in visible
    assert "完整建议行动第一步。" in visible


def test_business_dashboard_can_remove_unavailable_retest_recommendations(
    tmp_path: Path,
) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    chart = tmp_path / "chart.png"
    chart.write_bytes(PNG_1X1)
    review = CurveInterpretation(
        image_key="FV-NO-RETEST-run-step",
        test_key="run_step",
        observation="阶跃平台存在偏差。",
        diagnosis="定位响应存在异常。",
        rating=3,
        action="检查反馈连杆。",
        customer_text=(
            "阶跃平台存在偏差。定位响应存在异常。"
            "建议先检查反馈连杆，再按原幅值复测并补采数据。"
        ),
        confidence="high",
        experience_rule_ids=(),
    )
    valve = generator.ValveReport(
        folder="FV-NO-RETEST",
        rating=3,
        final_rating=3,
        main_issue="阶跃平台存在偏差。",
        action="检查反馈连杆。",
        issues=["阶跃最大误差7.9%。"],
        actions=["检查反馈连杆及定位器行程标定。"],
        summary_issues=["阶跃最大误差7.9%。"],
        summary_actions=["检查反馈连杆及定位器行程标定。"],
        images=[{
            "path": str(chart),
            "test_key": "run_step",
            "page": 1,
            "curve_review": review,
        }],
    )
    output = tmp_path / "no-retest.docx"

    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Demo",
        report_name="Report",
        report_date="2026-08-03",
        unit_name="Demo",
        recommendations_heading="排查、维护与处置建议",
        include_retest_recommendations=True,
        omit_retest_recommendations=True,
    )

    rendered = Document(output)
    visible = "\n".join(
        [paragraph.text for paragraph in rendered.paragraphs]
        + [cell.text for table in rendered.tables for row in table.rows for cell in row.cells]
    )
    assert "排查、维护与处置建议" in visible
    assert "后续复测、维护和检修建议" not in visible
    assert "复测" not in visible
    assert "补采" not in visible
    assert "重新测试" not in visible
    assert "建议先检查反馈连杆。" in visible


def test_first_valve_starts_with_page_break_before_not_a_hard_break(
    tmp_path: Path,
) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    valve = generator.ValveReport(
        folder="FV-PAGE-BREAK",
        rating=1,
        final_rating=1,
        main_issue="未见明显异常。",
        action="维持现状。",
    )
    output = tmp_path / "page-break-before.docx"

    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Demo",
        report_name="Report",
        report_date="2026-08-03",
        unit_name="Demo",
    )

    rendered = Document(output)
    dashboard = next(
        table for table in rendered.tables
        if len(table.columns) == 2
        and table.cell(0, 0).text.splitlines()[0] == "FV-PAGE-BREAK"
        and "健康状态" in table.cell(0, 1).text
    )
    previous = dashboard._tbl.getprevious()

    assert previous.tag == qn("w:p")
    assert previous.find("./w:pPr/w:pageBreakBefore", previous.nsmap) is not None
    assert previous.find(".//w:br[@w:type='page']", previous.nsmap) is None


def test_business_dashboard_shows_full_conclusions_and_test_execution_issues(
    tmp_path: Path,
) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    valve = generator.ValveReport(
        folder="FV-EVIDENCE",
        rating=3,
        final_rating=3,
        rating_conclusion_override="评级结论：阀门存在响应迟滞。",
        overall_assessment="综合诊断结论：建议结合现场工况计划复核。",
        main_issue="响应迟滞。",
        action="计划复核。",
        issues=["响应迟滞。"],
        actions=["计划复核。"],
        test_execution_issues=[{
            "test_key": "run_step",
            "test_name": "Run Step Test",
            "occurrence": 2,
            "status": "Failed",
        }, {
            "test_key": "performance_check",
            "test_name": "Performance Check",
            "occurrence": 1,
            "status": "Cancelled",
        }, {
            "test_key": "change_mode",
            "test_name": "Change Mode",
            "occurrence": 1,
            "status": "Aborted",
        }],
    )
    output = tmp_path / "full-conclusions-and-test-issues.docx"

    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Demo",
        report_name="Report",
        report_date="2026-07-24",
        unit_name="Demo",
    )

    rendered = Document(output)
    visible = "\n".join(
        [paragraph.text for paragraph in rendered.paragraphs]
        + [cell.text for table in rendered.tables for row in table.rows for cell in row.cells]
    )
    assert "评级结论：阀门存在响应迟滞。" in visible
    assert "综合诊断结论：建议结合现场工况计划复核。" in visible
    assert "测试执行异常（未完成/失败）" in visible
    assert "Run Step Test（第2次）：失败 [Failed]" in visible
    assert "原始报告状态" not in visible
    issue_table = next(
        table
        for table in rendered.tables
        if "Run Step Test（第2次）" in "\n".join(
            cell.text for row in table.rows for cell in row.cells
        )
    )
    assert len(issue_table.cell(0, 0).paragraphs) == 1


def test_business_dashboard_summary_only_omits_repeated_comprehensive_conclusion(
    tmp_path: Path,
) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    valve = generator.ValveReport(
        folder="FV-SUMMARY",
        rating=3,
        final_rating=3,
        rating_conclusion_override="阀门存在响应迟滞，评为3级。",
        overall_assessment="阀门存在响应迟滞，评为3级。主要问题：开向响应迟滞。",
        health_rating_detail_mode="summary_only",
        main_issue="开向响应迟滞。",
        action="计划复核。",
        issues=["开向响应迟滞。"],
        actions=["计划复核。"],
    )
    output = tmp_path / "summary-only.docx"

    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Demo",
        report_name="Report",
        report_date="2026-07-24",
        unit_name="Demo",
    )

    rendered = Document(output)
    visible = "\n".join(
        cell.text
        for table in rendered.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "阀门存在响应迟滞，评为3级。" in visible
    assert "综合诊断结论" not in visible
    assert "主要问题：开向响应迟滞" in visible


def test_business_dashboard_data_only_keeps_facts_and_removes_assessments(
    tmp_path: Path,
) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    chart = tmp_path / "chart.png"
    chart.write_bytes(PNG_1X1)
    review = CurveInterpretation(
        image_key="FV-DATA-run-step",
        test_key="run_step",
        observation="曲线存在迟滞。",
        diagnosis="需要检修。",
        rating=4,
        action="安排检修。",
        customer_text="INTERPRETATION-MUST-NOT-BE-VISIBLE",
        confidence="high",
        experience_rule_ids=(),
    )
    valve = generator.ValveReport(
        folder="FV-DATA",
        description="进料调节阀",
        unit="加氢装置",
        failure_consequence="进料波动",
        bypass="无副线",
        site_note="现场记录",
        rating=4,
        final_rating=4,
        rating_conclusion_override="CONCLUSION-MUST-NOT-BE-VISIBLE",
        overall_assessment="ASSESSMENT-MUST-NOT-BE-VISIBLE",
        main_issue="ISSUE-MUST-NOT-BE-VISIBLE",
        action="ACTION-MUST-NOT-BE-VISIBLE",
        issues=["ISSUE-MUST-NOT-BE-VISIBLE"],
        actions=["ACTION-MUST-NOT-BE-VISIBLE"],
        test_execution_issues=[{
            "test_key": "run_step",
            "test_name": "Run Step Test",
            "occurrence": 1,
            "status": "Failed",
        }],
        metrics={
            "Average Error": 3.2,
            "T86": 1.5,
            "T63": 0.8,
            "Td": 0.2,
            "Overshoot": 4.1,
            "Maximum Hysteresis + Dead Band": 2.0,
            "Average Hysteresis + Dead Band": 1.0,
            "Max Friction": 12.3,
            "Mean Friction": 6.2,
            "Supply Pressure": {"min": 50.0, "max": 55.0, "avg": 52.5},
        },
        positioner_report_info={
            "active_faults": ["Actuator Error"],
            "historical_faults": ["Position Error"],
            "strokes": "401",
            "cycles": "37605",
        },
        images=[{
            "path": chart,
            "page": 3,
            "test_key": "run_step",
            "curve_review": review,
        }],
        curve_interpretations=(review,),
    )
    output = tmp_path / "pure-data.docx"

    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Demo",
        report_name="Masoneilan ValVue3 离线检测数据报告",
        report_date="2026-07-24",
        report_period_label="2026年7月22日至今",
        review_label="纯数据版",
        unit_name="Demo",
        data_only=True,
        blank_ratings=True,
        blank_conclusions=True,
        blank_actions=True,
    )

    rendered = Document(output)
    visible = "\n".join(
        [paragraph.text for paragraph in rendered.paragraphs]
        + [cell.text for table in rendered.tables for row in table.rows for cell in row.cells]
    )
    for required in (
        "纯数据版",
        "离线检测数据报告",
        "设备与检测信息总览",
        "进料调节阀",
        "加氢装置",
        "进料波动",
        "无副线",
        "现场记录",
        "Run Step Test：失败 [Failed]",
        "平均误差",
        "3.2%",
        "Actuator Error",
        "测试曲线与原始数据",
        "图形判读",
        "原始报告第 3 页",
    ):
        assert required in visible
    for forbidden in (
        "诊断总览",
        "总体结论与检修建议",
        "健康状态",
        "诊断结论",
        "主要问题",
        "建议行动",
        "4级",
        "CONCLUSION-MUST-NOT-BE-VISIBLE",
        "ASSESSMENT-MUST-NOT-BE-VISIBLE",
        "ISSUE-MUST-NOT-BE-VISIBLE",
        "ACTION-MUST-NOT-BE-VISIBLE",
        "INTERPRETATION-MUST-NOT-BE-VISIBLE",
    ):
        assert forbidden not in visible

    title = next(
        table
        for table in rendered.tables
        if table.rows[0].cells[0].text.startswith("FV-DATA")
    )
    fills = [
        cell._tc.get_or_add_tcPr().find(qn("w:shd")).get(qn("w:fill"))
        for cell in title.rows[0].cells
    ]
    assert fills == [business_report.BRAND_PRIMARY, business_report.BRAND_PRIMARY]

    with ZipFile(output) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert curve_review_marker(review.image_key) in document_xml


def test_data_only_no_curve_notice_starts_on_a_new_page() -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    valve = generator.ValveReport(folder="FV-NO-CURVE")
    doc = Document()

    business_report._add_curve_pages(doc, valve, data_only=True)

    heading = next(
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text == "FV-NO-CURVE | 测试曲线与原始数据"
    )
    assert heading.paragraph_format.page_break_before


def test_business_no_curve_notice_starts_on_a_new_page() -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    valve = generator.ValveReport(folder="FV-NO-CURVE")
    doc = Document()

    business_report._add_curve_pages(doc, valve, data_only=False)

    heading = next(
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text == "FV-NO-CURVE | 曲线证据与图形判读"
    )
    assert heading.paragraph_format.page_break_before
    callout_text = "\n".join(
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "数据状态" in callout_text
    assert "判读限制" in callout_text
    assert "后续处理" in callout_text


def test_business_dashboard_keeps_auto_tune_fault_on_one_compact_token(
    tmp_path: Path,
) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    valve = generator.ValveReport(
        folder="FV-AUTO-TUNE",
        rating=4,
        final_rating=4,
        positioner_report_info={
            "historical_faults": ["Low Power", "Auto tune Error"],
        },
    )
    output = tmp_path / "auto-tune.docx"

    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Demo",
        report_name="Report",
        report_date="2026-07-24",
        unit_name="Demo",
    )

    rendered = Document(output)
    visible = "\n".join(
        cell.text for table in rendered.tables for row in table.rows for cell in row.cells
    )
    assert "Autotune Error" in visible
    assert "Auto tune Error" not in visible

    manifest = business_report.report_field_manifest(
        valve,
        {"source_record": {"historical_faults": ["Low Power", "Auto tune Error"]}},
    )
    historical = next(
        item for item in manifest if item["path"] == "status.historical_faults"
    )
    assert historical["display_text"] == "Low Power；Autotune Error"


def test_business_dashboard_compacts_conclusions_and_runtime_into_single_tables() -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    valve = generator.ValveReport(
        folder="FV-COMPACT",
        rating=3,
        final_rating=3,
        rating_conclusion_override="评级结论：阀门存在响应迟滞。",
        overall_assessment="综合诊断结论：建议结合现场工况计划复核。",
        main_issue="响应迟滞。",
        action="计划复核。",
        issues=["响应迟滞。"],
        actions=["计划复核。"],
        positioner_report_info={
            "strokes": "401",
            "cycles": "37605",
            "time_open": "32886",
            "time_closed": "37362",
            "time_near_closed": "0",
            "valve_position": "50.0 %",
            "manual_setpoint": "50.0 %",
            "input_signal": "12.0 mA",
            "supply_pressure": "380 kPa",
            "active_faults": ["Actuator Error"],
            "historical_faults": ["Position Error"],
        },
    )
    doc = Document()
    business_report._setup_document(doc, project_name="Demo")
    business_report._add_dashboard(
        doc, valve, "Demo", rating_conclusion=valve.rating_conclusion_override
    )

    table_texts = [
        "\n".join(cell.text for row in table.rows for cell in row.cells)
        for table in doc.tables
    ]
    conclusion_tables = [
        text
        for text in table_texts
        if "评级结论：阀门存在响应迟滞。" in text
        or "综合诊断结论：建议结合现场工况计划复核。" in text
    ]
    runtime_tables = [
        text
        for text in table_texts
        if "运行与故障摘要" in text
    ]
    key_metric_tables = [
        text
        for text in table_texts
        if all(
            field in text
            for field in ("阶跃测试", "斜坡测试", "执行机构性能测试")
        )
    ]

    assert len(conclusion_tables) == 1
    assert "评级结论：阀门存在响应迟滞。" in conclusion_tables[0]
    assert "综合诊断结论：建议结合现场工况计划复核。" in conclusion_tables[0]
    assert len(runtime_tables) == 1
    for required in ("活动故障", "快照", "历史故障"):
        assert required in runtime_tables[0]
    for friction_label in ("最大摩擦力", "平均摩擦力", "最小摩擦力"):
        assert friction_label not in runtime_tables[0]
    assert len(key_metric_tables) == 1
    for friction_label in ("最大摩擦力", "平均摩擦力", "最小摩擦力"):
        assert friction_label in key_metric_tables[0]
    runtime_table = next(
        table
        for table in doc.tables
        if "运行与故障摘要" in "\n".join(
            cell.text for row in table.rows for cell in row.cells
        )
    )
    assert len(runtime_table.rows) <= 4
    for index, row in enumerate(runtime_table.rows):
        assert row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                assert paragraph.paragraph_format.keep_with_next == (
                    index < len(runtime_table.rows) - 1
                )


def test_business_dashboard_preserves_human_report_level_text(tmp_path: Path) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    valve = generator.ValveReport(folder="FV-101", rating=2, final_rating=2)
    output = tmp_path / "human-text.docx"

    business_report.generate_business_dashboard_docx(
        [valve],
        rating_conclusions={"FV-101": "人工逐阀评级结论原文。"},
        overall_conclusion="人工总体结论原文。",
        recommendations=["人工建议一。", "人工建议二。"],
        output_path=output,
        project_name="Demo",
        report_name="Report",
        report_date="2026-07-23",
        unit_name="Demo",
        include_retest_recommendations=True,
    )

    rendered = Document(output)
    visible = "\n".join(
        [paragraph.text for paragraph in rendered.paragraphs]
        + [cell.text for table in rendered.tables for row in table.rows for cell in row.cells]
    )
    assert "人工总体结论原文。" in visible
    assert "人工建议一。" in visible
    assert "人工建议二。" in visible
    assert "测试冻结综合结论。" not in visible


def test_overall_priority_actions_render_compound_actions_as_complete_bullets(
    tmp_path: Path,
) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    valve = generator.ValveReport(
        folder="FV-ACTION-SUMMARY",
        rating=4,
        final_rating=4,
        main_issue="供气与机械响应需复核。",
        action=(
            "优先安排现场复核和窗口期检修；"
            "检查过滤减压阀、气路泄漏和执行机构密封；"
            "处理后复测扩展签名和阶跃，确认响应改善。"
        ),
    )
    output = tmp_path / "action-summary.docx"

    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Demo",
        report_name="Report",
        report_date="2026-08-26",
        unit_name="Demo",
    )

    rendered = Document(output)
    priority = next(
        table
        for table in rendered.tables
        if [cell.text for cell in table.rows[0].cells]
        == ["评级", "位号", "校验来源", "主要问题", "建议行动"]
    )
    action_lines = priority.rows[1].cells[4].text.splitlines()
    assert action_lines == [
        "• 优先安排现场复核和窗口期检修。",
        "• 检查过滤减压阀、气路泄漏和执行机构密封。",
        "• 处理后复测执行机构性能测试和阶跃，确认响应改善。",
    ]


def test_business_dashboard_sequencer_template_contract_and_attachment_dedup(
    tmp_path: Path,
) -> None:
    business_report = importlib.import_module("offline_engine.business_report")

    def png(path: Path, color: tuple[int, int, int]) -> None:
        Image.new("RGB", (8, 8), color).save(path)

    spec = tmp_path / "spec.png"
    air = tmp_path / "air.png"
    accessory = tmp_path / "accessory.png"
    png(spec, (10, 20, 30))
    png(air, (40, 50, 60))
    png(accessory, (70, 80, 90))
    valve = generator.ValveReport(
        folder="XV-001",
        description="",
        unit="",
        failure_consequence="",
        bypass="",
        site_note="",
        rating=2,
        final_rating=2,
        positioner_report_info={
            "firmware": "4.1",
            "diagnostics_level": "Advanced",
            "option_config": "Advanced Diagnostics; Double Acting;",
            "actuator_style": "Double",
            "air_action": "ATO",
            "file": {
                "report_start_time": "8/26/2026 9:45:45 AM",
                "report_end_time": "8/26/2026 10:21:24 AM",
            },
            "positioner": {
                "option_config": "Advanced Diagnostics; Double Acting;",
                "device_mode": "Normal",
                "valve_position": "-0.0",
                "manual_setpoint": "15.0",
                "input_signal": "4.00",
                "board_temperature": "31.2",
                "min_temperature": "-3.9",
                "max_temperature": "48.5",
            },
            "runtime": {
                "strokes": "241",
                "cycles": "14691",
                "time_open": "38479",
                "time_closed": "10811",
                "time_near_closed": "6319",
            },
            "status": {
                "active_faults": ["Low Air Supply Warning"],
                "historical_faults": ["Position Error"],
            },
        },
        metrics={
            "run_ramp": {
                "hysteresis": {"max": 1.73, "average": 0.87, "min": 0.0},
                "dead_band": {"max": 1.64, "average": 0.88, "min": 0.0},
                "linearity": {"max": 1.06, "average": 0.56, "min": 0.02},
            }
        },
    )
    valve.report_pages = [
        generator.SpecPage(
            valve="XV-001", tag="XV-001", source="specs/form.pdf",
            source_sha256="a" * 64, page=2, roles=("specification",),
            image_path=str(spec),
        ),
        generator.SpecPage(
            valve="XV-001", tag="XV-001", source="specs/form.pdf",
            source_sha256="a" * 64, page=5, roles=("air_piping",),
            image_path=str(air),
        ),
        generator.SpecPage(
            valve="XV-001", tag="XV-001", source="specs/form.pdf",
            source_sha256="a" * 64, page=6, roles=("accessory",),
            image_path=str(accessory),
        ),
    ]
    output = tmp_path / "sequencer-contract.docx"
    business_report.generate_business_dashboard_docx(
        [valve], **_prepared_report_text([valve]), output_path=output, project_name="Customer Name",
        report_name="Report", report_date="2026-08-26", unit_name="Customer Name",
    )

    rendered = Document(output)
    visible = "\n".join(
        [paragraph.text for paragraph in rendered.paragraphs]
        + [cell.text for table in rendered.tables for row in table.rows for cell in row.cells]
    )
    for forbidden in (
        "现场信息清单", "整体判断", "首要关注",
        "后续复测、维护和检修建议", "运行与 DTM 信息",
        "Maximum Hysteresis + Dead Band", "XV-001 | Accessory",
    ):
        assert forbidden not in visible
    for required in (
        "总体结论与检修建议", "建议优先处理清单", "设备概览", "待业主补充",
        "8/26/2026 9:45:45 AM – 8/26/2026 10:21:24 AM",
        "运行与定位器诊断信息", "累计行程次数", "累计循环次数",
        "累计全开时间", "累计全关时间", "累计近关时间",
        "Low Air Supply Warning", "Position Error", "双作用 / 气开",
        "阶跃测试", "斜坡测试", "执行机构性能测试",
        "斜坡指标", "回差", "死区", "线性度",
        "1.73%", "0.87%", "1.64%", "0.88%", "1.06%", "0.56%", "0.02%",
        "XV-001 | Specification", "XV-001 | Air Piping",
    ):
        assert required in visible
    device_table = next(
        table for table in rendered.tables
        if any(cell.text == "装置" for row in table.rows for cell in row.cells)
    )
    device_cells = [cell.text for row in device_table.rows for cell in row.cells]
    assert device_cells[device_cells.index("装置") + 1] == "待业主补充"
    with ZipFile(output) as archive:
        assert len([name for name in archive.namelist() if name.startswith("word/media/")]) == 2

    shared = tmp_path / "shared.png"
    png(shared, (100, 110, 120))
    valve.report_pages = [
        generator.SpecPage(
            valve="XV-001", tag="XV-001", source="specs/shared.pdf",
            source_sha256="b" * 64, page=3, roles=("specification",),
            image_path=str(shared),
        ),
        generator.SpecPage(
            valve="XV-001", tag="XV-001", source="specs/shared.pdf",
            source_sha256="b" * 64, page=3, roles=("air_piping",),
            image_path=str(shared),
        ),
    ]
    deduped = tmp_path / "deduped.docx"
    business_report.generate_business_dashboard_docx(
        [valve], **_prepared_report_text([valve]), output_path=deduped, project_name="Demo", report_name="Report",
        report_date="2026-08-26", unit_name="Demo",
    )
    deduped_doc = Document(deduped)
    bookmarks = deduped_doc._element.body.xpath(".//w:bookmarkStart/@w:name")
    assert spec_insertion_marker(valve.folder, "specification") in bookmarks
    assert spec_insertion_marker(valve.folder, "air_piping") in bookmarks
    assert generator.report_page_insertion_markers(
        valve.folder, valve.report_pages[0]
    ) == {
        "specification": spec_insertion_marker(
            valve.folder, "specification"
        ),
    }
    assert generator.report_page_insertion_markers(
        valve.folder, valve.report_pages[1]
    ) == {
        "air_piping": spec_insertion_marker(valve.folder, "air_piping"),
    }
    with ZipFile(deduped) as archive:
        assert len([name for name in archive.namelist() if name.startswith("word/media/")]) == 1


def test_business_dashboard_reads_nested_structured_evidence_without_whole_valve_text(
    tmp_path: Path,
) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    valve = generator.ValveReport(
        folder="XV-STRUCTURED",
        rating=2,
        final_rating=2,
        positioner_report_info={
            "file": {
                "report_start_time": "8/26/2026 9:45:45 AM",
                "report_end_time": "8/26/2026 10:21:24 AM",
            },
            "positioner": {
                "firmware": "4.1",
                "diagnostics_level": "Advanced",
                "option_config": "Advanced Diagnostics; Double Acting;",
                "actuator_style": "Double Acting",
                "air_action": "ATO",
                "device_mode": "Normal",
                "valve_position": "0.0 %",
                "manual_setpoint": "15.0 %",
                "input_signal": "4.00 mA",
                "board_temperature": "31.2 degC",
                "min_temperature": "-3.9 degC",
                "max_temperature": "48.5 degC",
            },
            "runtime": {
                "strokes": "241",
                "cycles": "14691",
                "time_open": "38479",
                "time_closed": "10811",
                "time_near_closed": "6319",
            },
            "status": {
                "active_faults": ["Reset"],
                "historical_faults": ["Position Error"],
                "task_execution": [
                    {
                        "task_name": "Run Ramp Test",
                        "test_key": "run_ramp",
                        "occurrence": 1,
                        "result": "Succeeded",
                    },
                    {
                        "task_name": "Extended Actuator Signature",
                        "test_key": "extended_signature",
                        "occurrence": 1,
                        "result": "Failed",
                    },
                ],
            },
        },
        metrics={
            "run_ramp": {
                "hysteresis": {"max": 1.73, "average": 0.87, "min": 0.0},
                "dead_band": {"max": 1.64, "average": 0.88, "min": 0.0},
                "linearity": {"max": 1.06, "average": 0.56, "min": 0.02},
            }
        },
    )

    output = tmp_path / "nested-structured.docx"
    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Demo",
        report_name="Report",
        report_date="2026-08-26",
        unit_name="Demo",
    )
    rendered = Document(output)
    visible = "\n".join(
        [paragraph.text for paragraph in rendered.paragraphs]
        + [
            cell.text
            for table in rendered.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    for required in (
        "4.1 / 未提供",
        "Advanced / Advanced Diagnostics; Double Acting;",
        "双作用 / 气开",
        "8/26/2026 9:45:45 AM – 8/26/2026 10:21:24 AM",
        "累计行程次数",
        "241 counts",
        "累计全开时间",
        "38479 h",
        "Position Error",
        "执行机构性能测试：失败",
        "失败 / 无有效曲线",
        "阶跃测试",
        "斜坡测试",
        "执行机构性能测试",
        "回差",
        "线性度",
        "1.73%",
        "0.02%",
    ):
        assert required in visible
    assert "Reset" not in visible
    assert "Run Ramp Test：成功" not in visible


def test_formal_summary_report_field_manifest_matches_docx_text(
    tmp_path: Path, monkeypatch
) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    positioner = {
        "option_config": "Advanced Diagnostics; Double Acting;",
        "device_mode": "Normal",
        "valve_position": "0.0 %",
        "manual_setpoint": "15.0 %",
        "input_signal": "4.00 mA",
        "board_temperature": "31.2 degC",
        "min_temperature": "-3.9 degC",
        "max_temperature": "48.5 degC",
    }
    runtime = {
        "strokes": "241",
        "cycles": "14691",
        "time_open": "38479",
        "time_closed": "10811",
        "time_near_closed": "6319",
    }
    status = {
        "active_faults": [],
        "historical_faults": ["Position Error"],
        "task_execution": [
            {
                "task_name": "Extended Actuator Signature",
                "test_key": "extended_signature",
                "occurrence": 1,
                "start_time": "8/26/2026 10:00:00 AM",
                "end_time": "8/26/2026 10:05:00 AM",
                "duration_seconds": 300,
                "result": "Failed",
            }
        ],
    }
    metrics = {
        "run_ramp": {
            "hysteresis": {"max": 1.73, "average": 0.87, "min": 0.0},
            "dead_band": {"max": 1.64, "average": 0.88, "min": 0.0},
            "linearity": {"max": 1.06, "average": 0.56, "min": 0.02},
        }
    }
    info = {
        "file": {
            "report_start_time": "8/26/2026 9:45:45 AM",
            "report_end_time": "8/26/2026 10:21:24 AM",
            "report_time": "2026-08-26 09:45:45–10:21:24",
        },
        "positioner": positioner,
        "runtime": runtime,
        "status": status,
    }
    valve = generator.ValveReport(
        folder="XV-MANIFEST",
        rating=2,
        final_rating=2,
        positioner_report_info=info,
        metrics=metrics,
    )
    valve.structured_evidence = {
        "schema_version": 1,
        "source_record": {
            "positioner_report_info": info,
            "metrics": metrics,
        },
        "source_files": [],
        "evidence_pages": [],
        "material_extraction": {"fields": {}},
    }
    summary_path = tmp_path / "analysis_summary.json"
    monkeypatch.setattr(generator, "SUMMARY_JSON", summary_path)
    generator.write_summary(
        [valve], report_field_manifest_builder=business_report.report_field_manifest
    )
    summary = __import__("json").loads(summary_path.read_text(encoding="utf-8"))
    entries = summary[0]["report_field_manifest"]
    assert {entry["path"] for entry in entries} == set(
        business_report.REPORT_FIELD_MANIFEST_PATHS
    )

    output = tmp_path / "manifest.docx"
    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Demo",
        report_name="Report",
        report_date="2026-08-26",
        unit_name="Demo",
    )
    rendered = Document(output)
    visible = "\n".join(
        [paragraph.text for paragraph in rendered.paragraphs]
        + [
            cell.text
            for table in rendered.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    assert all(entry["display_text"] in visible for entry in entries)
    assert entries[-1]["display_text"] == "0.02%"
    assert "执行机构性能测试：失败" in visible

    from validate_valvue3_report import docx_text, report_field_manifest_failures

    docx_ok, docx_body, docx_error = docx_text(output)
    assert docx_ok, docx_error
    assert report_field_manifest_failures(summary_path, docx_body) == (1, [])


def test_report_field_manifest_reads_authoritative_nested_material_fields(
    tmp_path: Path,
) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    info = {
        "file": {
            "report_start_time": "8/26/2026 9:45:45 AM",
            "report_end_time": "8/26/2026 10:21:24 AM",
            "report_time": "2026-08-26 09:45:45–10:21:24",
        },
        "positioner": {
            "option_config": "Advanced Diagnostics; Double Acting;",
            "device_mode": "Normal",
            "valve_position": 0.0,
            "manual_setpoint": 15.0,
            "input_signal": 4.0,
            "board_temperature": 31.2,
            "min_temperature": -3.9,
            "max_temperature": 48.5,
        },
        "runtime": {
            "strokes": 241,
            "cycles": 14691,
            "time_open": 38479,
            "time_closed": 10811,
            "time_near_closed": 6319,
        },
        "status": {
            "active_faults": [],
            "historical_faults": ["Position Error"],
            "task_execution": [
                {
                    "task_name": "Extended Actuator Signature",
                    "test_key": "extended_signature",
                    "occurrence": 1,
                    "start_time": "8/26/2026 10:00:00 AM",
                    "end_time": "8/26/2026 10:05:00 AM",
                    "result": "Failed",
                },
                {
                    "task_name": "Change Mode",
                    "test_key": "change_mode",
                    "occurrence": 1,
                    "start_time": "8/26/2026 10:05:00 AM",
                    "end_time": "8/26/2026 10:05:03 AM",
                    "result": "Failed",
                }
            ],
        },
    }
    metrics = {
        "run_ramp": {
            "hysteresis": {"max": 1.73, "average": 0.87, "min": 0.0},
            "dead_band": {"max": 1.64, "average": 0.88, "min": 0.0},
            "linearity": {"max": 1.06, "average": 0.56, "min": 0.02},
        }
    }
    valve = generator.ValveReport(
        folder="XV-NESTED",
        positioner_report_info=info,
        metrics=metrics,
    )
    values = {
        **{
            f"positioner.{key}": value
            for key, value in info["positioner"].items()
        },
        **{
            f"runtime.{key}": value
            for key, value in info["runtime"].items()
        },
        **{
            f"file.{key}": value
            for key, value in info["file"].items()
        },
        **{
            f"status.{key}": value
            for key, value in info["status"].items()
        },
        **{
            f"metrics.run_ramp.{metric}.{stat}": value
            for metric, stats in metrics["run_ramp"].items()
            for stat, value in stats.items()
        },
    }
    structured = {
        "schema_version": 1,
        "source_record": {
            "positioner_report_info": {
                **info,
                "positioner": {
                    **info["positioner"],
                    "manual_setpoint": 99.0,
                },
            },
            "metrics": {
                "run_ramp": {
                    **metrics["run_ramp"],
                    "hysteresis": {
                        **metrics["run_ramp"]["hysteresis"],
                        "min": 9.0,
                    },
                }
            },
        },
        "material_extraction": {
            "fields": {
                path: {"selected": {"normalized_value": value}}
                for path, value in values.items()
            }
        },
        "evidence_pages": [],
        "source_files": [],
    }

    manifest = business_report.report_field_manifest(valve, structured)
    assert len(manifest) == len(business_report.REPORT_FIELD_MANIFEST_PATHS)
    by_path = {entry["path"]: entry for entry in manifest}
    assert all(entry["evidence_value"] is not None for entry in manifest)
    assert by_path["positioner.manual_setpoint"]["evidence_value"] == 15.0
    assert by_path["metrics.run_ramp.hysteresis.min"]["evidence_value"] == 0.0
    assert by_path["status.active_faults"]["evidence_value"] == []
    assert by_path["status.task_execution"]["evidence_value"][0]["result"] == "Failed"
    expected_task_display = (
        "•  执行机构性能测试：失败 [Failed]；  "
        "•  Change Mode：失败 [Failed]"
    )
    assert by_path["status.task_execution"]["display_text"] == expected_task_display

    output = tmp_path / "report-field-manifest-task-display.docx"
    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Demo",
        report_name="Report",
        report_date="2026-08-26",
        unit_name="Demo",
    )
    rendered = Document(output)
    visible = "\n".join(
        [paragraph.text for paragraph in rendered.paragraphs]
        + [
            cell.text
            for table in rendered.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    assert expected_task_display in visible


def test_report_field_manifest_binds_combined_detection_window_as_end_time() -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    valve = generator.ValveReport(
        folder="FV-TIME",
        detection_date="2026-08-21 16:10:21–16:42:49",
    )
    structured = {
        "source_record": {
            "detection_time": "2026-08-21 16:10:21–16:42:49",
        },
        "material_extraction": {
            "fields": {
                "file.report_end_time": {
                    "selected": None,
                    "status": "missing",
                    "candidates": [],
                }
            }
        },
    }

    by_path = {
        entry["path"]: entry
        for entry in business_report.report_field_manifest(valve, structured)
    }

    assert by_path["file.report_end_time"] == {
        "path": "file.report_end_time",
        "evidence_value": "2026-08-21 16:10:21–16:42:49",
        "display_text": "2026-08-21 16:10:21–16:42:49",
    }


def test_repeated_curve_types_show_each_execution_window(tmp_path: Path) -> None:
    business_report = importlib.import_module("offline_engine.business_report")
    charts = [tmp_path / "first.png", tmp_path / "second.png"]
    for chart in charts:
        chart.write_bytes(PNG_1X1)
    images = []
    reviews = []
    windows = [
        ("8/21/2026 5:56:45 PM", "8/21/2026 5:58:37 PM"),
        ("8/25/2026 4:14:31 PM", "8/25/2026 4:16:23 PM"),
    ]
    for index, (chart, (start, end)) in enumerate(zip(charts, windows, strict=True), 1):
        review = CurveInterpretation(
            image_key=f"FV-TIME-{index}",
            test_key="run_step",
            observation="曲线已查看。",
            diagnosis="响应可复核。",
            rating=2,
            action="按计划复测。",
            customer_text="曲线已查看，建议按计划复测。",
            confidence="high",
            experience_rule_ids=(),
        )
        reviews.append(review)
        images.append({
            "path": str(chart),
            "test_key": "run_step",
            "page": index,
            "test_start_time": start,
            "test_end_time": end,
            "curve_review": review,
        })
    valve = generator.ValveReport(
        folder="FV-TIME",
        rating=2,
        final_rating=2,
        images=images,
        curve_interpretations=tuple(reviews),
    )
    output = tmp_path / "curve-times.docx"

    business_report.generate_business_dashboard_docx(
        [valve],
        **_prepared_report_text([valve]),
        output_path=output,
        project_name="Demo",
        report_name="Report",
        report_date="2026-08-26",
        unit_name="Demo",
    )

    rendered = Document(output)
    visible = "\n".join(
        cell.text
        for table in rendered.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "检测时间：2026-08-21 17:56:45–17:58:37" in visible
    assert "检测时间：2026-08-25 16:14:31–16:16:23" in visible
